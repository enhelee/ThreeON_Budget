import openpyxl
import pandas as pd
from budget import excel_plan_writer


def test_write_datasheet_basic_and_missing_highlight():
    df = pd.DataFrame([
        {"주관부서명": "A부", "부서코드": "1", "처지사": "화성지사", "부서부": "b",
         "속성": "제조", "예산코드": "1", "예산과목": "수선유지비-건물/구축물",
         "사업명": "공사1", "산출내역": "-", "연예산": 60000.0, "_row": 4},
        {"주관부서명": "B부", "부서코드": "2", "처지사": None, "부서부": "b",
         "속성": "제조", "예산코드": "2", "예산과목": "수선유지비-열원경상정비",
         "사업명": "공사2", "산출내역": "-", "연예산": 30000.0, "_row": 5},
    ])
    item_flags = {"수선유지비-건물/구축물": True, "수선유지비-열원경상정비": False}
    wb = openpyxl.Workbook()
    ws = wb.active
    result = excel_plan_writer.write_datasheet(ws, df, "손익", "2026", item_flags)
    assert result["written_rows"] == 2
    assert ws.cell(3, 7).value == "예산과목"
    assert ws.cell(4, 7).value == "수선유지비-건물/구축물"
    assert ws.cell(4, 12).value == "O"   # 심의대상 + 60000>=50000
    assert ws.cell(5, 12).value == ""    # 심의대상 아님
    assert ws.cell(5, 11).value == "처지사 누락"
    assert ws.cell(5, 3).fill.start_color.rgb == "00FFC7CE"
