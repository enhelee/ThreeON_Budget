# -*- coding: utf-8 -*-
"""실적 산출물 작성: {연도}년 {손익|자본}예산_실적.xlsx + zrfm2_V1.

- 양식1(월별): 계획행(실적금액·구분 태깅) + 신규행 하단 추가.
- 종합표: 계획본과 동일 레이아웃, SUMIFS로 실적금액(K열)을 지사별 집계.
- 검토리포트: 미매핑 처지사, 미분류 과목, 연도 불일치 등.
- zrfm2_V1: 원본 유지 + R열에 매칭된 계획/신규 사업명.
"""
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from . import loaders
from .config_store import SIMUI_THRESHOLD_THOUSAND
from . import excel_plan_writer as epw

ACTUAL_HEADER = [
    "주관부서명", "예산귀속 부서코드", "예산귀속 부서명(처.지사)", "예산귀속 부서명(부)",
    "속성", "예산코드", "예산과목", "사업명", "산출내역", "연예산(계획)",
    "최종 실적금액", "심의플래그", "구분", "비고",
]
VALUE_COL_LETTER = "K"   # 11번 = 최종 실적금액
NEW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # 신규(노랑)
MISSING_EXEC_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")  # 미시행(연빨강)

HEADER_ROW = epw.HEADER_ROW  # 3
DATA_START = 4


def _simui_flag(item_flags, item, base_amount):
    return "O" if item_flags.get(item) and base_amount is not None \
        and base_amount >= SIMUI_THRESHOLD_THOUSAND else ""


def write_actual_datasheet(ws, plan_rows, new_rows, budget, year, item_flags):
    ws.cell(1, 1, f"{year}년 {budget}예산 실적 (단위: 천원)")
    for i, h in enumerate(ACTUAL_HEADER, 1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = epw.HEADER_FILL

    r = DATA_START
    # 계획행
    for row in plan_rows:
        for i, key in enumerate(loaders.PLAN_COLUMNS, 1):
            val = row.get(key)
            ws.cell(r, i, None if (val is None or (isinstance(val, float) and pd.isna(val))) else val)
        ws.cell(r, 11, round(row.get("실적금액", 0.0)))
        flag = _simui_flag(item_flags, row.get("예산과목"), row.get("연예산"))
        ws.cell(r, 12, flag)
        ws.cell(r, 13, row.get("구분"))
        if row.get("구분") == "계획집행":
            note = f"매칭전표 {row.get('매칭전표수', 0)}건"
            low = row.get("저유사전표수", 0)
            if low:
                note += f" (저유사 {low}건)"
        elif row.get("구분") == "미시행":
            note = "계획O·실적X (실적 없음)"
        else:
            note = ""
        ws.cell(r, 14, note)
        if row.get("구분") == "미시행":
            for c in range(1, 15):
                ws.cell(r, c).fill = MISSING_EXEC_FILL
        r += 1

    # 신규행
    for nr in new_rows:
        ws.cell(r, 3, nr["처지사"])
        ws.cell(r, 7, nr["예산과목"])
        ws.cell(r, 8, nr["사업명"])
        ws.cell(r, 10, None)
        ws.cell(r, 11, round(nr["실적금액"]))
        ws.cell(r, 12, _simui_flag(item_flags, nr["예산과목"], nr["실적금액"]))
        ws.cell(r, 13, nr.get("구분", "신규"))
        ws.cell(r, 14, nr.get("비고", ""))
        for c in range(1, 15):
            ws.cell(r, c).fill = NEW_FILL
        r += 1

    return {"written_rows": r - DATA_START}


def write_review_sheet(wb, review):
    ws = wb.create_sheet("검토리포트")
    ws.append(["■ 실적분석 검토리포트"])
    ws.append([])
    ws.append(["대상연도", review.get("year")])
    ws.append(["ERP 연도분포", str(review.get("year_dist"))])
    if review.get("year_mismatch"):
        ws.append(["⚠ 연도 불일치", review["year_mismatch"]])
    ws.append([])
    ws.append([f"■ 미매핑 처지사 (종합표 집계 누락 {review.get('unmatched_dept_amt', 0):,}천원 — 처지사_별칭.json 보강 필요)"])
    for d in review.get("unmatched_dept", []):
        ws.append([d])
    ws.append([])
    ws.append([f"■ 미분류(결측) 예산과목 — 사용자 예산과목 구성에 없음 ({review.get('unclassified_item_amt', 0):,}천원, 종합표 미반영)"])
    for it in review.get("unclassified_item", []):
        ws.append([it])
    ws.append([])
    ws.append(["■ 요약"])
    for k, v in review.get("summary", {}).items():
        ws.append([k, v])


def write_actual_workbook(out_path, plan_rows, new_rows, dept_columns, item_rows,
                          budget, year, review):
    item_flags = {it["과목"]: it["심의대상"] for cat in item_rows for it in cat["과목들"]}
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "양식1(월별)"
    write_actual_datasheet(ds, plan_rows, new_rows, budget, year, item_flags)
    sm = wb.create_sheet("종합표")
    layout = epw.build_summary_layout(sm, dept_columns, item_rows, budget, year, kind="실적")
    epw.write_summary_formulas(sm, layout, value_col=VALUE_COL_LETTER)
    write_review_sheet(wb, review)
    wb.save(out_path)
    wb.close()


def write_zrfm2_v1(src_path, out_path, erp_annotated):
    """원본 zrfm2 + R열(18)에 매칭 사업명. _erp_row 기준으로 매핑."""
    label_by_row = {int(r): lab for r, lab in
                    zip(erp_annotated["_erp_row"], erp_annotated["매칭사업명"])}
    wb = openpyxl.load_workbook(src_path)
    ws = wb[wb.sheetnames[0]]
    ws.cell(1, 18, "사업명(매칭)")
    for r in range(2, ws.max_row + 1):
        lab = label_by_row.get(r)
        if lab:
            ws.cell(r, 18, lab)
    wb.save(out_path)
    wb.close()
