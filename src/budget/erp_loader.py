# -*- coding: utf-8 -*-
"""ERP(SAP) zrfm2 실적 데이터 로더.

zrfm2 컬럼(위치 기준):
  B(2) 약정항목=계정코드, C(3) 약정항목텍스트=예산과목 원문, D(4) 기간/연도,
  F(6) 참조 전표 번호, G(7) 지급예산금액(원), H(8) 텍스트=사업명, I(9) 이름=지사.
금액 단위는 원 → 천원으로 환산(/1000). 부호(역인식 음수)는 그대로 순액 집계.
"""
import pandas as pd
import openpyxl

ERP_COLUMNS = [
    "계정코드", "예산과목원문", "연도", "전표번호",
    "금액원", "금액천원", "사업명", "지사원문", "_erp_row",
]

# zrfm2 시트 내 위치(1-index)
_COL = {
    "계정코드": 2, "예산과목원문": 3, "기간연도": 4, "전표번호": 6,
    "금액": 7, "사업명": 8, "지사": 9,
}


def _parse_year(v):
    """'2025.001' / 2025 / '2025' → '2025'."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return s.split(".")[0].strip()


def load_erp(path, sheet=None):
    """zrfm2 워크북을 DataFrame(ERP_COLUMNS)으로 로드."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    recs = []
    r = 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        r += 1
        # 완전 빈 행 skip
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        amt = row[_COL["금액"] - 1]
        amt = float(amt) if isinstance(amt, (int, float)) else None
        recs.append({
            "계정코드": (str(row[_COL["계정코드"] - 1]).strip()
                      if row[_COL["계정코드"] - 1] is not None else None),
            "예산과목원문": (str(row[_COL["예산과목원문"] - 1]).strip()
                       if row[_COL["예산과목원문"] - 1] is not None else None),
            "연도": _parse_year(row[_COL["기간연도"] - 1]),
            "전표번호": (str(row[_COL["전표번호"] - 1]).strip()
                     if row[_COL["전표번호"] - 1] is not None else None),
            "금액원": amt,
            "금액천원": (amt / 1000.0 if amt is not None else None),
            "사업명": (str(row[_COL["사업명"] - 1]).strip()
                    if row[_COL["사업명"] - 1] is not None else None),
            "지사원문": (str(row[_COL["지사"] - 1]).strip()
                     if row[_COL["지사"] - 1] is not None else None),
            "_erp_row": r,
        })
    wb.close()
    return pd.DataFrame(recs, columns=ERP_COLUMNS)


def detect_years(df):
    """데이터에 등장한 연도 분포 {연도: 건수}."""
    out = {}
    for y in df["연도"]:
        if y:
            out[y] = out.get(y, 0) + 1
    return out
