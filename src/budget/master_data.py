# -*- coding: utf-8 -*-
"""마스터 데이터 로드 (손익예산26_최종본의 예산코드/부서코드 시트).

- 예산코드 시트: 약정항목(계정코드) -> 표준 예산과목명 + 속성
- 부서코드 시트: 부서코드 -> 처지사
ERP(zrfm2)의 코드값을 계획본 표준 표기로 환원하는 데 사용한다.
"""
import openpyxl

CODE_SHEET = "예산코드"
DEPT_SHEET = "부서코드"


def load_master(master_path):
    """마스터 워크북에서 코드맵을 읽어 dict로 반환.

    반환: {
      "code_to_item": {약정항목코드(str): 예산과목명(str)},
      "code_to_attr": {약정항목코드(str): 속성(str)},
      "deptcode_to_branch": {부서코드(str): 처지사(str)},
    }
    """
    wb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    code_to_item = {}
    code_to_attr = {}
    if CODE_SHEET in wb.sheetnames:
        ws = wb[CODE_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = row[0]
            if code is None:
                continue
            key = str(code).strip()
            code_to_item[key] = (str(row[1]).strip() if row[1] is not None else None)
            code_to_attr[key] = (str(row[4]).strip() if len(row) > 4 and row[4] is not None else None)

    deptcode_to_branch = {}
    if DEPT_SHEET in wb.sheetnames:
        ws = wb[DEPT_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = row[0]
            if code is None:
                continue
            branch = row[2] if len(row) > 2 else None
            deptcode_to_branch[str(code).strip()] = (str(branch).strip() if branch is not None else None)
    wb.close()
    return {
        "code_to_item": code_to_item,
        "code_to_attr": code_to_attr,
        "deptcode_to_branch": deptcode_to_branch,
    }
