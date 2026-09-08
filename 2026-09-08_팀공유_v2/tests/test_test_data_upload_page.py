"""Test 모드 '계획 및 실적 업로드' 화면(views/test_data_upload_page.py)을 AppTest로 검증한다."""
import io
import pandas as pd
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
import longterm_forecast as lf

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    yield tmp_path


def _run_render():
    import views.test_data_upload_page as page
    page.render()


def _budget_workbook_bytes(dept_code: str, account_name: str, amount_thousand: float, project_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    ws.append(["[단위 : 천원, 부가세 별도]"])
    ws.append([])
    ws.append(["사업명", "예산귀속 \n부서코드", "예산코드", "예산과목", "연예산 합계"])
    ws.append([project_name, dept_code, "9999", account_name, amount_thousand])

    ws_account = wb.create_sheet("예산코드")
    ws_account.append(["예산코드", "예산과목명"])
    ws_account.append(["9999", account_name])

    ws_dept = wb.create_sheet("부서코드")
    ws_dept.append(["부서코드", "부서명"])
    ws_dept.append([dept_code, "화성지사"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_uploading_budget_saves_to_shared_test_budget_store(isolated):
    """예산계획을 올리고 저장하면 lf.load_test_budget_plan()에 반영되고, 중장기예산 Test 모드와
    같은 저장소(TEST_BUDGET_PLAN_PATH)를 쓴다."""
    budget_bytes = _budget_workbook_bytes("700001", "수선유지비-열원경상정비", 100_000.0, "테스트사업")

    at = AppTest.from_function(_run_render, default_timeout=30)
    at.run()
    assert not at.exception

    uploader = next(u for u in at.file_uploader if u.key == "tdata_budget_upload")
    uploader.set_value([("예산계획.xlsx", budget_bytes, _XLSX_MIME)])
    at.run()
    assert not at.exception

    save_button = next(b for b in at.button if b.key == "tdata_budget_apply")
    save_button.click().run()
    assert not at.exception

    saved = lf.load_test_budget_plan()
    assert len(saved) == 1
    assert saved.iloc[0]["사업명"] == "테스트사업"


def test_without_shared_actuals_shows_warning(isolated):
    """'투자유형 예측 테스트'에서 아직 실적을 확정하지 않았으면 경고를 보여주고 예외 없이 렌더링된다."""
    at = AppTest.from_function(_run_render, default_timeout=30)
    at.run()
    assert not at.exception
    assert any("아직 실적이 없습니다" in w.value for w in at.warning)


def test_ready_message_shown_when_budget_and_actuals_both_present(isolated):
    """예산계획과 (세션에 미리 채워진) 실적이 모두 있으면 '배정 대비 실적'을 볼 수 있다는 안내가 뜬다."""
    budget_df = pd.DataFrame([
        {"예산귀속 \n부서코드": "700001", "예산과목": "수선유지비-열원경상정비", "연예산 합계": 100_000_000.0},
    ])
    lf.save_test_budget_plan(budget_df)

    shared_df = pd.DataFrame([
        {"사업명": "테스트사업", "예산과목": "수선유지비-열원경상정비", "사업장": "700001",
         "금액": 90_000_000.0, "투자유형_확정": "기타기계장치", "투자유형세부_확정": ""},
    ])

    at = AppTest.from_function(_run_render, default_timeout=30)
    at.session_state["ptest_shared_df"] = shared_df
    at.run()
    assert not at.exception
    assert any("모두 준비되었습니다" in s.value for s in at.success)
