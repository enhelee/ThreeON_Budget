"""예산 표준화 결과를 재무팀 참고 양식(builtin_template_표준화.xlsx, 25개 시트)에 채워 넣는다.

셀 좌표를 하드코딩하지 않고, 시트 안의 라벨(예산과목명, '손익소계', 'Standization' 등)을
스캔해서 위치를 찾는다 - template_fill.py와 같은 방식이다. '합계'/'중대형chp'/'DH'/
'소형CHP(여기부터)' 시트는 각 지사 시트 값을 더하는 수식이 이미 들어있어(=동탄!B4+화성!B4+...)
지사 시트만 채우면 엑셀에서 열 때 자동으로 집계된다 - 따로 채우지 않는다.
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from builtin_categories import normalize_account_name
from standardization import ALL_ACCOUNTS, METHOD_OPTIONS, _site_sheet_name_map

BUILTIN_TEMPLATE_PATH = "builtin_template_표준화.xlsx"


def _set_cell(ws, row: int, col: int, value):
    """ws.cell(...).value = value의 안전한 버전. 일부 지사 시트에는 빈 병합 셀이 그대로 남아있어
    (예: 삼송 D34:D39) 그 위치에 쓰려고 하면 MergedCell이 read-only라 오류가 난다.
    병합된 자리라면 병합을 풀고 쓴다(내용이 원래 비어있던 장식용 병합이라 안전하다)."""
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if (merged_range.min_row <= row <= merged_range.max_row
                    and merged_range.min_col <= col <= merged_range.max_col):
                ws.unmerge_cells(str(merged_range))
                break
        cell = ws.cell(row=row, column=col)
    cell.value = value


def _find_year_row(ws) -> dict:
    """연도 헤더 행을 찾아 {열번호: 연도} 반환. 없으면 빈 dict."""
    for r in range(1, min(ws.max_row, 10) + 1):
        year_by_col = {
            c: int(ws.cell(row=r, column=c).value)
            for c in range(2, ws.max_column + 1)
            if isinstance(ws.cell(row=r, column=c).value, (int, float)) and 2000 < ws.cell(row=r, column=c).value < 2100
        }
        if year_by_col:
            return year_by_col
    return {}


def _find_row_by_exact_label(ws, label: str, max_row: int = None):
    max_row = max_row or ws.max_row
    for r in range(1, max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == label:
            return r
    return None


def _match_history_account(label: str) -> str | None:
    """실적 히스토리 섹션(자본/손익 라인 아이템) 라벨 -> 우리 시스템의 정식 계정과목명."""
    normalized = normalize_account_name(label)
    return normalized if normalized in ALL_ACCOUNTS else None


def _match_standization_account(label: str) -> str | None:
    """'Standization' 표 안의 라벨 -> 정식 계정과목명. 그룹 헤더/장식성 문구는 None."""
    label = (label or "").strip()
    if not label or label == "기계장치":
        return None  # 자본예산 Standization 블록의 그룹 헤더일 뿐 - 실제 값은 '기타기계장치' 행에 들어간다
    if "기타기계장치" in label:
        return "기계장치"
    if "외주비-열원정기점검" in label:
        return "외주비-열원정기점검"
    if label == "공구와기구" or label.startswith("공구와기구"):
        return "공구와기구-열원시설공기구"
    normalized = normalize_account_name(label)
    if normalized in ALL_ACCOUNTS:
        return normalized
    for acct in ALL_ACCOUNTS:
        if label.startswith(acct):
            return acct
    return None


def _fill_history_section(ws, year_by_col: dict, yearly_by_account: dict):
    """자본/손익 라인 아이템 행(연도별 실적)을 채운다. yearly_by_account: {계정과목: {연도: 금액}}"""
    subtotal_row = _find_row_by_exact_label(ws, "손익소계")
    last_row = subtotal_row if subtotal_row else ws.max_row
    for r in range(1, last_row + 1):
        label = ws.cell(row=r, column=1).value
        account = _match_history_account(label)
        if account is None or account not in yearly_by_account:
            continue
        for col, year in year_by_col.items():
            amount = yearly_by_account[account].get(year)
            if amount is not None:
                _set_cell(ws, r, col, round(amount / 1000, 2))  # 양식 단위: 천원


def _fill_grade_history_rows(ws, grade_by_year: dict):
    """정비등급 이력 행(손익소계 다음, 보조소계 건너뛴 뒤)에 등급을 채운다."""
    subtotal_row = _find_row_by_exact_label(ws, "손익소계")
    if subtotal_row is None or not grade_by_year:
        return
    year_by_col = _find_year_row(ws)
    r = subtotal_row + 2
    while r <= ws.max_row:
        label = ws.cell(row=r, column=1).value
        if label is None or str(label).strip() == "":
            break
        for col, year in year_by_col.items():
            grade = grade_by_year.get(year)
            if grade is not None:
                _set_cell(ws, r, col, grade)
        r += 1  # 장비 행이 여러 개(#1 GT/#2 GT/ST)라도 동일한 이력을 반복 기입


def _fill_standization_block(ws, block_label: str, grades: list, standard_by_account_grade: dict):
    """'손익예산 Standization' / '자본예산 Standization' 표를 채운다.
    grades: 그 지사에 실제 존재하는 등급 목록(정렬됨). 헤더가 비어있는 빈 양식이므로 직접 써 넣는다."""
    header_row = _find_row_by_exact_label(ws, block_label)
    if header_row is None:
        return

    grade_start_col = 4  # D열부터 (A=라벨, B/C는 빈 칸 - 참고 양식 레이아웃과 동일)
    for i, grade in enumerate(grades):
        _set_cell(ws, header_row, grade_start_col + i, grade)
    note_col = grade_start_col + len(grades)
    _set_cell(ws, header_row, note_col, "비고")

    r = header_row + 1
    while r <= ws.max_row:
        label = ws.cell(row=r, column=1).value
        if label is None or str(label).strip() == "":
            break
        account = _match_standization_account(label)
        if account is not None:
            notable = []
            for i, grade in enumerate(grades):
                entry = standard_by_account_grade.get((account, grade))
                if entry is not None:
                    amount, note = entry
                    _set_cell(ws, r, grade_start_col + i, round(amount / 1000, 2))
                    if note and note not in METHOD_OPTIONS:
                        notable.append(f"{grade}:{note}")
            if notable:
                _set_cell(ws, r, note_col, "; ".join(dict.fromkeys(notable)))
        r += 1


def fill_standardization_workbook(actuals_df: pd.DataFrame, grade_history_df: pd.DataFrame,
                                   standard_df: pd.DataFrame, output_path: str,
                                   template_path: str = BUILTIN_TEMPLATE_PATH) -> str:
    """
    actuals_df: 사업장(표시명), 연도, 예산과목, 금액
    grade_history_df: 사업장(표시명), 연도, 등급
    standard_df: standardization.compute_standard_amounts()의 결과 (사업장, 구분, 예산과목, 등급, 표준금액, 산출방식, 비고)
    """
    wb = load_workbook(template_path)
    site_to_sheet = {v: k for k, v in _site_sheet_name_map().items()}

    actuals_df = actuals_df.copy()
    actuals_df["예산과목"] = actuals_df["예산과목"].apply(normalize_account_name)

    for site, sheet_name in site_to_sheet.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]

        site_actuals = actuals_df[actuals_df["사업장"] == site]
        yearly_by_account = {
            account: g.groupby("연도")["금액"].sum().to_dict()
            for account, g in site_actuals.groupby("예산과목")
        }
        year_by_col = _find_year_row(ws)
        _fill_history_section(ws, year_by_col, yearly_by_account)

        site_grades_hist = grade_history_df[grade_history_df["사업장"] == site]
        grade_by_year = dict(zip(site_grades_hist["연도"], site_grades_hist["등급"]))
        _fill_grade_history_rows(ws, grade_by_year)

        site_standard = standard_df[standard_df["사업장"] == site]
        if site_standard.empty:
            continue
        grades = sorted(site_standard["등급"].unique())

        for block_label, category in (("손익예산 Standization", "손익"), ("자본예산 Standization", "자본")):
            cat_standard = site_standard[site_standard["구분"] == category]
            lookup = {
                (row["예산과목"], row["등급"]): (row["표준금액"], row["비고"])
                for _, row in cat_standard.iterrows()
            }
            _fill_standization_block(ws, block_label, grades, lookup)

    wb.save(output_path)
    return output_path
