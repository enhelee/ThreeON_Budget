# -*- coding: utf-8 -*-
"""전망 엑셀 파서 4종 — v2 `standardization.py`·`longterm_forecast.py` 의 import 함수 승계.

Phase 6-6. 정비등급 이력(표준화 참고 엑셀)·정기점검보수공사 일정(미래 등급)·
고온부품 계획·본사 원가분배 마스터를 읽는다. **파싱만 한다** — 저장은 화면이 미리보기를
확인한 뒤 `/api/forecast/state/*` 로 한다(v2 도 미리보기 → 적용 두 단계였다).

v2 와 다른 것 둘.
  ① 시트명 → 지사명 매핑. v2 는 site_type_map.csv 의 표시명에서 «지사»·«사업소»를 뗀
     것을 썼다. 여기서는 호출측이 dept_config 의 지사 이름을 넘기고 같은 규칙으로 만든다.
  ② 단위 천원. v2 고온부품 파서는 양식(천원)을 원으로 환산했다. 이 앱은 천원이 표준이다.
"""
import io
import re

import pandas as pd
from openpyxl import load_workbook

from .forecast_calc import VALID_GRADE_TOKENS, _classify_grade_row

GRADE_COLS = ["사업장", "연도", "등급"]
HOT_PARTS_COLS = ["사업장", "연도", "항목", "금액"]
HQ_RATIO_COLS = ["사업장", "계약체결금액"]


def site_sheet_map(site_names) -> dict:
    """{'화성': '화성지사', '수원': '수원사업소', …} — 시트·라벨의 짧은 이름을 지사명으로."""
    mapping = {}
    for name in site_names:
        short = str(name).replace("지사", "").replace("사업소", "").strip()
        if short:
            mapping[short] = str(name)
    return mapping


def _wb(data: bytes):
    return load_workbook(io.BytesIO(data), data_only=True)


# ---------------------------------------------------------------- 정비등급 이력 (표준화 참고 엑셀)
def grades_from_workbook(data: bytes, site_names) -> pd.DataFrame:
    """지사 시트마다 '손익소계' 다음 장비 행에서 GT 계열 행(없으면 첫 행)의 연도별 등급."""
    wb = _wb(data)
    sheet_to_site = site_sheet_map(site_names)
    rows = []
    for sheet_name in wb.sheetnames:
        site = sheet_to_site.get(sheet_name.strip())
        if site is None:
            continue
        ws = wb[sheet_name]
        year_row = None
        for r in range(1, min(ws.max_row, 10) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(2, ws.max_column + 1)]
            if any(isinstance(v, (int, float)) and 2000 < v < 2100 for v in vals):
                year_row = r
                break
        if year_row is None:
            continue
        year_by_col = {c: int(ws.cell(row=year_row, column=c).value)
                       for c in range(2, ws.max_column + 1)
                       if isinstance(ws.cell(row=year_row, column=c).value, (int, float))}
        subtotal_row = None
        for r in range(year_row + 1, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value or "").strip() == "손익소계":
                subtotal_row = r
                break
        if subtotal_row is None:
            continue
        equipment_rows = []
        r = subtotal_row + 2                     # 손익소계 다음(보조소계)은 건너뛴다
        while r <= ws.max_row:
            label = ws.cell(row=r, column=1).value
            if label is None or str(label).strip() == "":
                break
            equipment_rows.append((r, str(label).strip()))
            r += 1
        if not equipment_rows:
            continue
        chosen_row = next((rl for rl in equipment_rows if "GT" in rl[1]), equipment_rows[0])[0]
        for col, year in year_by_col.items():
            grade = ws.cell(row=chosen_row, column=col).value
            if grade is None:
                continue
            grade = str(grade).strip()
            if not grade or grade == "-":
                continue
            rows.append({"사업장": site, "연도": year, "등급": grade})
    return pd.DataFrame(rows, columns=GRADE_COLS)


# ---------------------------------------------------------------- 정기점검보수공사 일정 (미래 등급)
def schedule_from_workbook(data: bytes, site_names) -> pd.DataFrame:
    """'정기점검보수공사' 시트에서 지사별 미래 정비등급. GT 계열 행 우선, 없으면 ST 계열 행."""
    wb = _wb(data)
    sheet_name = next((n for n in wb.sheetnames if "정기점검보수공사" in n), None)
    if sheet_name is None:
        return pd.DataFrame(columns=GRADE_COLS)
    ws = wb[sheet_name]

    year_row = None
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and 2020 < v < 2050:
                year_row = r
                break
        if year_row:
            break
    if year_row is None:
        return pd.DataFrame(columns=GRADE_COLS)

    year_by_gt_col = {c: int(ws.cell(row=year_row, column=c).value)
                      for c in range(1, ws.max_column + 1)
                      if isinstance(ws.cell(row=year_row, column=c).value, (int, float))
                      and 2020 < ws.cell(row=year_row, column=c).value < 2050}
    data_start_row = year_row + 3           # 연도행, '등급'행, GT/ST/기간행 다음부터

    short_to_full = site_sheet_map(site_names)
    short_names_sorted = sorted(short_to_full.keys(), key=len, reverse=True)

    site_rows, current_site, blank_run = {}, None, 0
    r = data_start_row
    while r <= ws.max_row:
        row_values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v in (None, "") for v in row_values):
            blank_run += 1
            if blank_run >= 3:
                break
            r += 1
            continue
        blank_run = 0
        label = ws.cell(row=r, column=1).value
        if label is not None and str(label).strip():
            raw = str(label).replace(" ", "").strip()
            match = next((s for s in short_names_sorted if raw.startswith(s)), None)
            current_site = short_to_full.get(match) if match else None
        if current_site:
            site_rows.setdefault(current_site, []).append(r)
        r += 1

    rows = []
    for site, row_list in site_rows.items():
        candidates = []
        for row_idx in row_list:
            gt_values = {y: ws.cell(row=row_idx, column=c).value for c, y in year_by_gt_col.items()}
            st_values = {y: ws.cell(row=row_idx, column=c + 1).value for c, y in year_by_gt_col.items()}
            if _classify_grade_row(gt_values) == "GT":
                candidates.append(("GT", gt_values))
                continue
            if _classify_grade_row(st_values) == "ST":
                candidates.append(("ST", st_values))
        chosen = next((c for c in candidates if c[0] == "GT"), None) or (candidates[0] if candidates else None)
        if chosen is None:
            continue
        for year, grade in chosen[1].items():
            if grade is None:
                continue
            grade = str(grade).strip()
            if grade not in VALID_GRADE_TOKENS:
                continue
            rows.append({"사업장": site, "연도": year, "등급": grade})
    return pd.DataFrame(rows, columns=GRADE_COLS)


# ---------------------------------------------------------------- 고온부품
def hot_parts_from_workbook(data: bytes, site_names) -> pd.DataFrame:
    """'고온부품(NN)' 시트에서 지사×연도×항목(고온부품재생/신품구매) 금액. **천원 그대로.**"""
    wb = _wb(data)
    sheet_name = next((n for n in wb.sheetnames if n.strip().startswith("고온부품")), None)
    if sheet_name is None:
        return pd.DataFrame(columns=HOT_PARTS_COLS)
    ws = wb[sheet_name]

    header_row, year_by_col = None, {}
    for r in range(1, 8):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            m = re.match(r"^(\d{2})년", str(v).strip())
            if m:
                found[c] = 2000 + int(m.group(1))
        if found:
            header_row, year_by_col = r, found
            break
    if header_row is None:
        return pd.DataFrame(columns=HOT_PARTS_COLS)

    short_to_full = site_sheet_map(site_names)
    rows, current_site, blank_run = [], None, 0
    r = header_row + 1
    while r <= ws.max_row:
        site_label = ws.cell(row=r, column=1).value
        item = ws.cell(row=r, column=5).value
        if site_label:
            current_site = short_to_full.get(str(site_label).strip())
        if item is None or current_site is None:
            blank_run += 1
            if blank_run >= 3:
                break
            r += 1
            continue
        blank_run = 0
        item = str(item).strip()
        for col, year in year_by_col.items():
            amount = ws.cell(row=r, column=col).value
            if not amount:
                continue
            rows.append({"사업장": current_site, "연도": year, "항목": item, "금액": float(amount)})
        r += 1
    return pd.DataFrame(rows, columns=HOT_PARTS_COLS)


# ---------------------------------------------------------------- 본사 원가분배
def hq_master_from_workbook(data: bytes, site_names, base_year: int = 2026) -> tuple:
    """'본사 원가분배' 시트 → (마스터 넓은 표, 배분비율 표). 라벨을 스캔해 위치를 찾는다.

    헤더 «NN년 예산» 을 찾아 그 열을 기준연도 예산 열(«NN년예산»)로, 오른쪽 지사 열들을
    지사명으로 읽는다. 세부내역 '합계'(중분류 소계)는 빼고, 대분류 '합계' 에서 멈춘다.
    이어지는 '경상정비 원가배부기준' 블록의 '합계' 행이 계약체결금액(배분비율)이다.
    """
    budget_label_re = re.compile(r"^(\d{2})년\s*예산$")
    empty_master = pd.DataFrame(columns=["대분류", "중분류", "세부내역", f"{int(base_year) % 100:02d}년예산"])
    empty_ratio = pd.DataFrame(columns=HQ_RATIO_COLS)
    wb = _wb(data)
    sheet_name = next((n for n in wb.sheetnames if "본사 원가분배" in n), None)
    if sheet_name is None:
        return empty_master, empty_ratio
    ws = wb[sheet_name]

    header_row, budget_col, budget_name = None, None, None
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            m = budget_label_re.match(str(ws.cell(row=r, column=c).value or "").strip())
            if m:
                header_row, budget_col, budget_name = r, c, f"{m.group(1)}년예산"
                break
        if header_row:
            break
    if header_row is None:
        return empty_master, empty_ratio

    short_to_full = site_sheet_map(site_names)
    site_col = {}
    for c in range(budget_col + 1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        name = short_to_full.get(str(v).strip())
        if name:
            site_col[c] = name
    site_order = [site_col[c] for c in sorted(site_col)]

    rows, major, minor = [], None, None
    r = header_row + 1
    while r <= ws.max_row:
        c1, c2, c3 = (ws.cell(row=r, column=k).value for k in (1, 2, 3))
        budget = ws.cell(row=r, column=budget_col).value
        label1 = str(c1).strip() if c1 else ""
        if label1 in ("합계", "경상정비 원가배부기준"):
            break
        if c1:
            major = label1
        if c2:
            minor = str(c2).strip()
        detail = str(c3).strip() if c3 else ""
        if detail == "합계":                    # 중분류 소계 — 개별 라인과 중복
            r += 1
            continue
        if budget not in (None, ""):
            row = {"대분류": major or "", "중분류": minor or "", "세부내역": detail, budget_name: budget}
            for col, site in site_col.items():
                row[site] = ws.cell(row=r, column=col).value or 0
            rows.append(row)
        r += 1
    master_cols = ["대분류", "중분류", "세부내역", budget_name] + site_order
    master_df = pd.DataFrame(rows, columns=master_cols) if rows else empty_master

    ratio_rows = []
    while r <= ws.max_row:
        c3 = ws.cell(row=r, column=3).value
        if c3 and str(c3).strip() == "합계":
            for col, site in site_col.items():
                val = ws.cell(row=r, column=col).value
                if val not in (None, ""):
                    ratio_rows.append({"사업장": site, "계약체결금액": float(val)})
            break
        r += 1
    ratio_df = pd.DataFrame(ratio_rows, columns=HQ_RATIO_COLS) if ratio_rows else empty_ratio
    return master_df, ratio_df
