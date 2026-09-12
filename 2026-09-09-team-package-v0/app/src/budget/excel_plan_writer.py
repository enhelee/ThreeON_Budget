# -*- coding: utf-8 -*-
"""동적 종합표(SUMIFS)와 양식1(월별) 데이터시트를 작성 (원본 워크북 없이 신규 생성)."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import loaders
from . import missing as missing_mod
from .config_store import SIMUI_THRESHOLD_THOUSAND

DATA_HEADER = [
    "주관부서명", "예산귀속 부서코드", "예산귀속 부서명(처.지사)", "예산귀속 부서명(부)",
    "속성", "예산코드", "예산과목", "사업명", "산출내역", "연예산 합계",
]
MISSING_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
SUBTOTAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_ROW = 3
NAME_ROW = 4

_MISSING_COL_INDEX = {"처지사": 3, "예산과목": 7, "연예산": 10, "사업명": 8}


def write_datasheet(ws, df, budget, year, item_flags):
    ws.cell(1, 1, f"{year}년 {budget}예산 계획 (단위: 천원)")
    for i, h in enumerate(DATA_HEADER, 1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for i, h in ((11, "검증"), (12, "심의플래그")):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL

    r = 4
    for _, row in df.iterrows():
        for i, key in enumerate(loaders.PLAN_COLUMNS, 1):
            val = row[key]
            ws.cell(r, i, None if pd.isna(val) else val)
        mf = missing_mod.missing_fields(row)
        for f in mf:
            ws.cell(r, _MISSING_COL_INDEX[f]).fill = MISSING_FILL
        ws.cell(r, 11, (", ".join(mf) + " 누락") if mf else "")
        item = row["예산과목"]
        amt = row["연예산"]
        flag = "O" if item_flags.get(item) and not pd.isna(amt) and amt >= SIMUI_THRESHOLD_THOUSAND else ""
        ws.cell(r, 12, flag)
        r += 1
    return {"written_rows": r - 4}


def build_summary_layout(ws, dept_columns, item_rows, budget, year, kind="계획"):
    ws.cell(1, 1, f"{year}년 {budget}예산 {kind} 종합표 (단위: 천원)")
    ws.merge_cells(start_row=HEADER_ROW, start_column=1, end_row=NAME_ROW, end_column=3)
    c = ws.cell(HEADER_ROW, 1, "구 분")
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = HEADER_FILL

    col = 4
    groups = []
    for g in dept_columns:
        start = col
        dept_cols = {}
        for name in g["지사"]:
            dept_cols[name] = col
            ws.cell(NAME_ROW, col, name)
            col += 1
        subtotal_col = col
        ws.cell(NAME_ROW, col, "소계")
        col += 1
        ws.merge_cells(start_row=HEADER_ROW, start_column=start, end_row=HEADER_ROW, end_column=subtotal_col)
        gc = ws.cell(HEADER_ROW, start, g["그룹"])
        gc.font = Font(bold=True, color="FFFFFF")
        gc.fill = HEADER_FILL
        groups.append({"그룹": g["그룹"], "dept_cols": dept_cols, "subtotal_col": subtotal_col})

    total_col = col
    ws.merge_cells(start_row=HEADER_ROW, start_column=total_col, end_row=NAME_ROW, end_column=total_col)
    tc = ws.cell(HEADER_ROW, total_col, "합계")
    tc.font = Font(bold=True, color="FFFFFF")
    tc.fill = HEADER_FILL

    for c_ in range(1, total_col + 1):
        for r_ in (HEADER_ROW, NAME_ROW):
            cell = ws.cell(r_, c_)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = NAME_ROW + 1
    item_row_layout = []
    for cat in item_rows:
        cat_start = row
        rows = []
        for it in cat["과목들"]:
            if it["심의대상"]:
                ws.cell(row, 2, it["과목"])
                ws.cell(row, 3, "5천만원 이상(심의대상)")
                rows.append({"과목": it["과목"], "심의구분": "이상", "row": row})
                row += 1
                ws.cell(row, 2, it["과목"])
                ws.cell(row, 3, "5천만원 미만")
                rows.append({"과목": it["과목"], "심의구분": "미만", "row": row})
                row += 1
            else:
                ws.cell(row, 2, it["과목"])
                rows.append({"과목": it["과목"], "심의구분": None, "row": row})
                row += 1
        cat_end = row - 1
        ws.cell(cat_start, 1, cat["대분류"]).font = Font(bold=True)
        if cat_end > cat_start:
            ws.merge_cells(start_row=cat_start, start_column=1, end_row=cat_end, end_column=1)
        subtotal_row = row
        ws.merge_cells(start_row=subtotal_row, start_column=1, end_row=subtotal_row, end_column=2)
        ws.cell(subtotal_row, 1, f"{cat['대분류']} 소계").font = Font(bold=True)
        ws.cell(subtotal_row, 1).fill = SUBTOTAL_FILL
        row += 1
        item_row_layout.append({
            "대분류": cat["대분류"], "rows": rows, "subtotal_row": subtotal_row,
            "item_start": cat_start, "item_end": cat_end,
        })

    grand_total_row = row
    ws.merge_cells(start_row=grand_total_row, start_column=1, end_row=grand_total_row, end_column=2)
    ws.cell(grand_total_row, 1, f"{budget}예산 합계").font = Font(bold=True)
    ws.cell(grand_total_row, 1).fill = SUBTOTAL_FILL

    return {"groups": groups, "total_col": total_col, "item_rows": item_row_layout, "grand_total_row": grand_total_row}


def write_summary_formulas(ws, layout, value_col="J", data_sheet="양식1(월별)",
                           item_col="G", dept_col="C", flag_col="L"):
    """종합표 각 셀에 SUMIFS 수식 작성.

    value_col: 데이터시트에서 합산할 금액 열(계획=연예산 J, 실적=실적금액 I 등).
    item_col/dept_col/flag_col: 데이터시트의 예산과목/처지사/심의플래그 열.
        계획 데이터시트 기본값은 G/C/L. 실적(양식3 포맷) 데이터시트는 B/E/J.
    """
    V = value_col
    DS = data_sheet
    IC, DC, FC = item_col, dept_col, flag_col
    all_cols = []
    for g in layout["groups"]:
        all_cols.extend(g["dept_cols"].values())
        all_cols.append(g["subtotal_col"])
    all_cols.append(layout["total_col"])

    for cat in layout["item_rows"]:
        for item_row in cat["rows"]:
            r = item_row["row"]
            for g in layout["groups"]:
                for dept, col in g["dept_cols"].items():
                    dept_cell = f"{get_column_letter(col)}{NAME_ROW}"
                    if item_row["심의구분"] == "이상":
                        f = (f"=SUMIFS('{DS}'!${V}:${V},'{DS}'!${IC}:${IC},$B{r},"
                             f"'{DS}'!${DC}:${DC},{dept_cell},'{DS}'!${FC}:${FC},\"O\")")
                    elif item_row["심의구분"] == "미만":
                        f = (f"=SUMIFS('{DS}'!${V}:${V},'{DS}'!${IC}:${IC},$B{r},"
                             f"'{DS}'!${DC}:${DC},{dept_cell},'{DS}'!${FC}:${FC},\"<>O\")")
                    else:
                        f = (f"=SUMIFS('{DS}'!${V}:${V},'{DS}'!${IC}:${IC},$B{r},"
                             f"'{DS}'!${DC}:${DC},{dept_cell})")
                    cell = ws.cell(r, col, f)
                    cell.number_format = "#,##0"
                first = min(g["dept_cols"].values())
                last = max(g["dept_cols"].values())
                sc = g["subtotal_col"]
                cell = ws.cell(r, sc, f"=SUM({get_column_letter(first)}{r}:{get_column_letter(last)}{r})")
                cell.number_format = "#,##0"
            terms = "+".join(f"{get_column_letter(g['subtotal_col'])}{r}" for g in layout["groups"])
            cell = ws.cell(r, layout["total_col"], f"={terms}")
            cell.number_format = "#,##0"

        sr = cat["subtotal_row"]
        for col in all_cols:
            L = get_column_letter(col)
            cell = ws.cell(sr, col, f"=SUM({L}{cat['item_start']}:{L}{cat['item_end']})")
            cell.number_format = "#,##0"

    gr = layout["grand_total_row"]
    subtotal_rows = [cat["subtotal_row"] for cat in layout["item_rows"]]
    for col in all_cols:
        L = get_column_letter(col)
        terms = "+".join(f"{L}{sr}" for sr in subtotal_rows)
        cell = ws.cell(gr, col, f"={terms}")
        cell.number_format = "#,##0"


def write_missing_sheet(wb, missing_rows, unclassified_items):
    ws = wb.create_sheet("결측치검토")
    ws.append(["행번호", "사업명", "누락필드"])
    for m in missing_rows:
        name = m["사업명"]
        if pd.isna(name):
            name = None
        ws.append([m["행번호"], name, ", ".join(m["누락필드"])])
    ws.append([])
    ws.append(["미분류 예산과목"])
    for item in unclassified_items:
        ws.append([item])


def write_plan_workbook(out_path, df_filtered, dept_columns, item_rows, budget, year, missing_rows, unclassified_items):
    import openpyxl as _openpyxl
    item_flags = {it["과목"]: it["심의대상"] for cat in item_rows for it in cat["과목들"]}
    wb = _openpyxl.Workbook()
    ds = wb.active
    ds.title = "양식1(월별)"
    write_datasheet(ds, df_filtered, budget, year, item_flags)
    sm = wb.create_sheet("종합표")
    layout = build_summary_layout(sm, dept_columns, item_rows, budget, year)
    write_summary_formulas(sm, layout)
    if missing_rows or unclassified_items:
        write_missing_sheet(wb, missing_rows, unclassified_items)
    wb.save(out_path)
    wb.close()
