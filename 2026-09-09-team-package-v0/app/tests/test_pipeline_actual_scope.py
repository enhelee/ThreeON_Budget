# -*- coding: utf-8 -*-
"""실적분석 스코프 규칙 회귀 테스트(합성 데이터).

- 손익·자본이 섞인 계획 원본을 올려도 각 산출물엔 해당 예산 행만 남는다.
- zrfm2 마지막 총계행·타연도 전표는 집계에서 빠지고 사유가 보고된다.
- 산출물 zrfm2_V1 파일명에 대상연도가 들어간다.
"""
import os

import openpyxl
import pytest

from budget import pipeline_actual
from budget.excel_actual_writer import ACTUAL_SHEET, DATA_START

PLAN_HEADER = ["주관부서명", "부서코드", "처지사", "부서부", "속성",
               "예산코드", "예산과목", "사업명", "산출내역", "연예산"]

PLAN_ROWS = [
    ["기술부", None, "화성지사", "기술부", "제조", None,
     "수선유지비-열원정기점검", "화성 정기점검 보수공사", None, 10000],
    ["기술부", None, "화성지사", "기술부", "자산", None,
     "기계장치", "화성 기계장치 교체", None, 20000],          # 자본 과목
    ["기술부", None, "청주지사", "기술부", "제조", None,
     "수선유지비-열원정기점검", "청주 정기점검", None, 5000],   # 실적 없음 → 미시행
    ["기술부", None, "없는지사", "기술부", "제조", None,
     "수선유지비-열원정기점검", "미등록 지사 사업", None, 7000],
    ["기술부", None, "화성지사", "기술부", "제조", None,
     "듣도보도못한과목", "정체불명 사업", None, 3000],
]

# zrfm2: B계정코드 C과목원문 D기간/연도 E전기일 F전표 G금액(원) H텍스트 I지사 J부서(부)
ERP_ROWS = [
    [None, "60909002", "수선유지비-열원정기점검", "2023.001", "2023-03-02", "D1",
     8_000_000, "화성 정기점검 보수공사 1회", "화성지사", "(CHP)화성지사 기술부"],
    [None, "20704001", "기계장치", "2023.001", "2023-04-01", "D2",
     20_000_000, "화성 기계장치 교체", "화성지사", "(CHP)화성지사 기술부"],
    [None, "60909002", "수선유지비-열원정기점검", "2022.001", "2022-05-01", "D3",
     99_000_000, "작년 정기점검", "화성지사", "(CHP)화성지사 기술부"],
    ["", "", "", "", None, "", 127_000_000, "", "", ""],        # 총계행
]


@pytest.fixture()
def inputs(tmp_path):
    plan = tmp_path / "계획.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    for i, h in enumerate(PLAN_HEADER, 1):
        ws.cell(3, i, h)
    for r, row in enumerate(PLAN_ROWS, 4):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(plan)
    wb.close()

    erp = tmp_path / "zrfm2.XLSX"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 17):
        ws.cell(1, i, f"h{i}")
    for r, row in enumerate(ERP_ROWS, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(erp)
    wb.close()
    return str(plan), str(erp), str(tmp_path / "config"), str(tmp_path / "out")


def _reasons(rows):
    return {r[0] for r in rows}


def test_other_budget_plan_rows_are_excluded(inputs):
    plan, erp, cfg, out = inputs
    res = pipeline_actual.run_actual(plan, erp, None, cfg, out, "2023", "손익")

    # 계획행 = 손익 과목 & 등록 지사 = 화성/청주 정기점검 2행
    assert res["요약"]["계획행수"] == 2
    assert res["요약"]["계획 연예산(천원)"] == 15000
    reasons = _reasons(res["계획행제외"])
    assert "타예산(자본) 과목" in reasons
    assert "처지사 미등록(지사구성에 없음)" in reasons
    assert "미분류 과목(구성에 없음)" in reasons

    wb = openpyxl.load_workbook(res["output_path"])
    items = {wb[ACTUAL_SHEET].cell(r, 2).value
             for r in range(DATA_START, wb[ACTUAL_SHEET].max_row + 1)}
    wb.close()
    assert "기계장치" not in items          # 자본 과목이 손익 산출물에 섞이지 않음
    assert "듣도보도못한과목" not in items


def test_total_row_and_other_year_are_dropped(inputs):
    plan, erp, cfg, out = inputs
    res = pipeline_actual.run_actual(plan, erp, None, cfg, out, "2023", "손익")

    # 2023년 손익 전표는 D1(8,000천원) 하나뿐. 총계행·2022년 전표는 빠진다.
    assert res["요약"]["총 실적(천원)"] == 8000
    reasons = _reasons(res["제외전표"])
    assert "제외(합계행)" in reasons
    assert "제외(타연도)" in reasons

    v1 = openpyxl.load_workbook(res["zrfm2_v1_path"])
    vs = v1[v1.sheetnames[0]]
    gubun = {vs.cell(r, 20).value for r in range(2, vs.max_row + 1)}
    v1.close()
    assert "제외(합계행)" in gubun
    assert "제외(타연도)" in gubun


def test_output_filenames_carry_year(inputs):
    plan, erp, cfg, out = inputs
    res = pipeline_actual.run_actual(plan, erp, None, cfg, out, "2023", "자본")
    assert os.path.basename(res["zrfm2_v1_path"]) == "zrfm2_2023_V1(자본).xlsx"
    assert os.path.basename(res["output_path"]) == "2023년 자본예산_실적.xlsx"
    assert os.path.basename(res["matched_csv_path"]) == "matched_2023_자본.csv"
    # 자본 실행이면 손익 계획행이 빠지고 기계장치만 남는다
    assert res["요약"]["계획행수"] == 1
    assert res["요약"]["총 실적(천원)"] == 20000
