# -*- coding: utf-8 -*-
"""투자유형 축(«무엇 때문에 쓰는 돈인가») — 분류표·라벨 읽기·분류 우선순위·전망 배분·API."""
import io
import os
import sys

import openpyxl
import pandas as pd
import pytest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from budget import db as dbm, invest_type as it, ml_registry   # noqa: E402

HDR = ["연번", "예산과목", "속성", "주관부서명", "예산귀속부서명(처.지사)", "예산귀속부서명(팀)", "사업명",
       "연예산(A)", "최종 실적금액(B)", "구 분", "투자유형", "", "투자유형세부"]


def _book(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "25년 계획 대비 실적"
    ws.cell(1, 1, "자본예산 실적")
    for i, h in enumerate(HDR, 1):
        ws.cell(2, i, h)
    for r, row in enumerate(rows, 3):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(path)
    return str(path)


def _row(i, item, name, amt, kind, detail=""):
    return [i, item, "자산", "기술부", "화성지사", "기계팀", name, amt, amt, "기계장치", kind, "", detail]


def test_every_raw_label_maps_into_taxonomy():
    for raw, (t, _) in it.RAW_MAP.items():
        assert t in it.TYPES, f"«{raw}» 이 분류표에 없는 유형 «{t}» 로 간다"
    for item, (t, _) in it.ITEM_RULE.items():
        assert t in it.TYPES
    assert set(it.TYPE_GROUP.values()) == {"정기", "비정기"}


def test_labels_from_workbook_reads_year_budget_and_mapping(tmp_path):
    p = _book(tmp_path / "(25년) 자본예산 실적 분석_x.xlsx", [
        _row(1, "기계장치", "GT 제어반 교체", 500000, "제어/노후설비"),
        _row(2, "기계장치", "AI 운전 최적화", 300000, "설비개선", "AX/DX"),
        _row(3, "기계장치", "라벨 없는 사업", 100000, None),
    ])
    L = it.labels_from_workbook(p)
    assert len(L) == 2
    assert set(L["연도"]) == {2025} and set(L["예산"]) == {"자본"}
    first = L.iloc[0]
    assert (first["투자유형"], first["세부"]) == ("노후설비 교체", "제어")
    assert (L.iloc[1]["투자유형"], L.iloc[1]["세부"]) == ("설비개선", "AX/DX")
    assert L.iloc[1]["text"] == "AI 운전 최적화 기계장치"


@pytest.fixture()
def conn(tmp_path):
    c = dbm.connect(str(tmp_path / "t.db"))
    ml_registry.init_tables(c)
    return c


def test_classification_precedence(conn):
    ml_registry.add_examples(conn, it.TYPE_MODEL, [("주제어 DCS 업그레이드 기계장치", "노후설비 교체")],
                             source="test", confirmed=True)
    df = pd.DataFrame([
        {"예산과목": "수선유지비-열원정기점검", "사업명": "주제어 DCS 업그레이드"},   # 과목규칙이 사람보다 먼저
        {"예산과목": "기계장치", "사업명": "주제어 DCS 업그레이드"},               # 사람
        {"예산과목": "기계장치", "사업명": "GT LTSA 기성"},                       # 낱말규칙
        {"예산과목": "기계장치", "사업명": "처음 보는 사업"},                     # 모델 없음 → 미분류
    ])
    out = it.classify_frame(conn, df)
    assert list(out["분류출처"]) == ["과목규칙", "사람", "낱말규칙", "미분류"]
    assert list(out["투자유형"]) == ["정기점검·정비", "노후설비 교체", "LTSA", "미분류"]
    assert list(out["대분류"]) == ["정기", "비정기", "정기", "미분류"]


def test_forecast_split_uses_rules_and_past_shares(conn):
    actual = pd.DataFrame([
        {"예산과목": "기계장치", "투자유형": "노후설비 교체", "금액": 300.0},
        {"예산과목": "기계장치", "투자유형": "설비개선", "금액": 100.0},
    ])
    rows = [{"예산과목": "기계장치_기타기계장치", "2026": 1000.0, "2027": 2000.0},
            {"예산과목": "수선유지비-열원경상정비", "2026": 50.0, "2027": 60.0},
            {"예산과목": "기계장치_고온부품재생", "2026": 7.0, "2027": 8.0}]
    out = {r["투자유형"]: r for r in it.forecast_by_type(conn, "2026", rows, actual)["rows"]}
    assert out["노후설비 교체"]["2026"] == pytest.approx(750.0)
    assert out["설비개선"]["2027"] == pytest.approx(500.0)
    assert out["경상정비"]["2026"] == 50.0
    assert out["고온부품·재생정비"]["2027"] == 8.0


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "")
    monkeypatch.setenv("DATABASE_URL", "")
    monkeypatch.setattr(dbm, "DEFAULT_DB", str(tmp_path / "api.db"))
    import server
    from budget import auth as authm
    monkeypatch.setattr(server, "AUTH", authm.AuthConfig())
    from fastapi.testclient import TestClient
    return TestClient(server.app)


def test_api_import_labels_trains_and_report_shows_model(client, tmp_path):
    kinds = ["제어/노후설비", "설비개선", "정기점검", "전기/노후설비"]
    rows = [_row(i, "기계장치", f"{['제어반', 'AI 최적화', '점검', '변압기'][i % 4]} 사업 {i}", 1000 + i, kinds[i % 4],
                 "AX/DX" if kinds[i % 4] == "설비개선" else "") for i in range(40)]
    buf = io.BytesIO()
    _book(tmp_path / "(24년) 자본예산 실적 분석_a.xlsx", rows)
    data = open(tmp_path / "(24년) 자본예산 실적 분석_a.xlsx", "rb").read()
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = client.post("/api/forecast/invest-type/import-labels",
                    files=[("file", ("(24년) 자본예산 실적 분석_a.xlsx", data, XLSX))])
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["rows"] == 40 and j["added"]["투자유형"]["inserted"] == 40
    assert j["trained"][it.TYPE_MODEL]["version"] == 1
    rep = client.get("/api/forecast/invest-type?base_year=2026").json()
    assert rep["model"]["투자유형"]["n_samples"] == 40
    assert [t["투자유형"] for t in rep["taxonomy"]] == it.TYPES
    assert client.get("/api/forecast/invest-type/export?base_year=2026").status_code == 200


def test_api_import_rejects_unlabeled_workbook(client, tmp_path):
    p = _book(tmp_path / "(24년) 자본예산 실적 분석_b.xlsx", [_row(1, "기계장치", "라벨 없음", 1, None)])
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    r = client.post("/api/forecast/invest-type/import-labels", files=[("file", ("x.xlsx", open(p, "rb").read(), XLSX))])
    assert r.status_code == 400


def test_ui_has_invest_type_tab():
    src = os.path.join(APP_DIR, "web", "src")
    view = open(os.path.join(src, "views", "forecast.js"), encoding="utf-8").read()
    router = open(os.path.join(src, "router.js"), encoding="utf-8").read()
    assert '["type", "투자유형별"]' in view and "typeBody()" in view
    assert '"type"]' in router.split("FC_TABS", 1)[1].split("\n", 1)[0]
    assert "loadForecastType" in router
