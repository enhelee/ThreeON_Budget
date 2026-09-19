# -*- coding: utf-8 -*-
"""전망 엑셀 파서 4종 — v2 `test_standardization.py`·`test_longterm_forecast.py` 의 파서 테스트 승계.

바뀐 것 둘. ① 시트명 → 지사명 매핑을 site_type_map.csv 가 아니라 **호출측이 넘기는
지사 이름 목록**(dept_config)에서 만든다. ② **단위 천원** — v2 는 고온부품 양식(천원)을
원으로 환산했지만 이 앱은 천원이 표준이라 그대로 둔다.
"""
import io

import pytest
from openpyxl import Workbook

from budget import forecast_import as fi

SITES = ["화성지사", "동탄지사", "광주전남지사", "수원사업소"]


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_site_sheet_map_strips_suffix():
    m = fi.site_sheet_map(["화성지사", "수원사업소", "광주전남지사"])
    assert m == {"화성": "화성지사", "수원": "수원사업소", "광주전남": "광주전남지사"}


# ---------------------------------------------------------------- 정비등급 이력 (표준화 참고 엑셀)
def _grade_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("화성")
    ws.cell(row=3, column=2, value=2025)
    ws.cell(row=3, column=3, value=2024)
    ws.cell(row=7, column=1, value="기계장치")
    ws.cell(row=7, column=2, value=3628188.06)
    ws.cell(row=7, column=3, value=3798293)
    ws.cell(row=13, column=1, value="자본소계")
    ws.cell(row=15, column=1, value="수선유지비-열원경상정비")
    ws.cell(row=15, column=2, value=6074554.26)
    ws.cell(row=21, column=1, value="손익소계")
    ws.cell(row=22, column=1, value=None)          # 보조소계 (건너뜀)
    ws.cell(row=23, column=1, value="#1 GT")
    ws.cell(row=23, column=2, value="MI")
    ws.cell(row=23, column=3, value="TI")
    ws.cell(row=24, column=1, value="ST")
    ws.cell(row=24, column=2, value="A")
    ws.cell(row=24, column=3, value="C")
    return _bytes(wb)


def test_grades_from_workbook_prefers_gt_row():
    result = fi.grades_from_workbook(_grade_workbook(), SITES)
    assert set(zip(result["연도"], result["등급"])) == {(2025, "MI"), (2024, "TI")}
    assert (result["사업장"] == "화성지사").all()


def test_grades_from_workbook_ignores_unknown_sheets():
    wb = Workbook()
    wb.active.title = "안내"
    assert fi.grades_from_workbook(_bytes(wb), SITES).empty


# ---------------------------------------------------------------- 정기점검보수공사 일정 (미래 등급)
def _schedule_workbook(site_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "정기점검보수공사 일정(2026~2041)"
    ws.cell(row=1, column=1, value="구 분")
    ws.cell(row=1, column=3, value=2026)
    ws.cell(row=1, column=6, value=2027)
    ws.cell(row=2, column=3, value="등급")
    ws.cell(row=2, column=6, value="등급")
    ws.cell(row=3, column=3, value="GT")
    ws.cell(row=3, column=4, value="ST")
    ws.cell(row=3, column=5, value="기간")
    ws.cell(row=3, column=6, value="GT")
    ws.cell(row=3, column=7, value="ST")
    ws.cell(row=3, column=8, value="기간")
    r = 4
    for site, values in site_rows.items():
        ws.cell(row=r, column=1, value=site)
        (gt26, st26), (gt27, st27) = values
        ws.cell(row=r, column=3, value=gt26)
        ws.cell(row=r, column=4, value=st26)
        ws.cell(row=r, column=6, value=gt27)
        ws.cell(row=r, column=7, value=st27)
        r += 1
    return _bytes(wb)


def test_schedule_parser_prefers_gt_when_present():
    result = fi.schedule_from_workbook(_schedule_workbook({"화 성": [("간이", "C"), ("TI", "B")]}), SITES)
    assert set(result["사업장"]) == {"화성지사"}
    assert dict(zip(result["연도"], result["등급"])) == {2026: "간이", 2027: "TI"}


def test_schedule_parser_falls_back_to_st_when_no_gt_token():
    result = fi.schedule_from_workbook(_schedule_workbook({"광주전남": [("-", "A"), ("-", "C")]}), SITES)
    assert set(result["사업장"]) == {"광주전남지사"}
    assert dict(zip(result["연도"], result["등급"])) == {2026: "A", 2027: "C"}


def test_schedule_parser_skips_placeholder_years():
    result = fi.schedule_from_workbook(_schedule_workbook({"수 원": [("-", "-"), ("간이", "C")]}), SITES)
    assert list(result["연도"]) == [2027]


# ---------------------------------------------------------------- 고온부품
def hot_parts_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "고온부품(26)"
    ws.cell(row=4, column=1, value="구  분")
    ws.cell(row=4, column=6, value="26년")
    ws.cell(row=4, column=7, value="27년")
    ws.cell(row=5, column=1, value="화성")
    ws.cell(row=5, column=4, value="기계장치")
    ws.cell(row=5, column=5, value="고온부품재생")
    ws.cell(row=5, column=6, value=1000)
    ws.cell(row=5, column=7, value=0)
    ws.cell(row=6, column=4, value="건설중인자산")
    ws.cell(row=6, column=5, value="신품구매")
    ws.cell(row=6, column=6, value=0)
    ws.cell(row=6, column=7, value=2000)
    return _bytes(wb)


def test_hot_parts_parser_keeps_thousand_won():
    result = fi.hot_parts_from_workbook(hot_parts_workbook(), SITES)
    assert len(result) == 2
    recycled = result[result["항목"] == "고온부품재생"].iloc[0]
    assert recycled["사업장"] == "화성지사"
    assert recycled["연도"] == 2026
    assert recycled["금액"] == 1000           # 천원 그대로 — v2 는 ×1000 했다
    assert result[result["항목"] == "신품구매"].iloc[0]["연도"] == 2027


# ---------------------------------------------------------------- 본사 원가분배
def hq_master_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "26년 본사 원가분배"
    ws.cell(row=5, column=1, value="총계")
    ws.cell(row=5, column=4, value="26년 예산")
    ws.cell(row=5, column=5, value="화성")
    ws.cell(row=5, column=6, value="동탄")
    ws.cell(row=6, column=1, value="플랜트")
    ws.cell(row=6, column=2, value="경상정비")
    ws.cell(row=6, column=3, value="중대형")
    ws.cell(row=6, column=4, value=1000)
    ws.cell(row=6, column=5, value=600)
    ws.cell(row=6, column=6, value=400)
    ws.cell(row=7, column=3, value="합계")      # 소계 행 - 제외
    ws.cell(row=7, column=4, value=1000)
    ws.cell(row=7, column=5, value=600)
    ws.cell(row=7, column=6, value=400)
    ws.cell(row=8, column=1, value="합계")      # 전체 합계 - 중단
    ws.cell(row=8, column=4, value=1000)
    ws.cell(row=9, column=1, value="경상정비 원가배부기준")
    ws.cell(row=10, column=3, value="합계")
    ws.cell(row=10, column=4, value="계약체결금액")
    ws.cell(row=10, column=5, value=6000)
    ws.cell(row=10, column=6, value=4000)
    return _bytes(wb)


def test_hq_master_parser_excludes_subtotal_rows():
    master, ratio = fi.hq_master_from_workbook(hq_master_workbook(), SITES)
    assert len(master) == 1
    assert list(master.columns) == ["대분류", "중분류", "세부내역", "26년예산", "화성지사", "동탄지사"]
    assert master.iloc[0]["화성지사"] == 600
    assert master.iloc[0]["동탄지사"] == 400


def test_hq_ratio_parser_reads_contract_totals():
    _, ratio = fi.hq_master_from_workbook(hq_master_workbook(), SITES)
    assert dict(zip(ratio["사업장"], ratio["계약체결금액"])) == {"화성지사": 6000, "동탄지사": 4000}


def test_hq_master_parser_without_sheet_returns_empty_pair():
    wb = Workbook()
    wb.active.title = "다른시트"
    master, ratio = fi.hq_master_from_workbook(_bytes(wb), SITES)
    assert master.empty and ratio.empty
    assert list(master.columns) == ["대분류", "중분류", "세부내역", "26년예산"]
