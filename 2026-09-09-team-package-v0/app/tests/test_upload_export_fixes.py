# -*- coding: utf-8 -*-
"""2026-10-13 실데이터(2023·2024) 최종 테스트에서 나온 네 가지를 고정한다.

1. 현업 원본 주관부서예산(12·13열, 손익·자본·건설공사 3파일)을 그대로 올릴 수 있어야 한다.
2. 과목명에 공백이 섞여도(«건 물») 같은 과목으로 본다.
3. 같은 연도·예산의 다운로드를 동시에 눌러도 깨진 파일이 나가지 않는다.
4. 한 번도 내려받지 않은 연도를 마감해도 산출물이 보관된다.
"""
import io
import os
import sys
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from budget import db as dbm, loaders, normalize, pipeline_db   # noqa: E402
from tests.test_db_pipeline import ERP_ROWS, PLAN_HEADER, PLAN_ROWS   # noqa: E402

# 23·24년 원본 헤더 — «기준»·«ESG» 가 예산과목과 사업명 사이에 끼고, 연예산이 12열
HDR_2324 = ["주관부서명", "예산귀속\n부서코드", "예산귀속부서명\n(처.지사)", "예산귀속\n부서명(부)", "속성",
            "예산코드", "예산과목", "기준", "ESG 연계 자원배분 분류", "사업명", "산출내역", "연예산 합계",
            "1월", "2월"]
# 25년 원본 — «본부» 가 3열에 더 낀다
HDR_25 = ["주관부서명", "예산귀속 부서코드", "본부", "예산귀속부서명(처.지사)", "예산귀속 부서명(부)", "속성",
          "예산코드", "예산과목", "기준", "중점전략 연계", "사업명", "산출내역", "연예산 합계"]


def _book(path, header, rows):
    wb = openpyxl.Workbook()
    wb.active.title = "표지"
    ws = wb.create_sheet("양식1(월별)")
    for i, h in enumerate(header, 1):
        ws.cell(3, i, h)
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(path)
    return str(path)


def _row2324(dept, item, name, amt):
    return ["기술부", "1000093", dept, "기술부", "제조", "60909002", item, "물가", "해당없음", name, "-", amt, 0, 0]


# ── 1. 원본 양식을 이름으로 읽는다 ─────────────────────────────────────────

def test_original_2324_layout_is_read_by_header_name(tmp_path):
    p = _book(tmp_path / "손익23.xlsx", HDR_2324, [
        _row2324("화성지사", "수선유지비-열원정기점검", "화성 정기점검 보수공사", 10000),
        ["", "", "", "", "", "", 0, "", "", 0, "", 0, 0, 0],      # 양식의 빈 줄(수식이 0)
        _row2324("대구지사", "지급수수료-열원점검수수료", "대구 점검수수료", 2500),
    ])
    df = loaders.load_business_plan(p)
    assert len(df) == 2, "수식만 남은 빈 줄은 사업이 아니다"
    first = df.iloc[0]
    assert first["처지사"] == "화성지사"
    assert first["예산과목"] == "수선유지비-열원정기점검"
    assert first["사업명"] == "화성 정기점검 보수공사", "«기준»·«ESG» 열이 사업명 자리를 밀어내면 안 된다"
    assert first["연예산"] == 10000.0
    assert list(df["_row"]) == [4, 6]


def test_original_25_layout_with_extra_hq_column(tmp_path):
    p = _book(tmp_path / "자본25.xlsx", HDR_25, [
        ["기술부", "1000093", "열원본부", "광교지사", "기술부", "자산", "11101", "기계장치", "물가", "-",
         "광교 GT 부품 교체", "-", 7000],
    ])
    r = loaders.load_business_plan(p).iloc[0]
    assert (r["처지사"], r["예산과목"], r["사업명"], r["연예산"]) == ("광교지사", "기계장치", "광교 GT 부품 교체", 7000.0)


def test_construction_file_with_shifted_amount_column(tmp_path):
    """23년 건설공사 원본은 12열이 비고 «연예산 합계»가 13열에 있다."""
    hdr = HDR_2324[:11] + [None, "연예산 합계"]
    row = _row2324("건설처", "외주비-열원공사비", "열원 배관 공사", None)[:11] + [None, 3300]
    r = loaders.load_business_plan(_book(tmp_path / "건설23.xlsx", hdr, [row])).iloc[0]
    assert r["연예산"] == 3300.0


def test_app_template_layout_still_reads(tmp_path):
    p = _book(tmp_path / "템플릿.xlsx", PLAN_HEADER, PLAN_ROWS)
    df = loaders.load_business_plan(p)
    assert list(df["사업명"]) == [r[7] for r in PLAN_ROWS]


def test_several_files_merge_with_unique_rows(tmp_path):
    """_row 는 전표를 계획 행에 잇는 열쇠 — 합친 뒤에도 겹치면 안 된다."""
    a = _book(tmp_path / "손익.xlsx", HDR_2324, [_row2324("화성지사", "수선유지비-열원정기점검", "가", 1),
                                                _row2324("화성지사", "수선유지비-열원정기점검", "나", 2)])
    b = _book(tmp_path / "자본.xlsx", HDR_2324, [_row2324("화성지사", "기계장치", "다", 3)])
    df = loaders.load_business_plans([a, b])
    assert list(df["사업명"]) == ["가", "나", "다"]
    assert df["_row"].is_unique


# ── 2. 과목명 공백 ────────────────────────────────────────────────────────

@pytest.mark.parametrize("raw", ["건 물", "건　물", " 건물 "])
def test_item_name_spaces_are_ignored(raw):
    assert normalize.normalize_item(raw, {}) == "건물"


def test_item_alias_matches_regardless_of_spaces():
    alias = {"수선유지비-열원보완개선및기타": "수선유지비-열원보완및개선"}
    assert normalize.normalize_item("수선유지비 - 열원보완개선및기타", alias) == "수선유지비-열원보완및개선"


# ── 3·4. 서버: 여러 파일 업로드 · 동시 다운로드 · 마감 보관 ─────────────────

@pytest.fixture()
def srv(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "")
    monkeypatch.setenv("DATABASE_URL", "")
    monkeypatch.setattr(dbm, "DEFAULT_DB", str(tmp_path / "t.db"))
    import server
    from budget import auth as authm
    monkeypatch.setattr(server, "AUTH", authm.AuthConfig())
    monkeypatch.setattr(server, "OUT_DIR", str(tmp_path / "out"))
    os.makedirs(server.OUT_DIR, exist_ok=True)
    from fastapi.testclient import TestClient
    return server, TestClient(server.app), tmp_path


def _xlsx_bytes(header, rows, sheet="양식1(월별)"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    if sheet == "양식1(월별)":
        for i, h in enumerate(header, 1):
            ws.cell(3, i, h)
        start = 4
    else:
        for i, h in enumerate(header, 1):
            ws.cell(1, i, h)
        start = 2
    for r, row in enumerate(rows, start):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _analyzed(client):
    """원본 양식 손익+자본 두 파일을 한 번에 올리고 zrfm2 를 올려 분석까지."""
    loss = _xlsx_bytes(HDR_2324, [_row2324("화성지사", "수선유지비-열원정기점검", "화성 정기점검 보수공사", 10000),
                                  _row2324("화성지사", "수선유지비-열원정기점검", "화성 옥외배관 도색공사", 5000)])
    cap = _xlsx_bytes(HDR_2324, [_row2324("화성지사", "기계장치", "화성 GT 부품 교체", 7000)])
    r = client.post("/api/upload/plan?year=2023",
                    files=[("file", ("손익.xlsx", loss, XLSX)), ("file", ("자본.xlsx", cap, XLSX))])
    assert r.status_code == 200, r.text
    erp = _xlsx_bytes([f"h{i}" for i in range(1, 17)], ERP_ROWS, sheet="Sheet1")
    assert client.post("/api/upload/erp?year=2023",
                       files={"file": ("zrfm2.xlsx", erp, XLSX)}).status_code == 200
    assert client.post("/api/analyze", json={"year": "2023"}).status_code == 200
    return r.json()


def test_plan_upload_merges_several_files_into_one_dataset(srv):
    _, client, _ = srv
    info = _analyzed(client)
    assert info["rows"] == 3 and info["total"] == 22000
    plans = [d for d in client.get("/api/status?year=2023").json()["datasets"] if d["kind"] == "plan"]
    assert len(plans) == 1 and plans[0]["label"] == "손익.xlsx + 자본.xlsx"


def test_concurrent_downloads_build_files_once_and_stay_whole(srv, monkeypatch):
    server, client, _ = srv
    _analyzed(client)
    calls = []
    real = pipeline_db.run_actual_db
    monkeypatch.setattr(server.pipeline_db, "run_actual_db",
                        lambda *a, **k: (calls.append(1), real(*a, **k))[1])
    kinds = ["actual", "v1", "matched", "actual", "v1", "matched"]
    with ThreadPoolExecutor(len(kinds)) as ex:
        res = list(ex.map(lambda k: server.export("2023", "손익", k), kinds))
    assert len(calls) == 1, "같은 (연도,예산) 산출물은 한 번만 다시 만든다"
    for k, r in zip(kinds, res):
        if k == "matched":
            assert r.body.strip()
        else:
            openpyxl.load_workbook(io.BytesIO(r.body)).close()   # 반쪽 파일이면 여기서 깨진다
    assert res[0].body == res[3].body
    assert "filename*=utf-8''" in res[0].headers["content-disposition"]


def test_lock_archives_outputs_even_without_prior_download(srv):
    server, client, _ = srv
    _analyzed(client)
    r = client.post("/api/lock", json={"year": "2023"})
    assert r.status_code == 200
    snap = os.path.join(server.OUT_DIR, "마감", "2023")
    assert "2023년 손익예산_실적.xlsx" in os.listdir(snap)
    assert r.json()["snapshot"] >= 3
    assert client.get("/api/lock-state?year=2023").json()["locked"] is True


# ── 5. 라이브 속도: PostgreSQL 은 표 확인(init_db)을 프로세스당 한 번만 ─────────

def test_postgres_schema_check_runs_once_per_process(tmp_path, monkeypatch):
    """init_db 는 표가 있어도 문장 44개를 보낸다 — Supabase 에선 그게 요청마다 4~7초였다."""
    import sqlite3
    url = "postgresql://example/budget"
    monkeypatch.setattr(dbm.dbcore, "database_url", lambda: url)
    monkeypatch.setattr(dbm.dbcore, "is_postgres_url", lambda u: u == url)
    monkeypatch.setattr(dbm.dbcore, "connect", lambda *a, **k: sqlite3.connect(str(tmp_path / "pg.db")))
    monkeypatch.setattr(dbm, "_PG_READY", set())
    calls = []
    real = dbm.init_db
    monkeypatch.setattr(dbm, "init_db", lambda c: (calls.append(1), real(c))[1])
    for _ in range(3):
        dbm.connect().close()
    assert len(calls) == 1
    # SQLite 경로는 지금처럼 매번 확인한다(테스트가 파일을 지우고 다시 연다)
    dbm.connect(str(tmp_path / "a.db")).close()
    dbm.connect(str(tmp_path / "a.db")).close()
    assert len(calls) == 3


# ── 6. 마감·분석·업로드는 같은 연도에서 한 번에 하나만 ─────────────────────

def test_lock_is_refused_while_same_year_analysis_runs(srv, monkeypatch):
    """2024년 마감(1분)과 다른 창의 재분석이 겹쳐, 고정된 숫자와 보관 파일이 다른 실행에서
    나올 수 있었다(2026-10-13 라이브). 겹치면 뒤에 온 요청을 409 로 돌려보낸다."""
    import threading
    from fastapi import HTTPException
    server, client, _ = srv
    _analyzed(client)
    started, release = threading.Event(), threading.Event()
    real = pipeline_db.run_actual_db

    def slow(*a, **k):
        started.set()
        release.wait(10)
        return real(*a, **k)
    monkeypatch.setattr(server.pipeline_db, "run_actual_db", slow)
    t = threading.Thread(target=server.analyze, args=(server.AnalyzeReq(year="2023"),))
    t.start()
    assert started.wait(10)
    try:
        with pytest.raises(HTTPException) as e:
            server.lock(server.LockReq(year="2023"))
        assert e.value.status_code == 409 and "분석 실행 중" in e.value.detail
        # 다른 연도는 막지 않는다
        assert client.get("/api/lock-state?year=2024").status_code == 200
    finally:
        release.set()
        t.join(20)
    monkeypatch.setattr(server.pipeline_db, "run_actual_db", real)
    assert client.post("/api/lock", json={"year": "2023"}).status_code == 200


def test_relock_summary_counts_reuploads(srv):
    server, client, _ = srv
    _analyzed(client)
    client.post("/api/lock", json={"year": "2023"})
    client.request("DELETE", "/api/lock/2023", json={"reason": "계획 교체"})
    plan = _xlsx_bytes(HDR_2324, [_row2324("화성지사", "수선유지비-열원정기점검", "화성 정기점검 보수공사", 10000)])
    assert client.post("/api/upload/plan?year=2023",
                       files=[("file", ("새계획.xlsx", plan, XLSX))]).status_code == 200
    r = client.post("/api/lock", json={"year": "2023"})
    # 해제와 같은 초의 행도 센다(changes_since 의 >=) — 테스트는 1초 안에 다 돌아 앞선 업로드까지 잡힌다
    assert r.json()["changes"].get("자료업로드", 0) >= 1
