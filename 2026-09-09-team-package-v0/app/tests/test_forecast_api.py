# -*- coding: utf-8 -*-
"""3·4단계 API — 기준연도 · 상태 9종 라운드트립 · 복사 가드 · 가져오기 미리보기 · 계산 (Phase 6-6)."""
import json
import os
import sys

import pytest

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)

from budget import db as dbm   # noqa: E402
from tests.test_forecast_import import hot_parts_workbook, hq_master_workbook   # noqa: E402


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "")
    monkeypatch.setenv("DATABASE_URL", "")       # 빈 문자열 = SQLite (실 DB 로 붙지 않게)
    monkeypatch.setattr(dbm, "DEFAULT_DB", str(tmp_path / "t.db"))
    import server
    from budget import auth as authm
    monkeypatch.setattr(server, "AUTH", authm.AuthConfig())
    from fastapi.testclient import TestClient
    return TestClient(server.app)


def _conn():
    return dbm.connect()


def _seed_run(conn, year, budget, lines):
    cur = conn.execute("INSERT INTO run(year,budget,created_at,summary_json) VALUES(?,?,?,?)",
                       (year, budget, "2026-01-01 00:00:00", json.dumps({})))
    run_id = cur.lastrowid
    for dept, item, amt, gubun in lines:
        conn.execute("INSERT INTO biz_line(run_id,budget,예산과목,처지사,사업명,연예산,실적금액,구분)"
                     " VALUES(?,?,?,?,?,?,?,?)", (run_id, budget, item, dept, "x", 0.0, amt, gubun))
    conn.commit()


def _seed_plan(conn, year, rows):
    cur = conn.execute("INSERT INTO dataset(kind,year,label,uploaded_at) VALUES('plan',?,?,?)",
                       (year, "plan.xlsx", "2026-01-01 00:00:00"))
    ds = cur.lastrowid
    for i, (dept, item, amt) in enumerate(rows):
        conn.execute("INSERT INTO plan_row(dataset_id,주관부서명,부서코드,처지사,부서부,속성,예산코드,예산과목,"
                     "사업명,산출내역,연예산,src_row) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                     (ds, "부", "C", dept, "", "일반", "K", item, f"사업{i}", "", amt, i + 4))
    conn.commit()


def test_healthz_has_no_forecast_url(client):
    """iframe 이 사라졌으니 그 주소도 사라진다."""
    body = client.get("/healthz").json()
    assert body["ok"] is True and "forecast_url" not in body


def test_base_years_default_is_2026_when_nothing(client):
    assert client.get("/api/forecast/base-years").json() == {
        "base_years": [], "default": "2026", "locked_years": []}


def test_state_put_get_round_trip_for_every_table(client):
    payloads = {
        "grades": [{"사업장": "화성지사", "연도": 2027, "등급": "TI"}],
        "methods": [{"사업장": "화성지사", "예산과목": "기계장치", "방식": "최근실적"}],
        "overrides": [{"사업장": "화성지사", "예산과목": "기계장치", "등급": "표준", "표준금액": 5.0}],
        "factors": [{"팩터명": "물가상승률", "연간비율": 0.02, "활성": True},
                    {"팩터명": "노후화", "연간비율": 0.01, "활성": False}],
        "hot-parts": [{"사업장": "화성지사", "연도": 2027, "항목": "고온부품재생", "금액": 10.0}],
        "hq-ratio": [{"사업장": "화성지사", "계약체결금액": 6000.0}],
        "hq-temp": [{"사업명": "a", "예산과목": "기계장치", "연도": 2027, "금액": 1.0}],
        "surprise": [{"사업장": "화성지사", "연도": 2028, "예산과목": "기계장치", "금액": 2.0, "사유": "x"}],
    }
    for table, rows in payloads.items():
        r = client.put(f"/api/forecast/state/{table}", json={"base_year": "2026", "rows": rows})
        assert r.status_code == 200, (table, r.text)
        assert r.json()["table"] == table and len(r.json()["rows"]) == len(rows)
    r = client.put("/api/forecast/state/hq-master", json={
        "base_year": "2026", "columns": ["대분류", "중분류", "세부내역", "26년예산", "화성지사"],
        "rows": [{"대분류": "플랜트", "중분류": "경상정비", "세부내역": "", "26년예산": 600, "화성지사": 600}],
        "ratio_rows": [{"사업장": "화성지사", "계약체결금액": 1.0}]})
    assert r.status_code == 200, r.text

    st = client.get("/api/forecast/state?base_year=2026").json()
    assert st["grades"] == [{"사업장": "화성지사", "연도": 2027, "등급": "TI"}]
    assert st["factors"][1]["활성"] is False
    assert st["hq_master"]["columns"][-1] == "화성지사" and st["hq_master"]["rows"][0]["화성지사"] == 600.0
    assert st["hq_ratio"] == [{"사업장": "화성지사", "계약체결금액": 1.0}]   # hq-master PUT 이 비율도 갈아 끼웠다
    assert st["surprise"][0]["사유"] == "x"
    assert "화성지사" in st["sites"] and "플랜트기술처" not in st["sites"]
    assert st["site_groups"]["화성지사"] == "중대형CHP"
    assert "기계장치" in st["accounts"]["자본"] and "수선유지비-열원경상정비" in st["accounts"]["손익"]
    assert client.get("/api/forecast/base-years").json()["base_years"] == ["2026"]


def test_put_empty_rows_clears_the_table(client):
    client.put("/api/forecast/state/surprise", json={"base_year": "2026", "rows": [
        {"사업장": "화성지사", "연도": 2028, "예산과목": "기계장치", "금액": 2.0, "사유": "x"}]})
    r = client.put("/api/forecast/state/surprise", json={"base_year": "2026", "rows": []})
    assert r.status_code == 200 and r.json()["rows"] == []


def test_unknown_state_table_is_404(client):
    assert client.put("/api/forecast/state/nope", json={"base_year": "2026", "rows": []}).status_code == 404


def test_copy_base_year_guards_touched_target(client):
    client.put("/api/forecast/state/factors", json={"base_year": "2026", "rows": [
        {"팩터명": "물가상승률", "연간비율": 0.03, "활성": True}]})
    assert client.post("/api/forecast/copy-base-year",
                       json={"from_year": "2026", "to_year": "2026"}).status_code == 400
    r = client.post("/api/forecast/copy-base-year", json={"from_year": "2026", "to_year": "2027"})
    assert r.status_code == 200 and r.json()["base_years"] == ["2026", "2027"]
    assert client.get("/api/forecast/state?base_year=2027").json()["factors"][0]["연간비율"] == 0.03
    # 읽기만 한(팩터 시드만 있는) 대상은 손대지 않은 것 — 복사가 막히면 안 된다
    client.get("/api/forecast/state?base_year=2028")
    assert client.post("/api/forecast/copy-base-year",
                       json={"from_year": "2026", "to_year": "2028"}).status_code == 200
    # 사람이 쓴 것이 있으면 409, overwrite 로만 덮는다
    client.put("/api/forecast/state/surprise", json={"base_year": "2027", "rows": [
        {"사업장": "화성지사", "연도": 2028, "예산과목": "기계장치", "금액": 1.0, "사유": ""}]})
    assert client.post("/api/forecast/copy-base-year",
                       json={"from_year": "2026", "to_year": "2027"}).status_code == 409
    assert client.post("/api/forecast/copy-base-year",
                       json={"from_year": "2026", "to_year": "2027", "overwrite": True}).status_code == 200
    assert client.get("/api/forecast/state?base_year=2027").json()["surprise"] == []


def test_import_preview_does_not_save(client):
    data = hot_parts_workbook()
    r = client.post("/api/forecast/import?kind=hot-parts&base_year=2026",
                    files={"file": ("h.xlsx", data, "application/octet-stream")})
    assert r.status_code == 200, r.text
    assert r.json()["rows"][0] == {"사업장": "화성지사", "연도": 2026, "항목": "고온부품재생", "금액": 1000.0}
    assert client.get("/api/forecast/state?base_year=2026").json()["hot_parts"] == []

    r = client.post("/api/forecast/import?kind=hq-master&base_year=2026",
                    files={"file": ("m.xlsx", hq_master_workbook(), "application/octet-stream")})
    body = r.json()
    assert body["columns"] == ["대분류", "중분류", "세부내역", "26년예산", "화성지사", "동탄지사"]
    assert body["ratio_rows"] == [{"사업장": "화성지사", "계약체결금액": 6000.0},
                                  {"사업장": "동탄지사", "계약체결금액": 4000.0}]

    assert client.post("/api/forecast/import?kind=nope&base_year=2026",
                       files={"file": ("h.xlsx", data, "application/octet-stream")}).status_code == 400
    assert client.post("/api/forecast/import?kind=grades&base_year=2026",
                       files={"file": ("h.xlsx", b"not an xlsx", "application/octet-stream")}).status_code == 400


def test_benchmark_and_table_read_locked_runs(client):
    conn = _conn()
    try:
        _seed_run(conn, "2023", "손익", [("화성지사", "수선유지비-열원경상정비", 1000.0, "계획집행")])
        _seed_run(conn, "2025", "손익", [("화성지사", "수선유지비-열원경상정비", 1100.0, "계획집행")])
        dbm.lock_year(conn, "2023")
        dbm.lock_year(conn, "2025")
        _seed_plan(conn, "2026", [("화성지사", "수선유지비-열원경상정비", 1200.0)])
    finally:
        conn.close()
    assert client.get("/api/forecast/base-years").json()["default"] == "2026"

    b = client.get("/api/forecast/benchmark?base_year=2026").json()
    assert b["years_used"] == ["2023", "2025"] and len(b["rows"]) == 1
    assert b["rows"][0]["표준금액"] == 1100.0 and b["population"]["total"] == 2100.0

    t = client.get("/api/forecast/table?base_year=2026").json()
    assert t["years"][0] == 2026 and "화성지사" in t["tables"]
    line = next(r for r in t["tables"]["화성지사"] if r["예산과목"] == "수선유지비-열원경상정비")
    assert line["2026"] == 1200.0                       # 당해년도 = 계획본
    assert line["2027"] == pytest.approx(1100.0 * 1.015)  # 팩터 시드 1.5%
    assert t["notes"]["budget_plan_missing"] is False


# ---------------------------------------------------------------- 6-7 양식 내보내기
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.parametrize("kind", ["standardization", "longterm", "schedule", "hot-parts", "hq-master"])
def test_export_returns_xlsx_for_every_kind(client, kind):
    """다섯 종류 모두 실제 내장 양식을 열어 채운 xlsx 를 돌려준다(빈 상태에서도 깨지지 않아야 한다)."""
    r = client.get(f"/api/forecast/export?kind={kind}&base_year=2026")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX)
    assert "filename*=utf-8''" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"          # zip(xlsx) 시그니처


def test_export_unknown_kind_is_400(client):
    assert client.get("/api/forecast/export?kind=nope&base_year=2026").status_code == 400


def test_export_longterm_contains_forecast_numbers(client):
    """마감 run + 계획본을 심고 내보낸 중장기 양식의 화성 탭에 기준연도 값(백만원)이 들어간다."""
    from openpyxl import load_workbook
    import io
    conn = _conn()
    try:
        _seed_run(conn, "2025", "손익", [("화성지사", "수선유지비-열원경상정비", 1100.0, "계획집행")])
        dbm.lock_year(conn, "2025")
        _seed_plan(conn, "2026", [("화성지사", "수선유지비-열원경상정비", 1200.0)])
    finally:
        conn.close()
    r = client.get("/api/forecast/export?kind=longterm&base_year=2026")
    ws = load_workbook(io.BytesIO(r.content))["화성"]
    row = next(r_ for r_ in range(1, ws.max_row + 1)
               if str(ws.cell(row=r_, column=2).value or "").strip() == "열원경상정비")
    col_2026 = next(c for c in range(1, ws.max_column + 1) if ws.cell(row=4, column=c).value == 26)
    assert ws.cell(row=row, column=col_2026).value == pytest.approx(1200.0 / 1000)   # 천원 → 백만원
