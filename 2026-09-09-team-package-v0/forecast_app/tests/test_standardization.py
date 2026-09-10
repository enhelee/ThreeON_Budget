import pandas as pd
import pytest
from openpyxl import Workbook
import standardization as std


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(std, "GRADE_HISTORY_PATH", str(tmp_path / "site_year_grade.csv"))
    monkeypatch.setattr(std, "METHOD_MAP_PATH", str(tmp_path / "standardization_method_map.csv"))
    monkeypatch.setattr(std, "OVERRIDE_PATH", str(tmp_path / "standardization_overrides.csv"))
    yield tmp_path


def test_grade_history_round_trip(isolated):
    df = pd.DataFrame({"사업장": ["화성지사"], "연도": [2025], "등급": ["MI"]})
    std.save_grade_history(df)
    loaded = std.load_grade_history()
    assert loaded.iloc[0]["등급"] == "MI"


def test_load_grade_history_missing_file_returns_empty(isolated):
    result = std.load_grade_history()
    assert result.empty
    assert list(result.columns) == ["사업장", "연도", "등급"]


def _force_method(site, account, method):
    """recommend_method()의 자동 추천과 무관하게 특정 방식을 강제하는 method_map."""
    return pd.DataFrame({"사업장": [site], "예산과목": [account], "방식": [method]})


def test_compute_grade_conditional_average():
    actuals = pd.DataFrame({
        "사업장": ["화성지사"] * 4,
        "연도": [2022, 2023, 2024, 2025],
        "예산과목": ["수선유지비-열원경상정비"] * 4,
        "금액": [100, 200, 300, 400],
    })
    grades = pd.DataFrame({
        "사업장": ["화성지사"] * 4,
        "연도": [2022, 2023, 2024, 2025],
        "등급": ["MI", "TI", "MI", "TI"],
    })
    method_map = _force_method("화성지사", "수선유지비-열원경상정비", "등급별 평균")
    result = std.compute_standard_amounts(actuals, grades, method_map)
    mi_row = result[result["등급"] == "MI"].iloc[0]
    ti_row = result[result["등급"] == "TI"].iloc[0]
    assert mi_row["표준금액"] == pytest.approx((100 + 300) / 2)
    assert ti_row["표준금액"] == pytest.approx((200 + 400) / 2)
    assert mi_row["방식출처"] == "사용자지정"


def test_compute_falls_back_to_overall_average_when_grade_missing():
    actuals = pd.DataFrame({
        "사업장": ["화성지사"] * 2,
        "연도": [2024, 2025],
        "예산과목": ["기계장치"] * 2,
        "금액": [100, 300],
    })
    grades = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2024, 2025], "등급": ["MI", "MI"]})
    method_map = _force_method("화성지사", "기계장치", "등급별 평균")
    result = std.compute_standard_amounts(actuals, grades, method_map)
    ci_row = result[result["등급"] == "CI"]
    assert ci_row.empty  # 등급 이력에 CI가 아예 없으면 컬럼 자체가 안 생긴다

    # CI가 이력에 존재하지만 그 연도의 실적이 없는 경우
    grades2 = pd.DataFrame({"사업장": ["화성지사"] * 3, "연도": [2023, 2024, 2025], "등급": ["CI", "MI", "MI"]})
    result2 = std.compute_standard_amounts(actuals, grades2, method_map)
    ci_row2 = result2[result2["등급"] == "CI"].iloc[0]
    assert ci_row2["표준금액"] == pytest.approx((100 + 300) / 2)
    assert "전체평균" in ci_row2["비고"]


def test_recommend_method_prefers_grade_when_it_explains_variance():
    yearly = pd.Series({2021: 100, 2022: 100, 2023: 500, 2024: 500, 2025: 100})
    grade_by_year = {2021: "MI", 2022: "MI", 2023: "TI", 2024: "TI", 2025: "MI"}
    method, reason = std.recommend_method(yearly, grade_by_year)
    assert method == "등급별 평균"
    assert "편차" in reason


def test_recommend_method_prefers_recent_when_stable():
    yearly = pd.Series({2021: 100, 2022: 102, 2023: 99, 2024: 101, 2025: 100})
    grade_by_year = {y: "A" for y in yearly.index}
    method, reason = std.recommend_method(yearly, grade_by_year)
    assert method == "최근실적"
    assert "변동" in reason


def test_recommend_method_prefers_five_year_average_when_noisy_and_no_grade_pattern():
    yearly = pd.Series({2019: 80, 2020: 220, 2021: 60, 2022: 240, 2023: 90, 2024: 210, 2025: 70})
    grade_by_year = {y: "A" for y in yearly.index}
    method, _ = std.recommend_method(yearly, grade_by_year)
    assert method == "5개년 평균"


def test_recommend_method_prefers_three_year_average_when_sparse():
    yearly = pd.Series({2023: 80, 2024: 220, 2025: 60})
    grade_by_year = {2023: "A", 2024: "A", 2025: "A"}
    method, _ = std.recommend_method(yearly, grade_by_year)
    assert method == "3개년 평균"


def test_recommend_method_single_year_uses_recent_actual():
    yearly = pd.Series({2025: 500})
    method, reason = std.recommend_method(yearly, {2025: "A"})
    assert method == "최근실적"
    assert "1개년" in reason


def test_compute_uses_auto_recommendation_when_no_explicit_method():
    actuals = pd.DataFrame({
        "사업장": ["강남지사"] * 3,
        "연도": [2023, 2024, 2025],
        "예산과목": ["지급수수료-열원점검수수료"] * 3,
        "금액": [100, 102, 99],
    })
    grades = pd.DataFrame({"사업장": ["강남지사"] * 3, "연도": [2023, 2024, 2025], "등급": ["A", "A", "A"]})
    result = std.compute_standard_amounts(actuals, grades)  # method_map 없음
    row = result.iloc[0]
    assert row["방식출처"] == "자동추천"
    assert row["산출방식"] == "최근실적"
    assert row["추천사유"] != ""


def test_compute_recent_actual_method():
    actuals = pd.DataFrame({
        "사업장": ["강남지사"] * 3,
        "연도": [2023, 2024, 2025],
        "예산과목": ["지급수수료-열원점검수수료"] * 3,
        "금액": [10, 20, 30],
    })
    grades = pd.DataFrame({"사업장": ["강남지사"] * 3, "연도": [2023, 2024, 2025], "등급": ["A", "A", "A"]})
    method_map = pd.DataFrame({"사업장": ["강남지사"], "예산과목": ["지급수수료-열원점검수수료"], "방식": ["최근실적"]})
    result = std.compute_standard_amounts(actuals, grades, method_map)
    assert result.iloc[0]["표준금액"] == 30


def test_compute_n_year_average_methods():
    actuals = pd.DataFrame({
        "사업장": ["강남지사"] * 5,
        "연도": [2021, 2022, 2023, 2024, 2025],
        "예산과목": ["지급수수료-열원점검수수료"] * 5,
        "금액": [10, 20, 30, 40, 50],
    })
    grades = pd.DataFrame({"사업장": ["강남지사"] * 5, "연도": [2021, 2022, 2023, 2024, 2025], "등급": ["A"] * 5})
    method_map = pd.DataFrame({"사업장": ["강남지사"], "예산과목": ["지급수수료-열원점검수수료"], "방식": ["3개년 평균"]})
    result = std.compute_standard_amounts(actuals, grades, method_map)
    assert result.iloc[0]["표준금액"] == pytest.approx((30 + 40 + 50) / 3)


def test_ltsa_and_cri_excluded_when_detail_available():
    actuals = pd.DataFrame({
        "사업장": ["화성지사"] * 4,
        "연도": [2023, 2023, 2024, 2024],
        "예산과목": ["기계장치"] * 4,
        "금액": [1000, 500, 2000, 300],
        "투자유형세부_확정": ["LTSA", "기타", "CRI", "기타"],
    })
    grades = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2023, 2024], "등급": ["MI", "MI"]})
    result = std.compute_standard_amounts(actuals, grades)
    assert result.iloc[0]["표준금액"] == pytest.approx((500 + 300) / 2)


def test_machine_total_used_as_is_without_detail_column():
    actuals = pd.DataFrame({
        "사업장": ["화성지사"] * 2,
        "연도": [2023, 2024],
        "예산과목": ["기계장치"] * 2,
        "금액": [1500, 2300],
    })
    grades = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2023, 2024], "등급": ["MI", "MI"]})
    assert std.has_ltsa_detail(actuals) is False
    result = std.compute_standard_amounts(actuals, grades)
    assert result.iloc[0]["표준금액"] == pytest.approx((1500 + 2300) / 2)


def test_override_takes_precedence():
    actuals = pd.DataFrame({
        "사업장": ["화성지사"], "연도": [2025], "예산과목": ["기계장치"], "금액": [1000],
    })
    grades = pd.DataFrame({"사업장": ["화성지사"], "연도": [2025], "등급": ["MI"]})
    overrides = pd.DataFrame({"사업장": ["화성지사"], "예산과목": ["기계장치"], "등급": ["MI"], "표준금액": [9999]})
    result = std.compute_standard_amounts(actuals, grades, overrides=overrides)
    assert result.iloc[0]["표준금액"] == 9999
    assert result.iloc[0]["산출방식"] == "수동수정"


def test_get_current_grade_returns_latest_year():
    grades = pd.DataFrame({
        "사업장": ["화성지사", "화성지사", "화성지사"],
        "연도": [2023, 2024, 2025],
        "등급": ["TI", "간이", "MI"],
    })
    assert std.get_current_grade("화성지사", grades) == "MI"


def test_get_current_grade_no_history_returns_none():
    grades = pd.DataFrame(columns=["사업장", "연도", "등급"])
    assert std.get_current_grade("화성지사", grades) is None


def _build_sample_workbook(path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("화성")
    ws.cell(row=3, column=2, value=2025)
    ws.cell(row=3, column=3, value=2024)
    ws.cell(row=7, column=1, value="기계장치")
    ws.cell(row=7, column=2, value=3628188.06)  # 천원 단위
    ws.cell(row=7, column=3, value=3798293)
    ws.cell(row=13, column=1, value="자본소계")  # 소계 행 - 실적으로 안 들어가야 함
    ws.cell(row=13, column=2, value=99999)
    ws.cell(row=15, column=1, value="수선유지비-열원경상정비")
    ws.cell(row=15, column=2, value=6074554.26)
    ws.cell(row=15, column=3, value=0)  # 0은 실적으로 안 들어가야 함
    ws.cell(row=21, column=1, value="손익소계")
    ws.cell(row=22, column=1, value=None)  # 보조소계 (건너뜀)
    ws.cell(row=23, column=1, value="#1 GT")
    ws.cell(row=23, column=2, value="MI")
    ws.cell(row=23, column=3, value="TI")
    ws.cell(row=24, column=1, value="ST")
    ws.cell(row=24, column=2, value="A")
    ws.cell(row=24, column=3, value="C")
    wb.save(path)


def test_import_grade_history_prefers_gt_row(tmp_path, monkeypatch):
    site_map = pd.DataFrame([{"사업장": "3070", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900}])
    monkeypatch.setattr(std, "load_site_type_map", lambda: site_map)

    path = tmp_path / "sample.xlsx"
    _build_sample_workbook(path)
    with open(path, "rb") as f:
        result = std.import_grade_history_from_workbook(f.read())

    assert set(zip(result["연도"], result["등급"])) == {(2025, "MI"), (2024, "TI")}
    assert (result["사업장"] == "화성지사").all()


def test_import_actuals_converts_thousand_won_and_matches_accounts(tmp_path, monkeypatch):
    site_map = pd.DataFrame([{"사업장": "3070", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900}])
    monkeypatch.setattr(std, "load_site_type_map", lambda: site_map)

    path = tmp_path / "sample.xlsx"
    _build_sample_workbook(path)
    with open(path, "rb") as f:
        result = std.import_actuals_from_workbook(f.read())

    assert (result["사업장"] == "화성지사").all()
    lookup = {(r["연도"], r["예산과목"]): r["금액"] for _, r in result.iterrows()}
    assert lookup[(2025, "기계장치")] == pytest.approx(3628188.06 * 1000)
    assert lookup[(2024, "기계장치")] == pytest.approx(3798293 * 1000)
    assert lookup[(2025, "수선유지비-열원경상정비")] == pytest.approx(6074554.26 * 1000)


def test_import_actuals_excludes_subtotal_rows_and_zero_amounts(tmp_path, monkeypatch):
    site_map = pd.DataFrame([{"사업장": "3070", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900}])
    monkeypatch.setattr(std, "load_site_type_map", lambda: site_map)

    path = tmp_path / "sample.xlsx"
    _build_sample_workbook(path)
    with open(path, "rb") as f:
        result = std.import_actuals_from_workbook(f.read())

    assert "자본소계" not in result["예산과목"].values
    assert (2024, "수선유지비-열원경상정비") not in set(zip(result["연도"], result["예산과목"]))


def test_import_actuals_ignores_unmapped_labels(tmp_path, monkeypatch):
    site_map = pd.DataFrame([{"사업장": "3070", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900}])
    monkeypatch.setattr(std, "load_site_type_map", lambda: site_map)

    path = tmp_path / "sample.xlsx"
    _build_sample_workbook(path)
    with open(path, "rb") as f:
        result = std.import_actuals_from_workbook(f.read())

    # '#1 GT'/'ST' 정비등급 행(23~24행)은 손익소계 이후라 실적 대상 범위 밖이다.
    assert "#1 GT" not in result["예산과목"].values
    assert "ST" not in result["예산과목"].values
