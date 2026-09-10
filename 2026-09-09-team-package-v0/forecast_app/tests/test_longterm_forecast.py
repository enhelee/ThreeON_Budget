import pandas as pd
import pytest
from openpyxl import Workbook
import longterm_forecast as lf


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(lf, "FACTORS_PATH", str(tmp_path / "longterm_factors.csv"))
    monkeypatch.setattr(lf, "HOT_PARTS_PATH", str(tmp_path / "hot_parts_plan.csv"))
    monkeypatch.setattr(lf, "HQ_MASTER_PATH", str(tmp_path / "hq_master_table.csv"))
    monkeypatch.setattr(lf, "HQ_RATIO_PATH", str(tmp_path / "hq_allocation_ratio.csv"))
    monkeypatch.setattr(lf, "HQ_TEMP_PROJECTS_PATH", str(tmp_path / "hq_temp_projects.csv"))
    monkeypatch.setattr(lf, "SURPRISE_PATH", str(tmp_path / "surprise_projects.csv"))
    yield tmp_path


# ---------------------------------------------------------------- 팩터
def test_year_multiplier_single_factor_compounds():
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.1, "활성": True}])
    assert lf.year_multiplier(2026, factors) == pytest.approx(1.0)
    assert lf.year_multiplier(2027, factors) == pytest.approx(1.1)
    assert lf.year_multiplier(2028, factors) == pytest.approx(1.21)


def test_year_multiplier_multiplies_multiple_active_factors():
    factors = pd.DataFrame([
        {"팩터명": "물가상승률", "연간비율": 0.1, "활성": True},
        {"팩터명": "노후화", "연간비율": 0.05, "활성": True},
    ])
    assert lf.year_multiplier(2027, factors) == pytest.approx(1.1 * 1.05)


def test_year_multiplier_ignores_inactive_factor():
    factors = pd.DataFrame([
        {"팩터명": "물가상승률", "연간비율": 0.1, "활성": True},
        {"팩터명": "노후화", "연간비율": 0.5, "활성": False},
    ])
    assert lf.year_multiplier(2027, factors) == pytest.approx(1.1)


def test_year_multiplier_before_base_year_is_one():
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.1, "활성": True}])
    assert lf.year_multiplier(2020, factors) == 1.0


def test_load_factors_seeds_default_inflation(isolated):
    factors = lf.load_factors()
    assert factors.iloc[0]["팩터명"] == "물가상승률"


# ---------------------------------------------------------------- 정비등급 스케줄 파서
def _build_schedule_workbook(site_rows: dict) -> bytes:
    """site_rows: {지사라벨: [(GT값,ST값), ...]} - 2026,2027 두 해만 구성한 최소 워크북."""
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "정기점검보수공사 일정(2026~2041)"
    ws.cell(row=1, column=1, value="구 분")
    ws.cell(row=1, column=3, value=2026)
    ws.cell(row=1, column=6, value=2027)
    ws.cell(row=2, column=3, value="등급")
    ws.cell(row=2, column=6, value="등급")
    ws.cell(row=3, column=3, value="GT")
    ws.cell(row=3, column=4, value="ST")
    ws.cell(row=3, column=5, value="기간")
    ws.cell(row=3, column=6, value="GT")
    ws.cell(row=3, column=7, value="ST")
    ws.cell(row=3, column=8, value="기간")

    r = 4
    for site, values in site_rows.items():
        ws.cell(row=r, column=1, value=site)
        (gt26, st26), (gt27, st27) = values
        ws.cell(row=r, column=3, value=gt26)
        ws.cell(row=r, column=4, value=st26)
        ws.cell(row=r, column=6, value=gt27)
        ws.cell(row=r, column=7, value=st27)
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_schedule_parser_prefers_gt_when_present(isolated):
    data = _build_schedule_workbook({"화 성": [("간이", "C"), ("TI", "B")]})
    result = lf.import_schedule_from_workbook(data)
    assert set(result["사업장"]) == {"화성지사"}
    grades = dict(zip(result["연도"], result["등급"]))
    assert grades == {2026: "간이", 2027: "TI"}


def test_schedule_parser_falls_back_to_st_when_no_gt_token(isolated):
    data = _build_schedule_workbook({"광주전남": [("-", "A"), ("-", "C")]})
    result = lf.import_schedule_from_workbook(data)
    assert set(result["사업장"]) == {"광주전남지사"}
    grades = dict(zip(result["연도"], result["등급"]))
    assert grades == {2026: "A", 2027: "C"}


def test_schedule_parser_skips_placeholder_years(isolated):
    data = _build_schedule_workbook({"수 원": [("-", "-"), ("간이", "C")]})
    result = lf.import_schedule_from_workbook(data)
    assert list(result["연도"]) == [2027]


# ---------------------------------------------------------------- 고온부품 파서
def test_hot_parts_parser_converts_thousand_won_to_won(isolated):
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "고온부품(26)"
    ws.cell(row=4, column=1, value="구  분")
    ws.cell(row=4, column=6, value="26년")
    ws.cell(row=4, column=7, value="27년")
    ws.cell(row=5, column=1, value="화성")
    ws.cell(row=5, column=4, value="기계장치")
    ws.cell(row=5, column=5, value="고온부품재생")
    ws.cell(row=5, column=6, value=1000)
    ws.cell(row=5, column=7, value=0)
    ws.cell(row=6, column=4, value="건설중인자산")
    ws.cell(row=6, column=5, value="신품구매")
    ws.cell(row=6, column=6, value=0)
    ws.cell(row=6, column=7, value=2000)
    buf = io.BytesIO()
    wb.save(buf)

    result = lf.import_hot_parts_from_workbook(buf.getvalue())
    assert len(result) == 2
    recycled = result[result["항목"] == "고온부품재생"].iloc[0]
    assert recycled["사업장"] == "화성지사"
    assert recycled["연도"] == 2026
    assert recycled["금액"] == 1_000_000  # 1000천원 -> 원


# ---------------------------------------------------------------- 본사 원가분배 마스터 파서
def _build_hq_master_workbook() -> bytes:
    import io
    wb = Workbook()
    ws = wb.active
    ws.title = "26년 본사 원가분배"
    ws.cell(row=5, column=1, value="총계")
    ws.cell(row=5, column=4, value="26년 예산")
    ws.cell(row=5, column=5, value="화성")
    ws.cell(row=5, column=6, value="동탄")

    ws.cell(row=6, column=1, value="플랜트")
    ws.cell(row=6, column=2, value="경상정비")
    ws.cell(row=6, column=3, value="중대형")
    ws.cell(row=6, column=4, value=1000)
    ws.cell(row=6, column=5, value=600)
    ws.cell(row=6, column=6, value=400)

    ws.cell(row=7, column=3, value="합계")  # 소계 행 - 제외되어야 함
    ws.cell(row=7, column=4, value=1000)
    ws.cell(row=7, column=5, value=600)
    ws.cell(row=7, column=6, value=400)

    ws.cell(row=8, column=1, value="합계")  # 전체 합계 라벨 - 여기서 중단
    ws.cell(row=8, column=4, value=1000)

    ws.cell(row=9, column=1, value="경상정비 원가배부기준")
    ws.cell(row=10, column=3, value="합계")
    ws.cell(row=10, column=4, value="계약체결금액")
    ws.cell(row=10, column=5, value=6000)
    ws.cell(row=10, column=6, value=4000)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hq_master_parser_excludes_subtotal_rows(isolated):
    master, ratio = lf.import_hq_master_from_workbook(_build_hq_master_workbook())
    assert len(master) == 1  # '합계' 소계 행은 제외
    assert master.iloc[0]["화성지사"] == 600
    assert master.iloc[0]["동탄지사"] == 400


def test_hq_ratio_parser_reads_contract_totals(isolated):
    master, ratio = lf.import_hq_master_from_workbook(_build_hq_master_workbook())
    ratio_dict = dict(zip(ratio["사업장"], ratio["계약체결금액"]))
    assert ratio_dict == {"화성지사": 6000, "동탄지사": 4000}


def test_hq_amount_by_site_account_maps_category_correctly(isolated):
    master, _ = lf.import_hq_master_from_workbook(_build_hq_master_workbook())
    amt = lf.hq_amount_by_site_account(master)
    assert amt[("화성지사", "수선유지비-열원경상정비")] == 600
    assert amt[("동탄지사", "수선유지비-열원경상정비")] == 400


# ---------------------------------------------------------------- 본사 일시적 사업 배분
def test_allocate_hq_temp_projects_uses_ratio():
    temp_df = pd.DataFrame([{"사업명": "긴급공사", "예산과목": "기계장치", "연도": 2026, "금액": 1000.0}])
    ratio_df = pd.DataFrame([
        {"사업장": "화성지사", "계약체결금액": 600.0},
        {"사업장": "동탄지사", "계약체결금액": 400.0},
    ])
    result = lf.allocate_hq_temp_projects(temp_df, ratio_df)
    amounts = dict(zip(result["사업장"], result["금액"]))
    assert amounts["화성지사"] == pytest.approx(600.0)
    assert amounts["동탄지사"] == pytest.approx(400.0)


def test_allocate_hq_temp_projects_empty_inputs_return_empty():
    empty = pd.DataFrame(columns=["사업명", "예산과목", "연도", "금액"])
    ratio_df = pd.DataFrame(columns=["사업장", "계약체결금액"])
    result = lf.allocate_hq_temp_projects(empty, ratio_df)
    assert result.empty


# ---------------------------------------------------------------- 지사별 표 계산
def test_compute_site_table_applies_grade_specific_standard_and_inflation():
    grade_hist = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2026, 2027], "등급": ["간이", "TI"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "TI", "표준금액": 200.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.1, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(100.0)
    assert row[2027] == pytest.approx(200.0 * 1.1)


def test_compute_site_table_uses_latest_known_grade_for_unlisted_future_years():
    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2030] == pytest.approx(100.0)  # 미래 연도에도 마지막으로 알려진 등급(간이)을 그대로 씀


def test_compute_site_table_hot_parts_passthrough_without_factor():
    grade_hist = pd.DataFrame(columns=["사업장", "연도", "등급"])
    standard_df = pd.DataFrame(columns=["사업장", "구분", "예산과목", "등급", "표준금액"])
    hot_parts = pd.DataFrame([{"사업장": "화성지사", "연도": 2027, "항목": "고온부품재생", "금액": 5000.0}])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.5, "활성": True}])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, hot_parts, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment)
    row = table[table["예산과목"] == "기계장치_고온부품재생"].iloc[0]
    assert row[2027] == 5000.0  # 팩터 미적용, 업로드값 그대로


# ---------------------------------------------------------------- 당해년도 예산계획 override
def test_budget_amount_by_site_account_aggregates_by_site_and_account(monkeypatch):
    budget_df = pd.DataFrame([
        {"예산귀속 \n부서코드": "H001", "예산과목": "수선유지비-열원경상정비", "연예산 합계": 100.0},
        {"예산귀속 \n부서코드": "H001", "예산과목": "수선유지비-열원경상정비", "연예산 합계": 50.0},
        {"예산귀속 \n부서코드": "H001", "예산과목": "기계장치", "연예산 합계": 30.0},
    ])
    monkeypatch.setattr(lf, "get_site_display_name", lambda code: "화성지사")
    lookup = lf.budget_amount_by_site_account(budget_df)
    assert lookup[("화성지사", "수선유지비-열원경상정비")] == 150.0
    assert lookup[("화성지사", "기계장치")] == 30.0


def test_budget_amount_by_site_account_empty_df_returns_empty_dict():
    assert lf.budget_amount_by_site_account(pd.DataFrame()) == {}


def test_compute_site_table_uses_budget_plan_for_base_year_when_available():
    grade_hist = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2026, 2027], "등급": ["간이", "간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.1, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])
    budget_lookup = {("화성지사", "수선유지비-열원경상정비"): 555.0}

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment,
                                   budget_lookup)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(555.0)  # 당해년도 - 예산계획 확정값
    assert row[2027] == pytest.approx(100.0 * 1.1)  # 다음년도부터 - 표준화 금액 x 팩터 (예산계획과 무관)


def test_compute_site_table_falls_back_to_standard_when_budget_plan_missing_entry():
    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])
    budget_lookup = {("동탄지사", "수선유지비-열원경상정비"): 999.0}  # 다른 지사 - 화성지사엔 매칭 안 됨

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment,
                                   budget_lookup)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(100.0)  # 예산계획에 없으면 표준화 금액으로 대체


def test_compute_site_table_adds_hq_master_allocation_on_top_of_budget_plan_for_base_year():
    """원본 양식의 "26년" 시트는 총원가배분(전체) = 본사 + 지사로 나뉜다 - 본사 원가분배 마스터 표
    (hq_amount_by_site_acct)가 이 계정을 다루면(경상정비 등 HQ_CATEGORY_TO_ACCOUNT에 있는 계정)
    그 배분값은 지사 자체 몫(예산계획)을 대체하는 게 아니라 더해진다. 서로 다른 부서코드에 잡힌
    별도 예산이라 합쳐야 그 지사의 26년 총액이 된다."""
    grade_hist = pd.DataFrame({"사업장": ["화성지사"] * 2, "연도": [2026, 2027], "등급": ["간이", "간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.1, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])
    hq_amount = {("화성지사", "수선유지비-열원경상정비"): 777.0}
    budget_lookup = {("화성지사", "수선유지비-열원경상정비"): 555.0}

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, hq_amount,
                                   empty_temp_alloc, factors, empty_surprise, empty_investment,
                                   budget_lookup)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(555.0 + 777.0)  # 지사 몫(예산계획) + 본사 배분 몫
    assert row[2027] == pytest.approx(100.0 * 1.1)  # 다음년도부터는 여전히 표준화 x 팩터 (본사배분과 무관)


def test_compute_site_table_hq_master_zero_when_other_site_has_entry():
    """본사 원가분배 마스터 표에 이 지사·계정 배분이 없으면(예: 마스터 표에 없는 신규 지사) 본사 몫은
    0으로 더해져 예산계획 금액만 그대로 남는다."""
    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])
    hq_amount = {("동탄지사", "수선유지비-열원경상정비"): 777.0}  # 다른 지사 - 화성지사엔 매칭 안 됨
    budget_lookup = {("화성지사", "수선유지비-열원경상정비"): 555.0}

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, hq_amount,
                                   empty_temp_alloc, factors, empty_surprise, empty_investment,
                                   budget_lookup)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(555.0)  # 본사배분에 없으면 0을 더한 것과 같아 예산계획 그대로


def test_compute_site_table_building_line_adds_hq_master_via_label_alias():
    """"수선유지비-건물/구축물"(표준화 라인)의 본사 배분 몫은 본사 원가분배 마스터 표에서
    "건물/구축물"(접두어 없는 라벨, HQ_CATEGORY_TO_ACCOUNT["건물구축물"])로 저장돼 있어 표기가
    다르다 - STANDARDIZED_ACCOUNT_TO_HQ_LABEL 별칭으로 찾아서 더해야 한다."""
    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-건물/구축물", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])
    hq_amount = {("화성지사", "건물/구축물"): 242.0}  # 마스터 표 라벨은 "건물/구축물"(접두어 없음)

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, hq_amount,
                                   empty_temp_alloc, factors, empty_surprise, empty_investment)
    row = table[table["예산과목"] == "수선유지비-건물/구축물"].iloc[0]
    assert row[2026] == pytest.approx(100.0 + 242.0)  # 표준화값(지사 몫) + 본사 배분(별칭으로 매칭)


def test_compute_site_table_without_budget_lookup_behaves_as_before():
    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    standard_df = pd.DataFrame([
        {"사업장": "화성지사", "구분": "손익", "예산과목": "수선유지비-열원경상정비", "등급": "간이", "표준금액": 100.0},
    ])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    empty_surprise = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액", "사유"])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, empty_surprise, empty_investment)
    row = table[table["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[2026] == pytest.approx(100.0)


def test_compute_site_table_adds_surprise_project():
    grade_hist = pd.DataFrame(columns=["사업장", "연도", "등급"])
    standard_df = pd.DataFrame(columns=["사업장", "구분", "예산과목", "등급", "표준금액"])
    empty_hot = pd.DataFrame(columns=["사업장", "연도", "항목", "금액"])
    factors = pd.DataFrame([{"팩터명": "물가상승률", "연간비율": 0.0, "활성": True}])
    empty_temp_alloc = pd.DataFrame(columns=["사업장", "연도", "예산과목", "금액"])
    surprise = pd.DataFrame([
        {"사업장": "화성지사", "연도": 2028, "예산과목": "공구와기구-열원시설공기구", "금액": 999.0, "사유": "돌발"},
    ])
    empty_investment = pd.DataFrame(columns=["사업장", "지사유형", "금액"])

    table = lf.compute_site_table("화성지사", grade_hist, standard_df, empty_hot, {},
                                   empty_temp_alloc, factors, surprise, empty_investment)
    row = table[table["예산과목"] == "공구와기구-열원시설공기구"].iloc[0]
    assert row[2028] == 999.0
    assert row[2027] == 0.0


# ---------------------------------------------------------------- Test 모드 예산계획(연도별 저장)
def test_save_and_load_test_budget_plan_is_per_year(isolated):
    """서로 다른 연도로 저장하면 각각 별도 파일에 저장되고, 같은 연도만 지정해서 불러온다."""
    df_2025 = pd.DataFrame([{"사업명": "2025예산", "예산과목": "기계장치", "연예산 합계": 100.0}])
    df_2026 = pd.DataFrame([{"사업명": "2026예산", "예산과목": "기계장치", "연예산 합계": 200.0}])
    lf.save_test_budget_plan(df_2025, year=2025)
    lf.save_test_budget_plan(df_2026, year=2026)

    assert lf.load_test_budget_plan(2025).iloc[0]["사업명"] == "2025예산"
    assert lf.load_test_budget_plan(2026).iloc[0]["사업명"] == "2026예산"
    assert lf.load_test_budget_plan_years() == [2025, 2026]

    both = lf.load_test_budget_plan_multi([2025, 2026])
    assert set(both["사업명"]) == {"2025예산", "2026예산"}


def test_load_test_budget_plan_defaults_to_base_year(isolated):
    """연도를 지정하지 않으면 BASE_YEAR(당해년도) 파일을 쓴다 - 중장기예산 Test 모드가 기존과 동일하게 동작."""
    df = pd.DataFrame([{"사업명": "당해년도예산", "예산과목": "기계장치", "연예산 합계": 100.0}])
    lf.save_test_budget_plan(df)  # year 생략

    assert lf.load_test_budget_plan().iloc[0]["사업명"] == "당해년도예산"  # year 생략
    assert lf.load_test_budget_plan(lf.BASE_YEAR).iloc[0]["사업명"] == "당해년도예산"


def test_load_test_budget_plan_migrates_legacy_single_file(isolated):
    """연도 구분 없이 쓰던 예전 단일 파일(TEST_BUDGET_PLAN_PATH)이 남아있으면, BASE_YEAR 파일로
    1회 옮겨써서 기존 백데이터를 잃지 않는다."""
    import os
    os.makedirs(lf.TEST_DATA_DIR, exist_ok=True)
    legacy_df = pd.DataFrame([{"사업명": "레거시예산", "예산과목": "기계장치", "연예산 합계": 300.0}])
    legacy_df.to_csv(lf.TEST_BUDGET_PLAN_PATH, index=False)

    loaded = lf.load_test_budget_plan(lf.BASE_YEAR)
    assert loaded.iloc[0]["사업명"] == "레거시예산"
    assert os.path.exists(lf.test_budget_plan_path(lf.BASE_YEAR))  # 새 경로로 마이그레이션됨
