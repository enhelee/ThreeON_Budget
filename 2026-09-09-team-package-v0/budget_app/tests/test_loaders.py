import openpyxl
from budget import loaders


def _make_plan(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    hdr = ["주관부서명", "부서코드", "처지사", "부서부", "속성", "예산코드", "예산과목", "사업명", "산출내역", "연예산"]
    for i, h in enumerate(hdr, 1):
        ws.cell(3, i, h)
    for j, row in enumerate(rows, 4):
        for i, v in enumerate(row, 1):
            ws.cell(j, i, v)
    wb.save(path)


def test_load_business_plan_basic(tmp_path):
    path = str(tmp_path / "plan.xlsx")
    rows = [
        ["플랜트기술처 기계기술부", "1000093", "화성지사", "고객지원부", "제조", "60909007",
         "수선유지비-열원경상정비", "경상정비공사", "-", 30000],
        [None, None, None, None, None, None, None, None, None, None],  # 완전 빈 행 -> 제외
        ["건설처 건설관리부", "1000029", "건설처", "건설관리부", "자산", "10101",
         "건물", "신축 건물", "-", 120000],
    ]
    _make_plan(path, rows)
    df = loaders.load_business_plan(path)
    assert len(df) == 2
    assert list(df.columns) == loaders.PLAN_COLUMNS + ["_row"]
    assert df.iloc[0]["연예산"] == 30000.0
    assert df.iloc[0]["처지사"] == "화성지사"
    assert df.iloc[0]["_row"] == 4
    assert df.iloc[1]["_row"] == 6


def test_load_business_plan_missing_amount_stays_none(tmp_path):
    path = str(tmp_path / "plan2.xlsx")
    rows = [
        ["주관", "1000001", "강남지사", "부", "제조", "1", "기계장치", "사업A", "-", None],
    ]
    _make_plan(path, rows)
    df = loaders.load_business_plan(path)
    assert df.iloc[0]["연예산"] is None
