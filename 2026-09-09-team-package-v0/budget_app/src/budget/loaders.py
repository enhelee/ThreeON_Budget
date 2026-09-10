# -*- coding: utf-8 -*-
"""사업별 예산 업로드 파일 로드 (손익/자본 구분 없는 전체 사업)."""
import pandas as pd
import openpyxl

PLAN_COLUMNS = [
    "주관부서명", "부서코드", "처지사", "부서부", "속성",
    "예산코드", "예산과목", "사업명", "산출내역", "연예산",
]


def load_business_plan(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["양식1(월별)"]
    recs = []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 11)]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        rec = dict(zip(PLAN_COLUMNS, vals))
        amt = rec["연예산"]
        rec["연예산"] = float(amt) if isinstance(amt, (int, float)) else None
        rec["_row"] = r
        recs.append(rec)
    wb.close()
    cols = PLAN_COLUMNS + ["_row"]
    return pd.DataFrame(recs, columns=cols)
