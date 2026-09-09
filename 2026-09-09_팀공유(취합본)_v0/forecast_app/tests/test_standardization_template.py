from openpyxl import Workbook
from standardization_template import _set_cell


def test_set_cell_writes_normal_cell():
    wb = Workbook()
    ws = wb.active
    _set_cell(ws, 1, 1, "값")
    assert ws.cell(row=1, column=1).value == "값"


def test_set_cell_unmerges_and_writes_when_target_is_merged():
    """일부 지사 시트에 남아있던 빈 병합 셀(예: 삼송 D34:D39)에 쓰려 하면 openpyxl이
    'MergedCell' object attribute 'value' is read-only 오류를 낸다 - 이를 재현하고 고쳤는지 확인한다.
    병합 범위의 왼쪽 위(anchor, D34)는 원래도 쓸 수 있는 진짜 Cell이라 문제가 없고,
    그 아래(D35 이하)가 MergedCell이 되어 오류가 나던 자리다."""
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("D34:D39")

    _set_cell(ws, 34, 4, "값1")  # anchor - 원래도 문제 없던 자리
    _set_cell(ws, 35, 4, "값2")  # MergedCell이라 그냥 두면 크래시 나는 자리

    assert len(ws.merged_cells.ranges) == 0  # 겹치는 병합은 풀렸다
    assert ws.cell(row=34, column=4).value == "값1"
    assert ws.cell(row=35, column=4).value == "값2"  # 병합이 풀렸으니 서로 다른 값을 가질 수 있다


def test_set_cell_only_unmerges_the_overlapping_range():
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("D34:D39")
    ws.merge_cells("F1:F3")  # 무관한 다른 병합은 그대로 남아있어야 한다

    _set_cell(ws, 35, 4, "값")  # D35는 anchor가 아닌 MergedCell이라 실제로 병합을 풀게 된다

    remaining = [str(r) for r in ws.merged_cells.ranges]
    assert remaining == ["F1:F3"]
