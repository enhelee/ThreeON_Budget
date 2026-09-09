import io
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
import longterm_forecast_template as lft


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    yield tmp_path


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _schedule_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "정기점검보수공사 일정(2026~2041)"
    ws.cell(row=1, column=1, value="구 분")
    ws.cell(row=1, column=3, value=2026)
    ws.cell(row=1, column=6, value=2027)
    ws.cell(row=2, column=3, value="등급")
    ws.cell(row=2, column=6, value="등급")
    ws.cell(row=3, column=3, value="GT")
    ws.cell(row=3, column=4, value="ST")
    ws.cell(row=3, column=5, value="기간")
    ws.cell(row=3, column=6, value="GT")
    ws.cell(row=3, column=7, value="ST")
    ws.cell(row=3, column=8, value="기간")
    ws.cell(row=4, column=1, value="화 성")  # 원본과의 미리채워진 값(수정 전)
    ws.cell(row=4, column=3, value="MI")
    ws.cell(row=4, column=6, value="MI")
    return _save(wb)


def test_fill_schedule_sheet_overwrites_with_current_grade_history(isolated):
    template_path = isolated / "sched.xlsx"
    template_path.write_bytes(_schedule_template_bytes())

    grade_hist = pd.DataFrame({"사업장": ["화성지사", "화성지사"], "연도": [2026, 2027], "등급": ["간이", "TI"]})
    data = lft.build_schedule_download(grade_hist, template_path=str(template_path))

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["정기점검보수공사 일정(2026~2041)"]
    assert ws.cell(row=4, column=3).value == "간이"  # 2026 - 원래 'MI'였던 자리를 덮어씀
    assert ws.cell(row=4, column=6).value == "TI"  # 2027


def test_fill_schedule_sheet_ignores_years_without_history(isolated):
    template_path = isolated / "sched.xlsx"
    template_path.write_bytes(_schedule_template_bytes())

    grade_hist = pd.DataFrame({"사업장": ["화성지사"], "연도": [2026], "등급": ["간이"]})
    data = lft.build_schedule_download(grade_hist, template_path=str(template_path))

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["정기점검보수공사 일정(2026~2041)"]
    assert ws.cell(row=4, column=3).value == "간이"
    assert ws.cell(row=4, column=6).value == "MI"  # 2027 이력 없음 - 원본 값 그대로 남음


def _hot_parts_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "고온부품(26)"
    ws.cell(row=4, column=1, value="구  분")
    ws.cell(row=4, column=6, value="26년")
    ws.cell(row=4, column=7, value="27년")
    ws.cell(row=5, column=1, value="화성")
    ws.cell(row=5, column=4, value="기계장치")
    ws.cell(row=5, column=5, value="고온부품재생")
    ws.cell(row=5, column=6, value=999)  # 원본 값(수정 전)
    ws.cell(row=5, column=7, value=0)
    return _save(wb)


def test_fill_hot_parts_sheet_round_trips_won_to_thousand_won(isolated):
    template_path = isolated / "hp.xlsx"
    template_path.write_bytes(_hot_parts_template_bytes())

    hot_parts = pd.DataFrame([{"사업장": "화성지사", "연도": 2026, "항목": "고온부품재생", "금액": 1_000_000}])
    data = lft.build_hot_parts_download(hot_parts, template_path=str(template_path))

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["고온부품(26)"]
    assert ws.cell(row=5, column=6).value == 1000  # 1,000,000원 -> 1000천원


def _hq_master_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "26년 본사 원가분배"
    ws.cell(row=5, column=1, value="총계")
    ws.cell(row=5, column=4, value="26년 예산")
    ws.cell(row=5, column=5, value="화성")
    ws.cell(row=5, column=6, value="동탄")
    ws.cell(row=6, column=1, value="플랜트")
    ws.cell(row=6, column=2, value="경상정비")
    ws.cell(row=6, column=3, value="중대형")
    ws.cell(row=6, column=4, value=1000)
    ws.cell(row=6, column=5, value=1)  # 원본 값(수정 전)
    ws.cell(row=6, column=6, value=1)
    ws.cell(row=7, column=1, value="합계")
    return _save(wb)


def test_fill_hq_master_sheet_overwrites_with_current_master_table(isolated):
    template_path = isolated / "hq.xlsx"
    template_path.write_bytes(_hq_master_template_bytes())

    master = pd.DataFrame([{
        "대분류": "플랜트", "중분류": "경상정비", "세부내역": "중대형", "26년예산": 1000,
        "화성지사": 600, "동탄지사": 400,
    }])
    data = lft.build_hq_master_download(master, template_path=str(template_path))

    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["26년 본사 원가분배"]
    assert ws.cell(row=6, column=5).value == 600
    assert ws.cell(row=6, column=6).value == 400


def _summary_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "총괄표"
    ws.cell(row=3, column=1, value="총괄표")
    ws.cell(row=3, column=3, value="26년")
    ws.cell(row=4, column=1, value="수선유지비")
    ws.cell(row=4, column=2, value="건물/구축물")
    ws.cell(row=5, column=2, value="열원정기점검")
    ws.cell(row=6, column=1, value="건물")  # 자본계정 - 배분 데이터 없어 손대면 안 됨
    return _save(wb)


def test_fill_summary_sheet_resolves_profit_loss_account_not_orphaned_hq_line(isolated):
    """'수선유지비-건물/구축물'(표준화 실적)과 orphan 라벨 '건물/구축물'(본사배분)이
    총괄표에서는 둘 다 같은 원본 셀 텍스트('건물/구축물')로 보이므로, 실제 값이
    들어있는 표준화 계정으로 정확히 매칭되어야 한다."""
    template_path = isolated / "summary.xlsx"
    template_path.write_bytes(_summary_template_bytes())

    site_table = pd.DataFrame([
        {"예산과목": "수선유지비-건물/구축물", 2026: 5_000_000},
        {"예산과목": "건물/구축물", 2026: 1},  # 본사배분 orphan 라벨 - 절대 이 값이 쓰이면 안 됨
        {"예산과목": "수선유지비-열원정기점검", 2026: 2_000_000},
    ])
    site_tables = {"화성지사": site_table}

    wb = load_workbook(str(template_path))
    ws = wb["총괄표"]
    lft.fill_summary_sheet(ws, site_tables)

    assert ws.cell(row=4, column=3).value == 5.0
    assert ws.cell(row=5, column=3).value == 2.0
    assert ws.cell(row=6, column=3).value is None  # 자본계정 건물 - 데이터 소스 없음, 그대로 유지
