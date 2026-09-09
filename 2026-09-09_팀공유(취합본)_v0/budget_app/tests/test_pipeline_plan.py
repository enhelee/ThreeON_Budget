import os
import openpyxl
from budget import pipeline_plan, config_store


def _make_plan(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    hdr = ["주관부서명", "부서코드", "처지사", "부서부", "속성", "예산코드", "예산과목", "사업명", "산출내역", "연예산"]
    for i, h in enumerate(hdr, 1):
        ws.cell(3, i, h)
    rows = [
        ["플랜트기술처 기계기술부", "1000093", "플랜트기술처", "기계기술부", "제조", "60909001",
         "수선유지비-건물/구축물", "본관 지붕 방수공사", "-", 60000],
        ["화성지사 고객지원부", "3050001", "화성지사", "고객지원부", "제조", "60909007",
         "수선유지비-열원경상정비", "경상정비공사", "-", 30000],
        ["건설처 건설관리부", "1000029", "건설처", "건설관리부", "자산", "10101",
         "건물", "신축 건물", "-", 120000],
        ["담당부", "", "강남지사", "부", "제조", "", "이상한과목", "미분류테스트", "-", 5],
    ]
    for j, row in enumerate(rows, 4):
        for i, v in enumerate(row, 1):
            ws.cell(j, i, v)
    wb.save(path)


def test_run_plan_pl_and_capital(tmp_path):
    plan_path = str(tmp_path / "plan.xlsx")
    _make_plan(plan_path)
    config_dir = str(tmp_path / "config")
    out_dir = str(tmp_path / "output")

    res_pl = pipeline_plan.run_plan(plan_path, config_dir, out_dir, "2026", "손익")
    assert os.path.exists(res_pl["output_path"])
    # 손익 시드 과목 2종(수선유지비-건물/구축물 60000 + 수선유지비-열원경상정비 30000) 매칭
    assert res_pl["요약"]["행수"] == 2
    assert res_pl["요약"]["총액"] == 90000.0
    assert "이상한과목" in res_pl["요약"]["미분류과목"]
    assert res_pl["요약"]["결측치건수"] == 0  # 이상한과목 행은 부서코드/예산코드만 비어있음(필수아님)

    res_cap = pipeline_plan.run_plan(plan_path, config_dir, out_dir, "2026", "자본")
    assert os.path.exists(res_cap["output_path"])
    assert res_cap["요약"]["행수"] == 1
    assert res_cap["요약"]["총액"] == 120000.0

    # 지사구성이 자동 생성/재사용되는지 확인
    assert os.path.exists(os.path.join(config_dir, "지사구성_2026.json"))
