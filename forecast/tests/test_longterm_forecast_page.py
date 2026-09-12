"""views/longterm_forecast_page.py의 render() 전체 경로(디스크의 classified_*.csv/budget_2026.csv 읽기 ->
compute_all_sites)를 AppTest로 검증한다. compute_site_table 자체의 override 로직은
tests/test_longterm_forecast.py에서 이미 단위 테스트로 커버되어 있으므로, 여기서는
'실제 화면이 그 값을 정확히 배선하는지'만 확인한다."""
import io
import shutil
import pandas as pd
import pytest
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest
import longterm_forecast as lf
import longterm_forecast_template as lft

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_PROJECT_ROOT = lft.__file__.rsplit("longterm_forecast_template.py", 1)[0]


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    # render()가 기본 template_path(BUILTIN_TEMPLATE_PATH, cwd 상대경로)로 재무팀 양식을 열기 때문에,
    # 격리된 tmp_path에도 실제 내장 양식 파일을 복사해둬야 한다.
    shutil.copy(_PROJECT_ROOT + lft.BUILTIN_TEMPLATE_PATH, tmp_path / lft.BUILTIN_TEMPLATE_PATH)
    monkeypatch.chdir(tmp_path)
    yield tmp_path


def _run_render():
    import views.longterm_forecast_page as page
    page.render()


def _run_render_test():
    import views.longterm_forecast_page as page
    page.render_test()


def _budget_workbook_bytes(dept_code: str, account_name: str, amount_thousand: float, project_name: str) -> bytes:
    """예산계획 업로드 양식(양식1(월별)/예산코드/부서코드) 최소본 - '연예산 합계'는 천원 단위(원본 양식과 동일)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    ws.append(["[단위 : 천원, 부가세 별도]"])
    ws.append([])
    ws.append(["사업명", "예산귀속 \n부서코드", "예산코드", "예산과목", "연예산 합계"])
    ws.append([project_name, dept_code, "9999", account_name, amount_thousand])

    ws_account = wb.create_sheet("예산코드")
    ws_account.append(["예산코드", "예산과목명"])
    ws_account.append(["9999", account_name])

    ws_dept = wb.create_sheet("부서코드")
    ws_dept.append(["부서코드", "부서명"])
    ws_dept.append([dept_code, "화성지사"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_render_base_year_column_uses_budget_plan_not_standardized_amount(isolated):
    """예산계획(budget_2026.csv)에 화성지사 '수선유지비-열원경상정비' 확정값(5.55억)이 있으면,
    실적 기반 표준화값(1억)이 아니라 그 확정값이 미리보기 표의 2026년 칸에 그대로 반영돼야 한다."""
    tmp_path = isolated

    site_type_map = pd.DataFrame([
        {"사업장": "9001", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900},
        {"사업장": "700001", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900},
    ])
    site_type_map.to_csv(tmp_path / "site_type_map.csv", index=False)

    classified = pd.DataFrame([
        {"전기일": "2024-03-15", "사업장": "9001", "계정과목": "수선유지비-열원경상정비", "금액": 100_000_000.0},
        {"전기일": "2025-03-15", "사업장": "9001", "계정과목": "수선유지비-열원경상정비", "금액": 100_000_000.0},
    ])
    classified.to_csv(tmp_path / "classified_2025.csv", index=False)

    grade_hist = pd.DataFrame([
        {"사업장": "화성지사", "연도": 2024, "등급": "간이"},
        {"사업장": "화성지사", "연도": 2025, "등급": "간이"},
    ])
    grade_hist.to_csv(tmp_path / "site_year_grade.csv", index=False)

    budget_df = pd.DataFrame([
        {"예산귀속 \n부서코드": "700001", "예산과목": "수선유지비-열원경상정비", "연예산 합계": 555_000_000.0},
    ])
    budget_df.to_csv(tmp_path / f"budget_{lf.BASE_YEAR}.csv", index=False)

    at = AppTest.from_function(_run_render, default_timeout=30)
    at.run()

    assert not at.exception

    preview = next(d.value for d in at.dataframe if "예산과목" in d.value.columns)
    row = preview[preview["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[str(lf.BASE_YEAR)] == pytest.approx(5.55)  # 억원 단위 - 예산계획 확정값, 표준화값(1억)이 아님


def test_render_without_budget_file_falls_back_to_standardized_amount(isolated):
    """budget_2026.csv가 아예 없으면(예산계획 미등록) 기존처럼 표준화 금액을 그대로 쓴다 -
    render()가 budget_df를 빈 DataFrame으로 넘겨도 compute_all_sites가 깨지지 않는지 확인."""
    tmp_path = isolated

    site_type_map = pd.DataFrame([
        {"사업장": "9001", "표시명": "화성지사", "지사유형": "중대형CHP", "적용시작연도": 1900},
    ])
    site_type_map.to_csv(tmp_path / "site_type_map.csv", index=False)

    classified = pd.DataFrame([
        {"전기일": "2024-03-15", "사업장": "9001", "계정과목": "수선유지비-열원경상정비", "금액": 100_000_000.0},
        {"전기일": "2025-03-15", "사업장": "9001", "계정과목": "수선유지비-열원경상정비", "금액": 100_000_000.0},
    ])
    classified.to_csv(tmp_path / "classified_2025.csv", index=False)

    grade_hist = pd.DataFrame([
        {"사업장": "화성지사", "연도": 2024, "등급": "간이"},
        {"사업장": "화성지사", "연도": 2025, "등급": "간이"},
    ])
    grade_hist.to_csv(tmp_path / "site_year_grade.csv", index=False)

    at = AppTest.from_function(_run_render, default_timeout=30)
    at.run()

    assert not at.exception

    preview = next(d.value for d in at.dataframe if "예산과목" in d.value.columns)
    row = preview[preview["예산과목"] == "수선유지비-열원경상정비"].iloc[0]
    assert row[str(lf.BASE_YEAR)] == pytest.approx(1.0)  # 예산계획 없음 - 표준화 금액(1억) 그대로


def test_render_test_merges_multiple_budget_files_on_save(isolated):
    """Test 모드 예산계획 업로드는 손익예산/자본예산/건설예산처럼 여러 파일을 한 번에 올려도
    모두 합쳐서 하나의 백데이터로 저장돼야 한다."""
    profit_file = _budget_workbook_bytes("700001", "수선유지비-열원경상정비", 100_000.0, "손익예산라인")
    capital_file = _budget_workbook_bytes("700001", "기계장치", 200_000.0, "자본예산라인")
    construction_file = _budget_workbook_bytes("700001", "건설중인자산-재생고온부품", 300_000.0, "건설예산라인")

    at = AppTest.from_function(_run_render_test, default_timeout=30)
    at.run()
    assert not at.exception

    uploader = next(u for u in at.file_uploader if u.key == "ltf_test_budget_upload")
    assert uploader.accept_multiple_files
    uploader.set_value([
        ("손익예산.xlsx", profit_file, _XLSX_MIME),
        ("자본예산.xlsx", capital_file, _XLSX_MIME),
        ("건설예산.xlsx", construction_file, _XLSX_MIME),
    ])
    at.run()
    assert not at.exception

    save_button = next(b for b in at.button if b.key == "ltf_test_budget_apply")
    save_button.click().run()
    assert not at.exception

    saved = lf.load_test_budget_plan()
    assert len(saved) == 3
    assert set(saved["사업명"]) == {"손익예산라인", "자본예산라인", "건설예산라인"}
    total = saved.set_index("사업명")["연예산 합계"]
    assert total["손익예산라인"] == pytest.approx(100_000.0 * 1000)  # 천원 -> 원 환산 확인
    assert total["자본예산라인"] == pytest.approx(200_000.0 * 1000)
    assert total["건설예산라인"] == pytest.approx(300_000.0 * 1000)


def _seed_test_mode_actuals(tmp_path):
    lf.save_test_actuals(pd.DataFrame([
        {"사업장": "화성지사", "연도": 2025, "예산과목": "수선유지비-열원경상정비", "금액": 100_000_000.0},
    ]))
    pd.DataFrame([{"사업장": "화성지사", "연도": 2025, "등급": "간이"}]).to_csv(
        tmp_path / "site_year_grade.csv", index=False)


def test_render_test_shows_status_before_computing(isolated):
    """표준화 실적 백데이터가 있어도, 버튼을 누르기 전에는 계산 결과 섹션이 나오지 않아야 한다."""
    _seed_test_mode_actuals(isolated)

    at = AppTest.from_function(_run_render_test, default_timeout=30)
    at.run()
    assert not at.exception
    assert any("계산 결과 미작성" in i.value for i in at.info)
    assert any(b.key == "ltf_test_build_btn" for b in at.button)
    assert not any(s.value == "계산 결과 미리보기 및 내보내기" for s in at.subheader)


def test_render_test_computes_after_build_button(isolated):
    """만들기 버튼을 누르면 지사별 계산 상태 블록과 함께 미리보기 섹션이 나온다."""
    _seed_test_mode_actuals(isolated)

    at = AppTest.from_function(_run_render_test, default_timeout=30)
    at.run()

    build_button = next(b for b in at.button if b.key == "ltf_test_build_btn")
    build_button.click().run()
    assert not at.exception

    assert any(s.value == "계산 결과 미리보기 및 내보내기" for s in at.subheader)
    compute_status = next(s for s in at.status if s.label == "지사별 계산 완료")
    assert compute_status.state == "complete"


def test_render_test_build_persists_across_fresh_app_runs(isolated):
    """앱을 재시작해도(=새 AppTest 인스턴스) 이전에 만든 계산 결과가 다시 계산할 필요 없이
    '작성됨'으로 남아있어야 하고, site_type_map.csv가 처음 생성되는 것만으로 '데이터가 변경되었다'고
    오판하면 안 된다(백데이터는 전부 파일 기반이라 세션 상태와 무관하게 그대로 남아있음)."""
    _seed_test_mode_actuals(isolated)

    at1 = AppTest.from_function(_run_render_test, default_timeout=30)
    at1.run()
    build_button = next(b for b in at1.button if b.key == "ltf_test_build_btn")
    build_button.click().run()
    assert not at1.exception

    at2 = AppTest.from_function(_run_render_test, default_timeout=30)
    at2.run()
    assert not at2.exception
    assert any("계산 결과 작성됨" in i.value for i in at2.info)
    assert not any("데이터가 변경되었습니다" in i.value for i in at2.info)
    assert any(s.value == "계산 결과 미리보기 및 내보내기" for s in at2.subheader)
    # 진짜로 다시 계산하지 않았어야 한다 - 지사별 계산 상태 블록 자체가 나타나면 안 된다.
    assert not any(s.label == "지사별 계산 완료" for s in at2.status)
    assert any("다시 계산하지 않음" in c.value for c in at2.caption)
