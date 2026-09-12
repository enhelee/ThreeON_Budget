"""중장기 예산(26~35년) 결과를 재무팀 참고 양식(builtin_template_중장기예산.xlsx)에 채워 넣는다.

standardization_template.py와 같은 방식(라벨 스캔 + 병합 셀 안전 처리)을 쓴다.

채우는 범위:
- 지사별 탭(19개): 정비등급 행 + 26~35년 각 예산과목 라인을 compute_site_table() 결과로 덮어쓴다.
- '26년 본사 원가분배': 가져온(또는 화면에서 수정한) hq_master_table을 그대로 되돌려 쓴다.
- '26년 총원가 배분': 지사별 탭의 26년 컬럼(본사배분+표준화+고온부품+투자비 합산 결과)로 채운다.
- '총괄표': 19개 지사별 탭의 연도별 합계를 예산과목 대분류 기준으로 모아 채운다.

'정기점검보수공사 일정'/'고온부품(26)' 원본 시트와 27~35년 개별 시트는 건드리지 않는다
(정책 결정 4: 26년만 상세 시트, 27~35년은 지사별 탭만으로 충분).

build_*_download() 함수들은 업로드 화면에서 "현재 저장된 값이 채워진 원본 양식"을 내려받을 때 쓴다 -
빈 참고 양식이 아니라 지금 이 앱에 저장된 정비등급/고온부품/본사원가분배를 그대로 반영해서
내보내므로, 사용자가 그 파일을 열어 수정한 뒤 그대로 다시 업로드할 수 있다.
"""
import io
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
import standardization as std
import longterm_forecast as lf

BUILTIN_TEMPLATE_PATH = "builtin_template_중장기예산.xlsx"


def _set_cell(ws, row: int, col: int, value):
    """병합 셀이면 풀고 쓴다(standardization_template.py의 동일 헬퍼와 같은 이유)."""
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if (merged_range.min_row <= row <= merged_range.max_row
                    and merged_range.min_col <= col <= merged_range.max_col):
                ws.unmerge_cells(str(merged_range))
                break
        cell = ws.cell(row=row, column=col)
    cell.value = value


def _find_bare_year_row(ws, max_scan_row: int = 6) -> dict:
    """지사별 탭은 연도를 '26'~'35' 같은 두 자리 정수로 쓴다. {열: 연도(2026~2035)} 반환."""
    for r in range(1, max_scan_row + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and 20 <= v <= 45:
                found[c] = 2000 + int(v)
        if len(found) >= 5:
            return found
    return {}


def fill_schedule_sheet(ws, grade_hist: pd.DataFrame):
    """'정기점검보수공사 일정' 시트의 GT열에 현재 저장된 미래(2026년 이후) 정비등급 이력을
    되돌려 쓴다 - longterm_forecast.import_schedule_from_workbook()과 같은 방식으로 지사별
    블록/장비 행을 찾고, 그 중 GT어휘를 쓰는 행(없으면 첫 행)에 쓴다."""
    year_row = None
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and 2020 < v < 2050:
                year_row = r
                break
        if year_row:
            break
    if year_row is None:
        return
    year_by_gt_col = {
        c: int(ws.cell(row=year_row, column=c).value)
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(row=year_row, column=c).value, (int, float))
        and 2020 < ws.cell(row=year_row, column=c).value < 2050
    }
    data_start_row = year_row + 3

    short_to_full = std._site_sheet_name_map()
    short_names_sorted = sorted(short_to_full.keys(), key=len, reverse=True)

    site_rows: dict[str, list[int]] = {}
    current_site = None
    r = data_start_row
    blank_run = 0
    while r <= ws.max_row:
        row_values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v in (None, "") for v in row_values):
            blank_run += 1
            if blank_run >= 3:
                break
            r += 1
            continue
        blank_run = 0
        label = ws.cell(row=r, column=1).value
        if label is not None and str(label).strip():
            raw = str(label).replace(" ", "").strip()
            match = next((s for s in short_names_sorted if raw.startswith(s)), None)
            current_site = short_to_full.get(match) if match else None
        if current_site:
            site_rows.setdefault(current_site, []).append(r)
        r += 1

    for site, row_list in site_rows.items():
        site_grades = grade_hist[(grade_hist["사업장"] == site) & (grade_hist["연도"] >= lf.BASE_YEAR)]
        grade_by_year = dict(zip(site_grades["연도"], site_grades["등급"]))
        if not grade_by_year:
            continue

        gt_rows = []
        for row_idx in row_list:
            existing_gt = {year: ws.cell(row=row_idx, column=c).value for c, year in year_by_gt_col.items()}
            if lf._classify_grade_row(existing_gt) == "GT":
                gt_rows.append(row_idx)
        target_row = gt_rows[0] if gt_rows else row_list[0]

        for col, year in year_by_gt_col.items():
            grade = grade_by_year.get(year)
            if grade:
                _set_cell(ws, target_row, col, grade)


def fill_hot_parts_sheet(ws, hot_parts_df: pd.DataFrame):
    """'고온부품(26)' 시트에 현재 저장된 고온부품 계획을 되돌려 쓴다(원 -> 천원 환산)."""
    header_row, year_by_col = None, {}
    for r in range(1, 8):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            m = re.match(r"^(\d{2})년", str(v).strip())
            if m:
                found[c] = 2000 + int(m.group(1))
        if found:
            header_row, year_by_col = r, found
            break
    if header_row is None:
        return

    short_to_full = std._site_sheet_name_map()
    lookup = {(r["사업장"], r["연도"], r["항목"]): r["금액"] for _, r in hot_parts_df.iterrows()}

    current_site = None
    r = header_row + 1
    blank_run = 0
    while r <= ws.max_row:
        site_label = ws.cell(row=r, column=1).value
        item = ws.cell(row=r, column=5).value
        if site_label:
            current_site = short_to_full.get(str(site_label).strip())
        if item is None or current_site is None:
            blank_run += 1
            if blank_run >= 3:
                break
            r += 1
            continue
        blank_run = 0
        item_s = str(item).strip()
        for col, year in year_by_col.items():
            amount = lookup.get((current_site, year, item_s))
            if amount is not None:
                _set_cell(ws, r, col, round(amount / 1000, 3))
        r += 1


def build_schedule_download(grade_hist: pd.DataFrame, template_path: str = BUILTIN_TEMPLATE_PATH) -> bytes:
    """현재 저장된 미래 정비등급이 채워진 '정기점검보수공사 일정' 양식을 bytes로 반환."""
    wb = load_workbook(template_path)
    sheet_name = next((n for n in wb.sheetnames if "정기점검보수공사" in n), None)
    if sheet_name:
        fill_schedule_sheet(wb[sheet_name], grade_hist)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_hot_parts_download(hot_parts_df: pd.DataFrame, template_path: str = BUILTIN_TEMPLATE_PATH) -> bytes:
    """현재 저장된 고온부품 계획이 채워진 '고온부품(26)' 양식을 bytes로 반환."""
    wb = load_workbook(template_path)
    sheet_name = next((n for n in wb.sheetnames if n.strip().startswith("고온부품")), None)
    if sheet_name:
        fill_hot_parts_sheet(wb[sheet_name], hot_parts_df)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_hq_master_download(hq_master_df: pd.DataFrame, template_path: str = BUILTIN_TEMPLATE_PATH) -> bytes:
    """현재 저장된 본사 원가분배 마스터 표가 채워진 '26년 본사 원가분배' 양식을 bytes로 반환."""
    wb = load_workbook(template_path)
    if "26년 본사 원가분배" in wb.sheetnames:
        fill_hq_master_sheet(wb["26년 본사 원가분배"], hq_master_df)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 지사별 탭 라벨(col1,col2 조합) -> longterm_forecast.SITE_TABLE_LINES 라벨
_SITE_ROW_LABELS = {
    ("수선유지비", "건물/구축물"): "수선유지비-건물/구축물",
    ("", "열원정기점검"): "수선유지비-열원정기점검",
    ("", "열원경상정비"): "수선유지비-열원경상정비",
    ("", "열원정기유지보수"): "수선유지비-열원정기유지보수",
    ("", "열원보완및개선"): "수선유지비-열원보완및개선",
    ("", "비저장품(보수자재)"): "비저장품(보수자재)",
    ("", "소모품(자재공기구)"): "소모품(자재공기구)",
    ("지급수수료", "지급수수료 열원"): "지급수수료-열원점검수수료",
    ("기계장치", "고온부품재생"): "기계장치_고온부품재생",
    ("", "기타기계장치"): "기계장치_기타기계장치",
    ("건가-외주비", "MI 및 A급 정비"): "외주비-열원정기점검",
    ("공구와기구", ""): "공구와기구-열원시설공기구",
    ("건설중인자산", "재생고온부품"): "건설중인자산_재생고온부품",
    ("", "자산화예비품"): "건설중인자산-자산화예비품",
    ("자산", "저장품"): "저장품-열원(보수)",
    ("투자비", "열원공사비 등"): "투자비",
}


def fill_site_sheet(ws, site_table: pd.DataFrame, grade_hist: pd.DataFrame, site: str):
    year_by_col = _find_bare_year_row(ws)
    if not year_by_col:
        return

    values_by_label = {row["예산과목"]: row for _, row in site_table.iterrows()}

    # 정비등급 행
    grade_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == "정비등급":
            grade_row = r
            break
    if grade_row:
        site_grades = grade_hist[grade_hist["사업장"] == site]
        grade_by_year = dict(zip(site_grades["연도"], site_grades["등급"]))
        latest = std.get_current_grade(site, grade_hist)
        for col, year in year_by_col.items():
            g = grade_by_year.get(year, latest)
            if g:
                _set_cell(ws, grade_row, col, g)

    # 라인 아이템 행 (col1 또는 col2 라벨로 매칭, 대분류는 이전 non-blank 값을 이어받음)
    major = ""
    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        if c1:
            major = str(c1).strip()
        label2 = str(c2).strip() if c2 else ""
        key = (major, label2) if (major, label2) in _SITE_ROW_LABELS else ("", label2)
        canonical = _SITE_ROW_LABELS.get(key)
        if canonical is None and label2 == "" and major:
            canonical = _SITE_ROW_LABELS.get((major, ""))
        if canonical is None:
            continue
        row_vals = values_by_label.get(canonical)
        if row_vals is None:
            continue
        for col, year in year_by_col.items():
            amount = row_vals.get(year)
            if amount is not None:
                _set_cell(ws, r, col, round(float(amount) / 1_000_000, 3))  # 시트 단위: 백만원


def fill_hq_master_sheet(ws, hq_master_df: pd.DataFrame):
    """'26년 본사 원가분배' 시트에 (가져왔거나 수정된) 마스터 표를 그대로 되돌려 쓴다."""
    if hq_master_df.empty:
        return
    header_row, budget_col = None, None
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(row=r, column=c).value or "").strip() == "26년 예산":
                header_row, budget_col = r, c
                break
        if header_row:
            break
    if header_row is None:
        return

    short_to_full = std._site_sheet_name_map()
    site_col = {}
    for c in range(budget_col + 1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        full = short_to_full.get(str(v).strip())
        if full:
            site_col[c] = full

    major = minor = None
    r = header_row + 1
    idx = 0
    rows = hq_master_df.reset_index(drop=True)
    while r <= ws.max_row and idx < len(rows):
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        c3 = ws.cell(row=r, column=3).value
        label1 = str(c1).strip() if c1 else ""
        if label1 in ("합계", "경상정비 원가배부기준"):
            break
        if c1:
            major = label1
        if c2:
            minor = str(c2).strip()
        detail = str(c3).strip() if c3 else ""
        if detail == "합계":
            r += 1
            continue
        src_row = rows.iloc[idx]
        if src_row["대분류"] == (major or "") and src_row["중분류"] == (minor or "") and src_row["세부내역"] == detail:
            for col, site in site_col.items():
                if site in rows.columns:
                    _set_cell(ws, r, col, float(src_row[site]))
            idx += 1
        r += 1


def fill_total_cost_sheet(ws, site_tables: dict):
    """'26년 총원가 배분' 시트를 지사별 탭의 26년 컬럼(본사+표준화+고온부품+투자비 합산 결과)으로 채운다."""
    header_row, budget_col = None, None
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(row=r, column=c).value or "").strip() == "26년 예산":
                header_row, budget_col = r, c
                break
        if header_row:
            break
    if header_row is None:
        return

    short_to_full = std._site_sheet_name_map()
    site_col = {}
    for c in range(budget_col + 1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        full = short_to_full.get(str(v).strip())
        if full:
            site_col[c] = full

    reverse_label = {v: k for k, v in lf.HQ_CATEGORY_TO_ACCOUNT.items()}
    major = None
    r = header_row + 1
    while r <= ws.max_row:
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        label1 = str(c1).strip() if c1 else ""
        if label1 == "합계":
            break
        if c1:
            major = label1
        label2 = str(c2).strip() if c2 else ""
        canonical = lf.HQ_CATEGORY_TO_ACCOUNT.get(label2) or lf.HQ_CATEGORY_TO_ACCOUNT.get(label1)
        if canonical:
            for col, site in site_col.items():
                table = site_tables.get(site)
                if table is None:
                    continue
                match = table[table["예산과목"] == canonical]
                if not match.empty and 2026 in match.columns:
                    _set_cell(ws, r, col, round(float(match.iloc[0][2026]) / 1_000_000, 3))
        r += 1


# '총괄표' 라벨(col1,col2 조합) -> longterm_forecast.SITE_TABLE_LINES 라벨
# 대분류(col1)는 이전 non-blank 값을 이어받으므로, "수선유지비" 그룹 안에서만 쓰이는
# "건물/구축물"은 반드시 (major,minor) 조합으로 매칭해야 자본계정의 동명 행(14~15행)과 안 섞인다.
_SUMMARY_ROW_LABELS = {
    ("수선유지비", "건물/구축물"): "수선유지비-건물/구축물",
    ("", "열원정기점검"): "수선유지비-열원정기점검",
    ("", "열원경상정비"): "수선유지비-열원경상정비",
    ("", "열원정기유지보수"): "수선유지비-열원정기유지보수",
    ("", "열원보완및개선"): "수선유지비-열원보완및개선",
    ("", "비저장품(보수자재)"): "비저장품(보수자재)",
    ("", "소모품(자재공기구)"): "소모품(자재공기구)",
    ("지급수수료", "지급수수료 열원"): "지급수수료-열원점검수수료",
    ("기계장치", "고온부품재생+LTSA"): "기계장치_고온부품재생",
    ("", "기타기계장치"): "기계장치_기타기계장치",
    ("공구와기구", ""): "공구와기구-열원시설공기구",
    ("예비품", "재생고온부품"): "건설중인자산_재생고온부품",
    ("", "자산화예비품"): "건설중인자산-자산화예비품",
    ("", "저장품"): "저장품-열원(보수)",
    ("투자비", "열원공사비 등"): "투자비",
}


def fill_summary_sheet(ws, site_tables: dict):
    """'총괄표' 시트 - 19개 지사 합계를 연도별로 채운다."""
    year_row = None
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip().endswith("년") and v.strip()[:-1].isdigit():
                year_row = r
                break
        if year_row:
            break
    if year_row is None:
        return
    year_by_col = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=year_row, column=c).value
        if isinstance(v, str) and v.strip().endswith("년") and v.strip()[:-1].isdigit():
            year_by_col[c] = 2000 + int(v.strip()[:-1])

    # 라인 아이템 행 (col1/col2 라벨 조합으로 매칭, 대분류는 이전 non-blank 값을 이어받음)
    major = ""
    r = year_row + 1
    while r <= ws.max_row:
        c1 = ws.cell(row=r, column=1).value
        c2 = ws.cell(row=r, column=2).value
        if c1:
            major = str(c1).strip()
        minor = str(c2).strip() if c2 else ""
        if minor in ("계", "합계"):
            r += 1
            continue

        key = (major, minor) if (major, minor) in _SUMMARY_ROW_LABELS else ("", minor)
        canonical = _SUMMARY_ROW_LABELS.get(key)
        if canonical is None and minor == "" and major:
            canonical = _SUMMARY_ROW_LABELS.get((major, ""))
        if canonical is None and major == "건가-외주비":
            canonical = "외주비-열원정기점검"  # 지사유형별로 세부 라벨 문구가 달라 대분류만으로 매칭

        if canonical:
            for col, year in year_by_col.items():
                total = 0.0
                for table in site_tables.values():
                    match = table[table["예산과목"] == canonical]
                    if not match.empty and year in match.columns:
                        total += float(match.iloc[0][year])
                _set_cell(ws, r, col, round(total / 1_000_000, 3))
        r += 1


def fill_longterm_workbook(site_tables: dict, grade_hist: pd.DataFrame, hq_master_df: pd.DataFrame,
                            output_path: str, template_path: str = BUILTIN_TEMPLATE_PATH) -> str:
    wb = load_workbook(template_path)
    short_to_full = std._site_sheet_name_map()

    for sheet_name in wb.sheetnames:
        site = short_to_full.get(sheet_name.strip())
        if site is None or site not in site_tables:
            continue
        fill_site_sheet(wb[sheet_name], site_tables[site], grade_hist, site)

    if "26년 본사 원가분배" in wb.sheetnames:
        fill_hq_master_sheet(wb["26년 본사 원가분배"], hq_master_df)
    if "26년 총원가 배분" in wb.sheetnames:
        fill_total_cost_sheet(wb["26년 총원가 배분"], site_tables)
    if "총괄표" in wb.sheetnames:
        fill_summary_sheet(wb["총괄표"], site_tables)

    wb.save(output_path)
    return output_path
