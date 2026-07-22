# -*- coding: utf-8 -*-
"""사업별 예산 입력 양식(빈 템플릿) 생성. 손익·자본 구분 없이 전체 사업 입력용."""
import openpyxl, os

HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "..", "templates", "사업별예산_템플릿.xlsx")

HEADER = [
    "주관부서명", "예산귀속 부서코드", "예산귀속 부서명(처.지사)", "예산귀속 부서명(부)",
    "속성", "예산코드", "예산과목", "사업명", "산출내역", "연예산 합계",
    "1월", "2월", "3월", "4월", "5월", "6월", "7월", "8월", "9월", "10월", "11월", "12월",
]

def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    ws.cell(1, 1, "※ 사업별 예산 입력 양식 — 손익/자본 구분 없이 전체 사업을 입력하세요.")
    for i, h in enumerate(HEADER, 1):
        ws.cell(3, i, h)
    out_path = os.path.abspath(OUT)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print("saved", out_path)

if __name__ == "__main__":
    main()
