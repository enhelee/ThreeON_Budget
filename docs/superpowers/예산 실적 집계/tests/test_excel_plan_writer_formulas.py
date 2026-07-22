import openpyxl
from openpyxl.utils import get_column_letter
from budget import excel_plan_writer


def _layout_inputs():
    dept_columns = [
        {"그룹": "본사", "지사": ["플랜트기술처", "안전처"]},
        {"그룹": "중대형CHP", "지사": ["화성지사"]},
    ]
    item_rows = [
        {"대분류": "수선유지비", "과목들": [
            {"과목": "수선유지비-건물/구축물", "심의대상": True},
            {"과목": "수선유지비-열원경상정비", "심의대상": False},
        ]},
        {"대분류": "지급수수료", "과목들": [
            {"과목": "지급수수료-열원점검수수료", "심의대상": False},
        ]},
    ]
    return dept_columns, item_rows


def test_build_summary_layout_columns_and_rows():
    dept_columns, item_rows = _layout_inputs()
    wb = openpyxl.Workbook()
    ws = wb.active
    layout = excel_plan_writer.build_summary_layout(ws, dept_columns, item_rows, "손익", "2026")

    assert layout["groups"][0]["dept_cols"] == {"플랜트기술처": 4, "안전처": 5}
    assert layout["groups"][0]["subtotal_col"] == 6
    assert layout["groups"][1]["dept_cols"] == {"화성지사": 7}
    assert layout["groups"][1]["subtotal_col"] == 8
    assert layout["total_col"] == 9

    cat0 = layout["item_rows"][0]
    assert cat0["item_start"] == 5 and cat0["item_end"] == 7
    assert cat0["subtotal_row"] == 8
    cat1 = layout["item_rows"][1]
    assert cat1["item_start"] == 9 and cat1["item_end"] == 9
    assert cat1["subtotal_row"] == 10
    assert layout["grand_total_row"] == 11

    assert ws.cell(4, 4).value == "플랜트기술처"
    assert ws.cell(4, 6).value == "소계"
    assert ws.cell(5, 2).value == "수선유지비-건물/구축물"
    assert ws.cell(5, 3).value == "5천만원 이상(심의대상)"
    assert ws.cell(6, 3).value == "5천만원 미만"
    assert ws.cell(8, 1).value == "수선유지비 소계"


def test_write_summary_formulas_references_datasheet():
    dept_columns, item_rows = _layout_inputs()
    wb = openpyxl.Workbook()
    ws = wb.active
    layout = excel_plan_writer.build_summary_layout(ws, dept_columns, item_rows, "손익", "2026")
    excel_plan_writer.write_summary_formulas(ws, layout)

    f_over = ws.cell(5, 4).value  # 이상 행, 플랜트기술처 열
    assert f_over.startswith("=SUMIFS(")
    assert "'양식1(월별)'!$L:$L,\"O\"" in f_over
    assert "$B5" in f_over and "D4" in f_over

    f_under = ws.cell(6, 4).value  # 미만 행
    assert "\"<>O\"" in f_under

    f_plain = ws.cell(7, 4).value  # 심의 없는 항목
    assert "$L:$L" not in f_plain

    f_subtotal = ws.cell(5, 6).value  # 본사 그룹 소계(이상 행)
    assert f_subtotal == "=SUM(D5:E5)"

    f_cat_subtotal = ws.cell(8, 4).value  # 수선유지비 소계행, D열
    assert f_cat_subtotal == "=SUM(D5:D7)"

    f_grand_total = ws.cell(11, 4).value
    assert f_grand_total == "=D8+D10"
