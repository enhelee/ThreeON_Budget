# -*- coding: utf-8 -*-
"""DB 기반 파이프라인(보안 저장소) 회귀 테스트 — 흡수·실행·오버라이드."""
import os

import openpyxl
import pytest

from budget import db as dbm
from budget import pipeline_db

PLAN_HEADER = ["주관부서명", "부서코드", "처지사", "부서부", "속성",
               "예산코드", "예산과목", "사업명", "산출내역", "연예산"]

PLAN_ROWS = [
    ["기술부", None, "화성지사", "기술부", "제조", None,
     "수선유지비-열원정기점검", "화성 정기점검 보수공사", None, 10000],
    ["기술부", None, "화성지사", "기술부", "제조", None,
     "수선유지비-열원정기점검", "화성 옥외배관 도색공사", None, 5000],
]

ERP_ROWS = [
    # 화성: 계획 있음 → 계획집행
    [None, "60909002", "수선유지비-열원정기점검", "2023.001", "2023-03-02", "D1",
     8_000_000, "화성 정기점검 보수공사 1회", "화성지사", "(CHP)화성지사 기술부"],
    # 대구: 계획 없음 → 신규
    [None, "60909002", "수선유지비-열원정기점검", "2023.002", "2023-04-01", "D2",
     3_000_000, "전혀 관련 없는 임시 긴급공사", "대구지사", "(CHP)대구지사 기술부"],
    # 화성: 계획 있음 그룹의 저유사 전표 → group 정책상 임의귀속(오버라이드 대상)
    [None, "60909002", "수선유지비-열원정기점검", "2023.003", "2023-05-01", "D3",
     2_500_000, "다른 무언가 공사", "화성지사", "(CHP)화성지사 기술부"],
]
TOTAL = 13500


@pytest.fixture()
def env(tmp_path):
    plan = tmp_path / "계획.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "양식1(월별)"
    for i, h in enumerate(PLAN_HEADER, 1):
        ws.cell(3, i, h)
    for r, row in enumerate(PLAN_ROWS, 4):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(plan)
    wb.close()

    erp = tmp_path / "zrfm2.XLSX"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 17):
        ws.cell(1, i, f"h{i}")
    for r, row in enumerate(ERP_ROWS, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    wb.save(erp)
    wb.close()

    conn = dbm.connect(str(tmp_path / "budget.db"))
    dbm.ingest_plan(conn, str(plan), "2023")
    dbm.ingest_erp(conn, str(erp), "2023")
    return conn, str(tmp_path / "config"), str(tmp_path / "out")


def test_ingest_and_run_from_db(env):
    conn, cfg, out = env
    ds = dbm.list_datasets(conn)
    assert {d["kind"] for d in ds} == {"plan", "erp"}

    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["총 실적(천원)"] == TOTAL     # 계획집행 + 신규
    assert os.path.exists(res["output_path"])
    assert os.path.exists(res["zrfm2_v1_path"])      # DB 모드에서도 V1 생성
    # 실행 이력 + 매칭 상세가 DB에 남는다
    run = dbm.latest_run(conn, "2023", "손익")
    assert run and run["summary"]["총 실적(천원)"] == TOTAL
    detail = dbm.load_match_detail(conn, run["id"])
    assert len(detail) == 3
    assert set(detail["구분"]) == {"계획집행", "신규"}   # 대구지사 전표만 신규


def test_override_reassigns_voucher(env):
    """group 정책상 임의귀속된 전표(D3)를 사용자가 올바른 계획 사업으로 재배정."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    assert d3["구분"] == "계획집행" and d3["매칭확신도"] < 0.8   # 저유사 임의귀속 상태

    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]),
                     "화성 옥외배관 도색공사", memo="테스트 재배정")
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["총 실적(천원)"] == TOTAL     # 총액 보존
    assert res["요약"]["수동지정 전표수"] == 1
    detail2 = dbm.load_match_detail(conn, res["run_id"])
    moved = detail2[detail2["erp_row_id"] == int(d3["erp_row_id"])].iloc[0]
    assert moved["구분"] == "계획집행"
    assert moved["매칭사업명"] == "화성 옥외배관 도색공사"
    assert moved["매칭확신도"] == 1.0


def test_manual_biz_added_as_plan_row(env):
    """수동 추가 사업이 계획행으로 편입되고(미시행), 재배정 대상이 된다."""
    conn, cfg, out = env
    bid = dbm.add_manual_biz(conn, "2023", "손익", "수선유지비-열원정기점검",
                             "화성지사", "수동 추가 검증 사업", attr="제조",
                             plan_amt=7777)
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    row = biz[biz["사업명"] == "수동 추가 검증 사업"]
    assert len(row) == 1
    assert row.iloc[0]["연예산"] == 7777
    assert row.iloc[0]["src_row"] == 1_000_000 + bid
    assert res["요약"]["총 실적(천원)"] == TOTAL       # 총액 불변(실적 없음)

    # 전표를 수동 추가 사업으로 재배정 → 계획집행으로 귀속
    detail = dbm.load_match_detail(conn, res["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]), "수동 추가 검증 사업")
    res2 = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz2 = dbm.load_biz_lines(conn, res2["run_id"])
    row2 = biz2[biz2["사업명"] == "수동 추가 검증 사업"].iloc[0]
    assert row2["구분"] == "계획집행" and row2["실적금액"] == 2500
    assert res2["요약"]["총 실적(천원)"] == TOTAL

    # 삭제하면 다음 분석부터 제외
    dbm.delete_manual_biz(conn, bid)
    for ov in dbm.list_overrides(conn, "2023", "손익"):
        dbm.delete_override(conn, ov["id"])
    res3 = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz3 = dbm.load_biz_lines(conn, res3["run_id"])
    assert not (biz3["사업명"] == "수동 추가 검증 사업").any()


def test_cross_dept_override_moves_voucher(env):
    """타지사 이동: target_dept를 지정하면 전표가 그 지사의 사업으로 귀속된다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]     # 화성지사 전표
    # 화성지사 전표를 청주지사의 신규 사업으로 이동
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]),
                     "청주 이관 검증 사업", target_dept="청주지사")
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["총 실적(천원)"] == TOTAL        # 총액 보존
    biz = dbm.load_biz_lines(conn, res["run_id"])
    moved = biz[biz["사업명"] == "청주 이관 검증 사업"]
    assert len(moved) == 1
    assert moved.iloc[0]["처지사"] == "청주지사"
    assert moved.iloc[0]["실적금액"] == 2500
    # 상세(match_line)에도 이동된 지사로 기록
    det2 = dbm.load_match_detail(conn, res["run_id"])
    mv = det2[det2["erp_row_id"] == int(d3["erp_row_id"])].iloc[0]
    assert mv["처지사정규"] == "청주지사"


def test_year_lock_and_attr_master(env):
    """연도 잠금 상태 저장/해제 + 마스터 속성맵으로 잘못된 속성이 교정된다."""
    conn, cfg, out = env
    # 잠금 토글
    assert not dbm.is_locked(conn, "2023")
    dbm.lock_year(conn, "2023")
    assert dbm.is_locked(conn, "2023")
    assert "2023" in dbm.locked_years(conn)
    dbm.unlock_year(conn, "2023")
    assert not dbm.is_locked(conn, "2023")

    # 마스터 속성 교정: 속성 칸에 과목명이 들어간 계획(25년 자본 원본 케이스)
    conn.execute("INSERT OR REPLACE INTO item_master VALUES(?,?,?,?,?,?)",
                 ("60909002", "수선유지비-열원정기점검", None, None, "제조", None))
    conn.commit()
    # 계획의 속성을 고의로 오염시킨 뒤 재분석
    ds = conn.execute("SELECT id FROM dataset WHERE kind='plan'").fetchone()[0]
    conn.execute("UPDATE plan_row SET 속성='수선유지비-열원정기점검' WHERE dataset_id=?", (ds,))
    conn.commit()
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    attrs = {str(a) for a in biz[biz["구분"] != "신규"]["속성"].dropna().unique()}
    assert attrs <= {"일반", "제조", "건가", "자산"}     # 4종으로 교정됨
    assert "제조" in attrs


def test_override_auto_learns(env):
    """재배정 저장 시 learn_from_override로 즉시 수동확정 학습이 축적된다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]), "화성 옥외배관 도색공사")
    n = dbm.learn_from_override(conn, "2023", "손익", [int(d3["erp_row_id"])],
                                "화성 옥외배관 도색공사")
    assert n == 1
    stats = dbm.learned_stats(conn)
    assert stats.get("수동확정", 0) == 1
    learned = dbm.load_learned(conn)
    key = ("수선유지비-열원정기점검", dbm.norm_text(d3["전표텍스트"]))
    assert learned.get(key) == "화성 옥외배관 도색공사"


def test_learn_from_confirmed_results(env):
    """사람이 재배정으로 확정한 결과를 학습하면, 재배정을 지워도
    다음 분석에서 같은 텍스트의 전표가 학습 맵으로 자동 확정된다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]      # 임의귀속 상태
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]),
                     "화성 옥외배관 도색공사")
    pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")

    # 확정 결과 학습 → 재배정 삭제(수동 지시 없이 학습만으로 재현되는지)
    counts = dbm.learn_from_year(conn, "2023")
    assert counts["수동확정"] >= 1
    for ov in dbm.list_overrides(conn, "2023", "손익"):
        dbm.delete_override(conn, ov["id"])

    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["학습확정 전표수"] >= 1
    detail2 = dbm.load_match_detail(conn, res["run_id"])
    moved = detail2[detail2["erp_row_id"] == int(d3["erp_row_id"])].iloc[0]
    assert moved["매칭사업명"] == "화성 옥외배관 도색공사"
    assert moved["매칭확신도"] == 0.99                    # 학습 확정 표시
    assert res["요약"]["총 실적(천원)"] == TOTAL          # 총액 보존


def test_biz_edit_persists_and_patches(env):
    """관리자 사업 수정(속성·사업명)이 즉시 패치되고 재분석에도 유지된다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    # '화성 정기점검 보수공사' 사업의 속성·사업명 수정 (빈칸 입력 포함)
    fields = {"사업명": "화성 정기점검 보수공사(개정)", "속성": "", "부서부": "정비1부"}
    dbm.add_biz_edit(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                     "화성 정기점검 보수공사", fields)
    n = dbm.patch_latest_run_biz(conn, "2023", "손익", "수선유지비-열원정기점검",
                                 "화성지사", "화성 정기점검 보수공사", fields)
    assert n == 1
    # 즉시 패치 확인 (재분석 없이)
    biz = dbm.load_biz_lines(conn, base["run_id"])
    row = biz[biz["사업명"] == "화성 정기점검 보수공사(개정)"]
    assert len(row) == 1
    import pandas as pd
    assert pd.isna(row.iloc[0]["속성"])                   # 빈칸 입력 → 지움
    assert row.iloc[0]["부서부"] == "정비1부"
    # match_line 라벨도 새 이름으로 이동(전표 트리 정합)
    det = dbm.load_match_detail(conn, base["run_id"])
    assert (det["매칭사업명"] == "화성 정기점검 보수공사(개정)").any()

    # 재분석해도 수정이 유지(매칭 전 계획행에 적용)
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz2 = dbm.load_biz_lines(conn, res["run_id"])
    assert (biz2["사업명"] == "화성 정기점검 보수공사(개정)").any()
    assert not (biz2["사업명"] == "화성 정기점검 보수공사").any()
    assert res["요약"]["총 실적(천원)"] == TOTAL


def test_patch_latest_run_biz_moves_dept(env):
    """✎ 수정으로 처지사를 바꾸면 즉시 패치가 biz_line·match_line 지사를 함께 옮긴다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    n = dbm.patch_latest_run_biz(conn, "2023", "손익", "수선유지비-열원정기점검",
                                 "화성지사", "화성 정기점검 보수공사",
                                 {"처지사": "청주지사"})
    assert n == 1
    biz = dbm.load_biz_lines(conn, base["run_id"])
    row = biz[biz["사업명"] == "화성 정기점검 보수공사"].iloc[0]
    assert row["처지사"] == "청주지사"
    det = dbm.load_match_detail(conn, base["run_id"])
    moved = det[det["매칭사업명"] == "화성 정기점검 보수공사"]
    assert len(moved) and set(moved["처지사정규"]) == {"청주지사"}


def test_biz_moved_whole_to_other_dept(env):
    """사업 통째 이동: 계획행(biz_edit.처지사) + 전표(override.target_dept)가 함께 옮겨진다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d1 = detail[detail["전표번호"] == "D1"].iloc[0]      # 화성 정기점검 보수공사 귀속 전표

    dbm.add_biz_edit(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                     "화성 정기점검 보수공사", {"처지사": "청주지사"})
    dbm.add_override(conn, "2023", "손익", int(d1["erp_row_id"]),
                     "화성 정기점검 보수공사", target_dept="청주지사")

    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["총 실적(천원)"] == TOTAL          # 총액 보존
    biz = dbm.load_biz_lines(conn, res["run_id"])
    moved = biz[biz["사업명"] == "화성 정기점검 보수공사"]
    assert len(moved) == 1
    m = moved.iloc[0]
    assert m["처지사"] == "청주지사"                      # 지사 이동
    assert m["연예산"] == 10000                          # 연예산(A)도 함께 이동
    assert m["실적금액"] == 8000                         # 실적(B)도 함께 이동
    assert m["구분"] == "계획집행"                        # 계획집행 관계 유지
    # 옛 지사에는 남지 않는다(다른 사업만 남음)
    assert not ((biz["처지사"] == "화성지사")
                & (biz["사업명"] == "화성 정기점검 보수공사")).any()


def test_biz_delete_excludes_plan_row_and_restores(env):
    """사업 삭제 = 계획행을 분석에서 제외. 원본은 남아 이력 삭제로 복구된다."""
    conn, cfg, out = env
    # 실적 없는 계획행(도색공사)에 붙은 임의귀속 전표를 다른 사업으로 치워 미시행로 만든다
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]),
                     "화성 정기점검 보수공사")
    dbm.add_biz_delete(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                       "화성 옥외배관 도색공사")
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    assert not (biz["사업명"] == "화성 옥외배관 도색공사").any()   # 삭제됨
    assert res["요약"]["총 실적(천원)"] == TOTAL                   # 실적 총액 보존
    # 검토리포트에 삭제분을 보고(조용한 누락 금지)
    reasons = [r[0] for r in res["계획행제외"]]
    assert any("관리자 사업 삭제" in r for r in reasons)

    # 이력을 지우면 되살아난다
    dl = dbm.list_biz_deletes(conn, "2023", "손익")
    assert len(dl) == 1
    dbm.delete_biz_delete(conn, dl[0]["id"])
    res2 = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz2 = dbm.load_biz_lines(conn, res2["run_id"])
    assert (biz2["사업명"] == "화성 옥외배관 도색공사").any()


def test_biz_delete_with_vouchers_moved_keeps_total(env):
    """전표가 붙은 사업 삭제 = 전표를 다른 사업(또는 미지정)으로 옮긴 뒤 제외 → 총액 보존."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    new_row = detail[detail["구분"] == "신규"].iloc[0]       # 대구지사 전표
    unassigned = "대구지사 수선유지비-열원정기점검 미지정"
    dbm.add_override(conn, "2023", "손익", int(new_row["erp_row_id"]), unassigned)
    dbm.add_biz_delete(conn, "2023", "손익", "수선유지비-열원정기점검", "대구지사",
                       new_row["매칭사업명"].replace("[신규] ", ""))
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    assert res["요약"]["총 실적(천원)"] == TOTAL              # 금액은 사라지지 않는다
    biz = dbm.load_biz_lines(conn, res["run_id"])
    moved = biz[biz["사업명"] == unassigned]
    assert len(moved) == 1 and moved.iloc[0]["실적금액"] == 3000


def test_rename_syncs_override_no_duplicate(env):
    """사업명을 바꾸면 그 이름을 가리키던 재배정도 함께 옮겨져 동명 중복 행이 안 생긴다."""
    conn, cfg, out = env
    bid = dbm.add_manual_biz(conn, "2023", "손익", "수선유지비-열원정기점검",
                             "화성지사", "옛이름 사업", attr="제조", plan_amt=0)
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, res["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]), "옛이름 사업")
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    assert len(biz[biz["사업명"] == "옛이름 사업"]) == 1        # 한 행(계획집행)

    # 이름 변경 + 참조 동기화
    dbm.add_biz_edit(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                     "옛이름 사업", {"사업명": "새이름 사업"})
    moved = dbm.rename_biz_references(conn, "2023", "손익", "수선유지비-열원정기점검",
                                      "화성지사", "옛이름 사업", "새이름 사업")
    assert moved["override"] == 1
    res2 = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz2 = dbm.load_biz_lines(conn, res2["run_id"])
    assert not (biz2["사업명"] == "옛이름 사업").any()          # 옛 이름 잔존 없음
    same = biz2[biz2["사업명"] == "새이름 사업"]
    assert len(same) == 1                                      # 중복 아님 — 한 행
    assert same.iloc[0]["구분"] == "계획집행" and same.iloc[0]["실적금액"] == 2500
    assert res2["요약"]["총 실적(천원)"] == TOTAL
    dbm.delete_manual_biz(conn, bid)


def test_biz_edit_changes_plan_amount(env):
    """연예산(A) 수정이 계획행에 반영되고(원본 plan_row는 보존) 이력 삭제로 복구된다."""
    conn, cfg, out = env
    dbm.add_biz_edit(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                     "화성 정기점검 보수공사", {"연예산": 12345.0})
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    row = biz[biz["사업명"] == "화성 정기점검 보수공사"].iloc[0]
    assert row["연예산"] == 12345                     # 10000 → 12345
    assert res["요약"]["총 실적(천원)"] == TOTAL       # 실적은 그대로

    ed = dbm.list_biz_edits(conn, "2023", "손익")[0]
    dbm.delete_biz_edit(conn, ed["id"])
    res2 = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz2 = dbm.load_biz_lines(conn, res2["run_id"])
    assert biz2[biz2["사업명"] == "화성 정기점검 보수공사"].iloc[0]["연예산"] == 10000


def test_biz_edit_changes_item(env):
    """예산과목 수정: 계획행은 새 과목으로 옮겨가고, 전표는 계정코드 과목에 남아 총액 보존."""
    conn, cfg, out = env
    dbm.add_biz_edit(conn, "2023", "손익", "수선유지비-열원정기점검", "화성지사",
                     "화성 옥외배관 도색공사", {"예산과목": "수선유지비-건물/구축물"})
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    biz = dbm.load_biz_lines(conn, res["run_id"])
    row = biz[biz["사업명"] == "화성 옥외배관 도색공사"]
    assert len(row) == 1
    assert row.iloc[0]["예산과목"] == "수선유지비-건물/구축물"      # 과목 이동
    assert row.iloc[0]["연예산"] == 5000                          # 연예산도 함께 이동
    assert res["요약"]["총 실적(천원)"] == TOTAL                   # 전표 총액 보존


def test_make_files_false_skips_outputs_same_numbers(env):
    """화면 갱신용 빠른 분석: 파일을 만들지 않지만 숫자는 동일해야 한다."""
    conn, cfg, out = env
    fast = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익", make_files=False)
    assert fast["files_written"] is False
    assert fast["output_path"] is None and fast["zrfm2_v1_path"] is None
    assert not os.path.exists(os.path.join(out, "2023년 손익예산_실적.xlsx"))
    assert not os.path.exists(os.path.join(out, "zrfm2_2023_V1(손익).xlsx"))
    assert not os.path.exists(os.path.join(out, "matched_2023_손익.csv"))
    # DB 기록은 정상(화면 조회는 run/biz_line/match_line만 쓴다)
    assert dbm.latest_run(conn, "2023", "손익")["id"] == fast["run_id"]
    assert len(dbm.load_biz_lines(conn, fast["run_id"])) > 0
    assert len(dbm.load_match_detail(conn, fast["run_id"])) == 3

    full = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익", make_files=True)
    assert full["files_written"] is True
    for k in ("output_path", "zrfm2_v1_path", "matched_csv_path"):
        assert os.path.exists(full[k])
    assert full["요약"] == fast["요약"]          # 파일 생성 여부가 숫자를 바꾸지 않는다


def test_export_regeneration_does_not_add_run(env):
    """내보내기용 파일 재생성은 실행 이력을 남기지 않는다 —
    남기면 방금 만든 파일이 'run보다 오래됨'이 되어 매번 다시 만들게 된다."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익", make_files=False)
    runs_before = conn.execute("SELECT COUNT(*) FROM run").fetchone()[0]
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익",
                                    make_files=True, record_run=False)
    assert conn.execute("SELECT COUNT(*) FROM run").fetchone()[0] == runs_before
    assert res["run_id"] == base["run_id"]          # 기존 실행을 그대로 가리킨다
    assert os.path.exists(res["output_path"])       # 파일은 만들어졌다
    assert dbm.latest_run(conn, "2023", "손익")["id"] == base["run_id"]


def test_override_does_not_learn_on_save(env):
    """재배정 저장은 학습하지 않는다(학습은 설정 탭 버튼으로 분리) — 학습 대기로 집계."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    d3 = detail[detail["전표번호"] == "D3"].iloc[0]
    before = dbm.learned_stats(conn).get("총계", 0)
    dbm.add_override(conn, "2023", "손익", int(d3["erp_row_id"]),
                     "화성 옥외배관 도색공사")
    assert dbm.learned_stats(conn).get("총계", 0) == before      # 저장만으로는 학습 없음
    assert dbm.count_learn_pending(conn, "2023") >= 1            # 학습 대기로 잡힌다

    # 분석 반영 후 버튼(learn_from_year)을 누르면 학습된다
    pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익", make_files=False)
    dbm.learn_from_year(conn, "2023")
    assert dbm.learned_stats(conn).get("총계", 0) > before
    assert dbm.count_learn_pending(conn, "2023") == 0


def test_override_forced_new_name(env):
    """신규 전표(D2)에 사용자가 직접 사업명 부여 → 그 이름의 신규 사업으로 강제."""
    conn, cfg, out = env
    base = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    detail = dbm.load_match_detail(conn, base["run_id"])
    row = detail[detail["구분"] == "신규"].iloc[0]
    dbm.add_override(conn, "2023", "손익", int(row["erp_row_id"]), "긴급 배관보수(수동)")
    res = pipeline_db.run_actual_db(conn, cfg, out, "2023", "손익")
    wb = openpyxl.load_workbook(res["output_path"])
    ws = wb["Sheet1"]
    names = {ws.cell(r, 7).value for r in range(3, ws.max_row + 1)}
    wb.close()
    assert "긴급 배관보수(수동)" in names
    assert res["요약"]["총 실적(천원)"] == TOTAL
