import json
import os

import openpyxl
import pytest

from budget import pipeline_plan, pipeline_actual
from budget.excel_actual_writer import (
    ACTUAL_HEADER, ACTUAL_SHEET, DATA_START, HEADER_ROW,
)

ROOT = r"C:\Users\User\Desktop\중장기 예산 소요 전망"
REAL_PLAN = os.path.join(ROOT, "양식1_월별_템플릿(25년 계획분)_R1.xlsx")
ZRFM2 = os.path.join(ROOT, "zrfm2.XLSX")
MASTER = os.path.join(ROOT, "(20251202) 손익예산26_최종본.xlsx")

pytestmark = pytest.mark.skipif(
    not (os.path.exists(REAL_PLAN) and os.path.exists(ZRFM2)),
    reason="실제 데이터 파일 없음",
)


def _make_plan(tmp_path, budget):
    cfg = str(tmp_path / "config")
    out = str(tmp_path / "output")
    res = pipeline_plan.run_plan(REAL_PLAN, cfg, out, "2025", budget)
    return res["output_path"], cfg, out


# 행별 천원 반올림 누적 오차 허용치. 요약·양식1·종합표는 모두 '행별 반올림 후 합'
# 이므로 서로 정확히 일치하고, ERP 원천 합계와는 이 범위 안에서만 차이 난다.
ROUND_TOL = 20


def test_actual_total_conserves_full_erp(tmp_path):
    """종합표 총 실적 + 미반영 = ERP 해당 예산 전체(유실 0).

    미매핑으로 남는 건 I열이 '본　　사'이고 J열로도 처지사를 특정할 수 없는
    전표뿐이며(플랜트기술처 강제 귀속 폐지), 그만큼 '미반영'에 잡혀 총액은 보존된다.
    """
    for budget, universe in [("손익", 98554801), ("자본", 49709178)]:
        plan_path, cfg, out = _make_plan(tmp_path, budget)
        res = pipeline_actual.run_actual(
            plan_path, ZRFM2, MASTER, cfg, out, "2025", budget, new_policy="group",
        )
        assert os.path.exists(res["output_path"])
        assert os.path.exists(res["zrfm2_v1_path"])
        s = res["요약"]
        assert abs(s["총 실적(천원)"] + s["미반영(천원)"] - universe) <= ROUND_TOL
        assert s["계획집행 실적(천원)"] + s["신규 실적(천원)"] == s["총 실적(천원)"]
        assert abs(s["반올림차이(천원)"]) <= ROUND_TOL
        # 산출물 이름에 대상연도가 들어간다
        assert os.path.basename(res["zrfm2_v1_path"]) == f"zrfm2_2025_V1({budget}).xlsx"
        # 계정코드 동일 개명 통일로 '열원보완및개선'은 더 이상 미분류가 아님
        assert res["미분류과목"] == []
        # 미매핑은 '본사' 전표에 한정(지사명 정규화 실패가 아님)
        assert set(res["미매핑처지사"]) <= {"본　　사"}


def test_actual_sheet_layout_matches_reference(tmp_path):
    """데이터시트가 사용자 분석본 [Sheet1] 포맷(2행 헤더·3행 데이터, A~J 동일)이고,
    실적금액은 I열(9)이며 zrfm2_V1에 반영구분(20열)이 있다."""
    plan_path, cfg, out = _make_plan(tmp_path, "자본")
    res = pipeline_actual.run_actual(
        plan_path, ZRFM2, MASTER, cfg, out, "2025", "자본", new_policy="group",
    )
    wb = openpyxl.load_workbook(res["output_path"])
    ws = wb[ACTUAL_SHEET]
    header = [ws.cell(HEADER_ROW, c).value for c in range(1, len(ACTUAL_HEADER) + 1)]
    assert header == ACTUAL_HEADER
    # 분석본 Sheet1의 A~J 헤더와 정확히 일치해야 한다
    assert header[:10] == [
        "연번", "예산과목", "속성", "주관부서명", "예산귀속 부서명(처.지사)",
        "예산귀속 부서명(팀)", "사업명", "연예산(A)", "최종 실적금액(B)", "구 분",
    ]
    # 데이터 첫 행 연번=1, 실적금액 숫자
    assert ws.cell(DATA_START, 1).value == 1
    assert isinstance(ws.cell(DATA_START, 9).value, (int, float))
    wb.close()

    v1 = openpyxl.load_workbook(res["zrfm2_v1_path"])
    vs = v1[v1.sheetnames[0]]
    assert vs.cell(1, 20).value == "반영구분"
    v1.close()


def test_deselect_item_becomes_unreflected(tmp_path):
    """예산과목 하나(기계장치)를 포함 해제하면 그 실적이 종합표에서 빠지고
    동일 금액이 미반영으로 이동한다(총액=ERP 자본유니버스 보존)."""
    plan_path, cfg, out = _make_plan(tmp_path, "자본")
    base = pipeline_actual.run_actual(
        plan_path, ZRFM2, MASTER, cfg, out, "2025", "자본", new_policy="group",
    )
    base_total = base["요약"]["총 실적(천원)"]
    base_unref = base["요약"]["미반영(천원)"]

    # 과목구성에서 '기계장치' 포함 해제
    item_path = os.path.join(cfg, "과목구성_2025_자본.json")
    with open(item_path, "r", encoding="utf-8") as f:
        items = json.load(f)
    for it in items:
        if it["과목"] == "기계장치":
            it["포함"] = False
    with open(item_path, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)

    res = pipeline_actual.run_actual(
        plan_path, ZRFM2, MASTER, cfg, out, "2025", "자본", new_policy="group",
    )
    moved = base_total - res["요약"]["총 실적(천원)"]
    assert moved > 0                                   # 기계장치 실적이 종합표에서 빠짐
    # 빠진 만큼 미반영으로 이동 (반영+미반영 = 자본 유니버스 보존)
    assert abs((res["요약"]["미반영(천원)"] - base_unref) - moved) <= 2
    assert base_total + base_unref == pytest.approx(
        res["요약"]["총 실적(천원)"] + res["요약"]["미반영(천원)"], abs=2)

    # 데이터시트·종합표에 '기계장치'가 남지 않아야 함(선택된 항목만 포함)
    wb = openpyxl.load_workbook(res["output_path"], data_only=True)
    items_in_sheet = {wb[ACTUAL_SHEET].cell(r, 2).value
                      for r in range(DATA_START, wb[ACTUAL_SHEET].max_row + 1)}
    assert "기계장치" not in items_in_sheet
    summary_items = {wb["종합표"].cell(r, 2).value
                     for r in range(5, wb["종합표"].max_row + 1)}
    assert "기계장치" not in summary_items
    wb.close()
    # 그 전표들은 zrfm2_V1에서 '미반영'으로 표기
    v1 = openpyxl.load_workbook(res["zrfm2_v1_path"], data_only=True)
    vs = v1[v1.sheetnames[0]]
    assert any(vs.cell(r, 20).value == "미반영" for r in range(2, vs.max_row + 1))
    v1.close()


def test_strict_policy_more_new(tmp_path):
    plan_path, cfg, out = _make_plan(tmp_path, "손익")
    grp = pipeline_actual.run_actual(plan_path, ZRFM2, MASTER, cfg, out, "2025", "손익", new_policy="group")
    strt = pipeline_actual.run_actual(plan_path, ZRFM2, MASTER, cfg, out, "2025", "손익", new_policy="strict_name")
    # 총액은 정책 무관 동일(행 구성이 달라 행별 반올림 오차만 차이)
    assert abs(grp["요약"]["총 실적(천원)"] - strt["요약"]["총 실적(천원)"]) <= ROUND_TOL
    # 엄격정책은 신규 실적이 더 큼
    assert strt["요약"]["신규 실적(천원)"] > grp["요약"]["신규 실적(천원)"]
