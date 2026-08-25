# -*- coding: utf-8 -*-
"""실적 산출물 작성: {연도}년 {손익|자본}예산_실적.xlsx + zrfm2_V1.

- Sheet1: 사용자 분석본과 동일 포맷. 계획행(실적금액·집행구분 태깅) + 신규행 하단 추가.
- 종합표: 계획본과 동일 레이아웃, SUMIFS로 실적금액(I열)을 지사별 집계.
- 검토리포트: 미매핑 처지사, 미분류 과목, 연도 불일치 등.
- zrfm2_V1: 원본 유지 + R열에 매칭된 계획/신규 사업명.
"""
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, column_index_from_string

from . import loaders
from .config_store import SIMUI_THRESHOLD_THOUSAND
from . import excel_plan_writer as epw

# 사용자 분석본 `(23년) 손익예산 실적 분석.xlsx`의 [Sheet1] 포맷을 그대로 따른다.
#   A~J = 분석본과 동일(J '구 분'은 사용자가 손으로 채우는 과목 축약 분류),
#   K·L = 분석본의 무제 두 열(사업유형 / 분류키=J&K),
#   M~P = 우리 엔진의 분석열.
ACTUAL_SHEET = "Sheet1"
ACTUAL_HEADER = [
    "연번", "예산과목", "속성", "주관부서명", "예산귀속 부서명(처.지사)",
    "예산귀속 부서명(팀)", "사업명", "연예산(A)", "최종 실적금액(B)", "구 분",
    "사업유형", "분류키",
    "집행구분", "심의플래그", "비고", "매칭확신도",
]
NCOL = len(ACTUAL_HEADER)                 # 16
VALUE_COL_LETTER = "I"    # 9번 = 최종 실적금액(B)
ITEM_COL_LETTER = "B"     # 2번 = 예산과목
DEPT_COL_LETTER = "E"     # 5번 = 예산귀속 부서명(처.지사)
FLAG_COL_LETTER = "N"     # 14번 = 심의플래그
GUBUN_COL = 10            # J '구 분'(수기 분류)
TYPE_COL = 11             # K 사업유형(수기 분류)
KEY_COL = 12              # L 분류키 = J&K
EXEC_COL = 13             # M 집행구분(계획집행/미시행/신규)
FLAG_COL_INDEX = column_index_from_string(FLAG_COL_LETTER)
NEW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # 신규(노랑)
MISSING_EXEC_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")  # 미시행(연빨강)

# 분석본 Sheet1과 동일: 1행 비움, 2행 헤더, 3행부터 데이터.
HEADER_ROW = 2
DATA_START = 3


def _simui_flag(item_flags, item, base_amount):
    return "O" if item_flags.get(item) and base_amount is not None \
        and base_amount >= SIMUI_THRESHOLD_THOUSAND else ""


def _clean(val):
    return None if (val is None or (isinstance(val, float) and pd.isna(val))) else val


def write_actual_datasheet(ws, plan_rows, new_rows, budget, year, item_flags):
    """분석본 [Sheet1] 포맷으로 계획행 + 신규행을 기록.

    연예산(A)은 계획행만, 최종 실적금액(B)은 전 행. 신규행은 분석본과 같이
    연예산 0으로 두어 '계획 없이 집행된 사업'임을 나타낸다.
    """
    ws.cell(1, 9, "[단위 : 천원, 부가세 별도]")
    for i, h in enumerate(ACTUAL_HEADER, 1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = epw.HEADER_FILL

    r = DATA_START
    seq = 0
    for row in plan_rows:
        seq += 1
        ws.cell(r, 1, seq)
        ws.cell(r, 2, _clean(row.get("예산과목")))
        ws.cell(r, 3, _clean(row.get("속성")))
        ws.cell(r, 4, _clean(row.get("주관부서명")))
        ws.cell(r, 5, _clean(row.get("처지사")))
        ws.cell(r, 6, _clean(row.get("부서부")))
        ws.cell(r, 7, _clean(row.get("사업명")))
        ws.cell(r, 8, _clean(row.get("연예산")))
        ws.cell(r, 9, round(row.get("실적금액", 0.0)))
        ws.cell(r, KEY_COL, f"={get_column_letter(GUBUN_COL)}{r}&{get_column_letter(TYPE_COL)}{r}")
        ws.cell(r, EXEC_COL, row.get("구분"))
        ws.cell(r, FLAG_COL_INDEX,
                _simui_flag(item_flags, row.get("예산과목"), row.get("연예산")))
        if row.get("구분") == "계획집행":
            note = f"매칭전표 {row.get('매칭전표수', 0)}건"
            low = row.get("저유사전표수", 0)
            if low:
                note += f" (저유사 {low}건)"
            amb = row.get("임의귀속전표수", 0)
            if amb:
                note += f" ※임의귀속 {amb}건 — 사업 단위 배분 검증 필요"
        elif row.get("구분") == "미시행":
            note = "계획O·실적X (실적 없음)"
        else:
            note = ""
        ws.cell(r, NCOL - 1, note)
        conf = row.get("매칭확신도")
        ws.cell(r, NCOL, conf if conf is not None else None)
        if row.get("구분") == "미시행":
            for c in range(1, NCOL + 1):
                ws.cell(r, c).fill = MISSING_EXEC_FILL
        r += 1

    # 신규행(계획에 없는 집행) — 분석본과 같이 사업 단위 1행, 연예산 0
    for nr in new_rows:
        seq += 1
        ws.cell(r, 1, seq)
        ws.cell(r, 2, nr["예산과목"])
        ws.cell(r, 5, nr["처지사"])
        ws.cell(r, 6, _clean(nr.get("부서부")))     # ERP 부서명(부) 최빈값
        ws.cell(r, 7, nr["사업명"])
        ws.cell(r, 8, 0)
        ws.cell(r, 9, round(nr["실적금액"]))
        ws.cell(r, KEY_COL, f"={get_column_letter(GUBUN_COL)}{r}&{get_column_letter(TYPE_COL)}{r}")
        ws.cell(r, EXEC_COL, nr.get("구분", "신규"))
        ws.cell(r, FLAG_COL_INDEX,
                _simui_flag(item_flags, nr["예산과목"], nr["실적금액"]))
        ws.cell(r, NCOL - 1, nr.get("비고", ""))
        for c in range(1, NCOL + 1):
            ws.cell(r, c).fill = NEW_FILL
        r += 1

    # 가독성: 열 너비 · 금액 서식 · 헤더 고정
    widths = [6, 24, 6, 24, 18, 26, 46, 14, 14, 16, 14, 16, 12, 8, 44, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for rr in range(DATA_START, r):
        for cc in (8, 9):
            ws.cell(rr, cc).number_format = "#,##0"
    ws.freeze_panes = ws.cell(DATA_START, 1)

    return {"written_rows": r - DATA_START}


def write_review_sheet(wb, review):
    ws = wb.create_sheet("검토리포트")
    ws.append(["■ 실적분석 검토리포트"])
    ws.append([])
    ws.append(["대상연도", review.get("year")])
    ws.append(["ERP 연도분포", str(review.get("year_dist"))])
    if review.get("year_mismatch"):
        ws.append(["⚠ 연도 불일치", review["year_mismatch"]])
    ws.append([])

    ws.append(["■ 점검 필요 사항"])
    for w in review.get("warnings", []) or ["(없음)"]:
        ws.append([w])
    ws.append([])

    ws.append(["■ 예산과목별 계획 대비 실적 (종합표와 동일 모집단, 단위: 천원)"])
    ws.append(["예산과목", "계획(A)", "실적(B)", " └ 계획집행", " └ 신규",
               "집행률(%)", "계획행수", "집행행수", "미시행행수"])
    for row in review.get("item_table", []):
        ws.append(list(row))
    ws.append([])

    ws.append(["■ 계획행 제외 내역 — 이번 예산의 실적 대상이 아니어서 양식1에서 뺀 행"])
    ws.append(["사유", "행수", "연예산(천원)", "해당 예산과목"])
    for row in review.get("plan_excluded_rows", []):
        ws.append(list(row))
    ws.append([])

    ws.append(["■ 집계 제외 전표 — 합계행·타연도(zrfm2_V1 '반영구분'에 표기)"])
    ws.append(["사유", "금액(천원)", "전표수"])
    for row in review.get("excluded_rows", []):
        ws.append(list(row))
    ws.append([])
    ws.append([f"■ 미매핑 처지사 (종합표 집계 누락 {review.get('unmatched_dept_amt', 0):,}천원 — "
               "지사구성에 추가하거나 처지사_별칭.json 보강 필요)"])
    ws.append(["※ I열이 '본　　사'처럼 뭉뚱그려진 전표는 처지사_별칭.json에 "
               "\"I열|J열\": \"처지사\" 형태(예: \"본　　사|경영지원처 총무부\")로 지정할 수 있습니다."])
    detail = review.get("unmatched_dept_detail") or []
    if detail:
        ws.append(["지사명(I열)", "부서명(부)(J열)", "예산과목", "금액(천원)", "전표수"])
        for row in detail:
            ws.append(list(row))
    else:
        for d in review.get("unmatched_dept", []):
            ws.append([d])
    ws.append([])
    ws.append([f"■ 미분류(결측) 예산과목 — 사용자 예산과목 구성에 없음 ({review.get('unclassified_item_amt', 0):,}천원, 종합표 미반영)"])
    for it in review.get("unclassified_item", []):
        ws.append([it])
    ws.append([])
    ws.append([f"■ 미반영 전표 — 선택 안 된 예산과목·지사(구성 미포함/미분류/미매핑) ({review.get('unreflected_amt', 0):,}천원, 종합표·양식1 미포함)"])
    ws.append(["구분(과목/지사)", "금액(천원)", "전표수"])
    for row in review.get("unreflected_top", []):
        ws.append(list(row))
    ws.append([])
    ws.append(["■ 요약"])
    for k, v in review.get("summary", {}).items():
        ws.append([k, v])


def write_actual_workbook(out_path, plan_rows, new_rows, dept_columns, item_rows,
                          budget, year, review):
    item_flags = {it["과목"]: it["심의대상"] for cat in item_rows for it in cat["과목들"]}
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = ACTUAL_SHEET
    write_actual_datasheet(ds, plan_rows, new_rows, budget, year, item_flags)
    sm = wb.create_sheet("종합표")
    layout = epw.build_summary_layout(sm, dept_columns, item_rows, budget, year, kind="실적")
    epw.write_summary_formulas(
        sm, layout, value_col=VALUE_COL_LETTER, data_sheet=ACTUAL_SHEET,
        item_col=ITEM_COL_LETTER, dept_col=DEPT_COL_LETTER, flag_col=FLAG_COL_LETTER,
    )
    write_review_sheet(wb, review)
    wb.save(out_path)
    wb.close()


UNREFLECTED_FILL = PatternFill(start_color="E2E2E2", end_color="E2E2E2", fill_type="solid")  # 미반영(연회색)


def write_zrfm2_v1(src_path, out_path, erp_annotated):
    """원본 zrfm2 + R열(18) 매칭 사업명 + S열(19) 매칭확신도 + T열(20) 반영구분.

    반영구분: 계획집행 / 신규 / 신규(소액집행) / 미반영 / (공란=타예산).
    '미반영' = 선택 안 된 예산과목·지사, 미분류, 미매핑 처지사 전표(종합표 미포함).
    """
    label_by_row = {int(r): lab for r, lab in
                    zip(erp_annotated["_erp_row"], erp_annotated["매칭사업명"])}
    conf_by_row = {int(r): c for r, c in
                   zip(erp_annotated["_erp_row"], erp_annotated["매칭확신도"])}
    gubun_by_row = {int(r): g for r, g in
                    zip(erp_annotated["_erp_row"], erp_annotated["구분"])}
    wb = openpyxl.load_workbook(src_path)
    ws = wb[wb.sheetnames[0]]
    ws.cell(1, 18, "사업명(매칭)")
    ws.cell(1, 19, "매칭확신도")
    ws.cell(1, 20, "반영구분")
    for r in range(2, ws.max_row + 1):
        lab = label_by_row.get(r)
        if lab:
            ws.cell(r, 18, lab)
        c = conf_by_row.get(r)
        if c is not None:
            ws.cell(r, 19, c)
        g = gubun_by_row.get(r)
        if g:
            ws.cell(r, 20, g)
            if g == "미반영":
                ws.cell(r, 20).fill = UNREFLECTED_FILL
    wb.save(out_path)
    wb.close()


def write_zrfm2_v1_from_df(out_path, erp_annotated):
    """DB 모드: 원본 zrfm2 파일 없이 ERP DataFrame에서 zrfm2_V1을 생성.

    컬럼 = ERP 원장 주요 열 + 매칭 주석(사업명(매칭)·매칭확신도·반영구분).
    """
    cols = ["_erp_row", "계정코드", "예산과목원문", "연도", "전기일", "전표번호",
            "금액원", "사업명", "지사원문", "부서부원문", "사업장코드", "거래처코드",
            "매칭사업명", "매칭확신도", "구분"]
    header = ["행ID", "계정코드", "예산과목", "기간/연도", "전기일", "참조 전표 번호",
              "지급예산금액(원)", "텍스트(사업명)", "지사명", "부서명(부)",
              "사업장코드", "거래처코드", "사업명(매칭)", "매칭확신도", "반영구분"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "zrfm2_V1"
    for c, h in enumerate(header, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = epw.HEADER_FILL
    r = 2
    for _, row in erp_annotated.iterrows():
        for c, k in enumerate(cols, 1):
            v = row.get(k)
            ws.cell(r, c, None if (v is None or (isinstance(v, float) and pd.isna(v))) else v)
        g = row.get("구분")
        if g == "미반영":
            ws.cell(r, len(cols)).fill = UNREFLECTED_FILL
        r += 1
    ws.freeze_panes = "A2"
    ws.column_dimensions["H"].width = 44
    ws.column_dimensions["M"].width = 40
    for rr in range(2, r):
        ws.cell(rr, 7).number_format = "#,##0"
    wb.save(out_path)
    wb.close()


# matched_{year}.csv 컬럼(데이터 계약 §08). ERP 전표 + 매칭 결과.
MATCHED_CSV_COLUMNS = [
    "_erp_row", "계정코드", "예산과목원문", "과목정규", "지사원문", "처지사정규",
    "연도", "전표번호", "사업명", "금액원", "금액천원",
    "매칭사업명", "매칭확신도", "구분",
]


def write_matched_csv(out_path, erp_annotated):
    """전표-사업 매칭 결과를 CSV로 저장(데이터 계약: matched_{year}.csv).
    3·4단계 파이프라인/외부 도구가 소비. Excel 한글 호환 위해 utf-8-sig."""
    cols = [c for c in MATCHED_CSV_COLUMNS if c in erp_annotated.columns]
    erp_annotated.to_csv(out_path, columns=cols, index=False, encoding="utf-8-sig")
