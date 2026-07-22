# -*- coding: utf-8 -*-
"""ERP(SAP) zrfm2 실적 데이터 로더.

zrfm2 컬럼(위치 기준):
  B(2) 약정항목=계정코드, C(3) 약정항목텍스트=예산과목 원문, D(4) 기간/연도,
  E(5) FM 전기일=전기일(날짜), F(6) 참조 전표 번호, G(7) 지급예산금액(원),
  H(8) 텍스트=사업명, I(9) 이름=지사, J(10) 이름=부서명(부),
  L(12) 손익 센터=사업장 코드, O(15) 공급업체=거래처 코드.
금액 단위는 원 → 천원으로 환산(/1000). 부호(역인식 음수)는 그대로 순액 집계.
전기일/사업장코드/거래처코드는 연동규격(data_{연도}.csv) 매핑용 부가 컬럼(매칭 로직 비사용).
"""
import pandas as pd
import openpyxl

ERP_COLUMNS = [
    "계정코드", "예산과목원문", "연도", "전표번호",
    "금액원", "금액천원", "사업명", "지사원문", "부서부원문",
    "전기일", "사업장코드", "거래처코드", "_erp_row",
]

# zrfm2 시트 내 위치(1-index)
_COL = {
    "계정코드": 2, "예산과목원문": 3, "기간연도": 4, "전기일": 5, "전표번호": 6,
    "금액": 7, "사업명": 8, "지사": 9, "부서부": 10, "사업장": 12, "거래처": 15,
}


def _cell_str(row, pos):
    """1-index 위치의 셀을 문자열로(없으면 None). 범위를 벗어나면 None."""
    i = pos - 1
    if i >= len(row):
        return None
    v = row[i]
    return str(v).strip() if v is not None and str(v).strip() != "" else None


def _parse_year(v):
    """'2025.001' / 2025 / '2025' → '2025'."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    return s.split(".")[0].strip()


def _fmt_date(v):
    """전기일 → 'YYYY-MM-DD' 문자열(datetime/문자열 모두 허용). 실패 시 원문 유지."""
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return s[:10] if s else None


def load_erp(path, sheet=None):
    """zrfm2 워크북을 DataFrame(ERP_COLUMNS)으로 로드."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    recs = []
    r = 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        r += 1
        # 완전 빈 행 skip
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        amt = row[_COL["금액"] - 1]
        amt = float(amt) if isinstance(amt, (int, float)) else None
        raw_date = row[_COL["전기일"] - 1] if _COL["전기일"] - 1 < len(row) else None
        recs.append({
            "계정코드": _cell_str(row, _COL["계정코드"]),
            "예산과목원문": _cell_str(row, _COL["예산과목원문"]),
            "연도": _parse_year(row[_COL["기간연도"] - 1]),
            "전표번호": _cell_str(row, _COL["전표번호"]),
            "금액원": amt,
            "금액천원": (amt / 1000.0 if amt is not None else None),
            "사업명": _cell_str(row, _COL["사업명"]),
            "지사원문": _cell_str(row, _COL["지사"]),
            "부서부원문": _cell_str(row, _COL["부서부"]),
            "전기일": _fmt_date(raw_date),
            "사업장코드": _cell_str(row, _COL["사업장"]),
            "거래처코드": _cell_str(row, _COL["거래처"]),
            "_erp_row": r,
        })
    wb.close()
    return pd.DataFrame(recs, columns=ERP_COLUMNS)


def detect_years(df):
    """데이터에 등장한 연도 분포 {연도: 건수}."""
    out = {}
    for y in df["연도"]:
        if y:
            out[y] = out.get(y, 0) + 1
    return out
