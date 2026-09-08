"""예산 표준화 - 과년도 실적을 지사×예산과목×정비등급 단위로 통계 내어 표준 금액(벤치마크)을 만든다.

정비등급 이력(지사·연도별 MI/TI/CI/간이 또는 A/B/C 등)은 재무팀이 만든 참고 엑셀
(양식/표준화/★(최종본) 2026_예산표준화(25년실적 반영)_V3.xlsx 같은 파일)에서 가져올 수 있다.
표준 금액 산출 방식(최근실적/N개년평균/등급별평균)은 forecast.py의 방식-맵 패턴과 동일하게
사업장×예산과목별로 저장하고, 계산 결과는 화면에서 직접 수정(override)할 수 있다.
"""
import os
import pandas as pd
from openpyxl import load_workbook
from builtin_categories import PROFIT_LOSS_ACCOUNTS, CAPITAL_ACCOUNTS, normalize_account_name
from mapping_config import load_site_type_map

ALL_ACCOUNTS = PROFIT_LOSS_ACCOUNTS | CAPITAL_ACCOUNTS

GRADE_HISTORY_PATH = "site_year_grade.csv"  # 컬럼: 사업장, 연도, 등급
METHOD_MAP_PATH = "standardization_method_map.csv"  # 컬럼: 사업장, 예산과목, 방식
OVERRIDE_PATH = "standardization_overrides.csv"  # 컬럼: 사업장, 예산과목, 등급, 표준금액

METHOD_OPTIONS = ["등급별 평균", "최근실적", "3개년 평균", "5개년 평균"]
DEFAULT_METHOD = "등급별 평균"

# 자본예산 "기계장치" 계정 중 정기(LTSA/CRI) 성격의 투자유형세부는 등급별 표준화 대상에서 제외한다
# (연도별 집행계획으로 별도 산정 - 표준화방향 시트 설명 참고).
LTSA_LIKE_DETAILS = {"LTSA", "CRI"}

CURRENT_YEAR_FALLBACK = 2025  # 정비등급 이력이 하나도 없을 때만 쓰는 기본값


def reference_year_for_undated(grade_history_df: pd.DataFrame) -> int:
    """연도 정보가 없는 테스트모드 업로드분(예: '투자유형 예측 테스트' 결과)에 붙일 '현재' 연도.
    정비등급 이력의 가장 최근 연도를 쓴다 - 임의의 값(예: 9999)을 쓰면 엑셀 양식의 실제 연도 칸과
    매칭이 안 돼 연도별 실적 표가 비어 보이는 문제가 있다. 표준화·중장기예산 테스트 모드가 공유해서 쓴다."""
    if grade_history_df.empty:
        return CURRENT_YEAR_FALLBACK
    return int(grade_history_df["연도"].max())


# ---------------------------------------------------------------- 정비등급 이력
def load_grade_history() -> pd.DataFrame:
    if os.path.exists(GRADE_HISTORY_PATH):
        return pd.read_csv(GRADE_HISTORY_PATH, dtype={"사업장": str})
    return pd.DataFrame(columns=["사업장", "연도", "등급"])


def save_grade_history(df: pd.DataFrame):
    df.to_csv(GRADE_HISTORY_PATH, index=False)


def _site_sheet_name_map() -> dict:
    """엑셀 시트명(예: '화성') -> 사업장 표시명(예: '화성지사') 매핑.
    표준화 참고 엑셀의 지사 시트명은 site_type_map.csv 표시명에서 '지사'/'사업소' 접미어를 뗀 것과 같다."""
    site_map = load_site_type_map()
    mapping = {}
    if site_map.empty or "표시명" not in site_map.columns:
        return mapping
    for name in site_map["표시명"].dropna().unique():
        short = str(name).replace("지사", "").replace("사업소", "").strip()
        if short:
            mapping[short] = name
    return mapping


def import_grade_history_from_workbook(file_bytes) -> pd.DataFrame:
    """정비등급 이력이 채워진 표준화 참고 엑셀(예: V3본)에서 지사×연도×등급을 읽어온다.

    각 지사 시트는 '손익소계' 행 다음에 보조소계(빈 라벨) 행이 하나 있고, 그 다음
    한두 개의 장비별 정비등급 이력 행(예: '#1 GT'/'#2 GT'/'ST', 또는 DH의 '등급')이
    이어진다. 라벨에 'GT'가 들어간 행을 그 지사의 대표 등급 이력으로 쓰고
    (실제 참고 엑셀에서도 Standization 표가 GT 계열 등급만 축으로 쓴다), 없으면
    첫 번째 장비 행을 쓴다.
    """
    import io
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_to_site = _site_sheet_name_map()

    rows = []
    for sheet_name in wb.sheetnames:
        site = sheet_to_site.get(sheet_name.strip())
        if site is None:
            continue
        ws = wb[sheet_name]

        year_row = None
        for r in range(1, min(ws.max_row, 10) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(2, ws.max_column + 1)]
            if any(isinstance(v, (int, float)) and 2000 < v < 2100 for v in vals):
                year_row = r
                break
        if year_row is None:
            continue
        year_by_col = {
            c: int(ws.cell(row=year_row, column=c).value)
            for c in range(2, ws.max_column + 1)
            if isinstance(ws.cell(row=year_row, column=c).value, (int, float))
        }

        subtotal_row = None
        for r in range(year_row + 1, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value or "").strip() == "손익소계":
                subtotal_row = r
                break
        if subtotal_row is None:
            continue

        equipment_rows = []
        r = subtotal_row + 2  # 손익소계 다음(보조소계)은 건너뛴다
        while r <= ws.max_row:
            label = ws.cell(row=r, column=1).value
            if label is None or str(label).strip() == "":
                break
            equipment_rows.append((r, str(label).strip()))
            r += 1
        if not equipment_rows:
            continue

        chosen = next((rl for rl in equipment_rows if "GT" in rl[1]), equipment_rows[0])
        chosen_row = chosen[0]
        for col, year in year_by_col.items():
            grade = ws.cell(row=chosen_row, column=col).value
            if grade is None:
                continue
            grade = str(grade).strip()
            if not grade or grade == "-":
                continue
            rows.append({"사업장": site, "연도": year, "등급": grade})

    return pd.DataFrame(rows, columns=["사업장", "연도", "등급"])


def import_actuals_from_workbook(file_bytes) -> pd.DataFrame:
    """정비등급 이력과 같은 표준화 참고 엑셀에서 지사×연도×예산과목 실적 금액도 함께 읽어온다.

    각 지사 시트는 연도 헤더 행 다음부터 '손익소계' 행까지 자본/손익 라인아이템이 나열되어
    있다(중간의 '자본소계'는 소계 행이라 건너뛴다). 라벨이 ALL_ACCOUNTS와 매칭되는 행만 가져온다.
    양식 단위는 천원이라 원 단위로 환산한다."""
    import io
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_to_site = _site_sheet_name_map()

    rows = []
    for sheet_name in wb.sheetnames:
        site = sheet_to_site.get(sheet_name.strip())
        if site is None:
            continue
        ws = wb[sheet_name]

        year_row = None
        for r in range(1, min(ws.max_row, 10) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(2, ws.max_column + 1)]
            if any(isinstance(v, (int, float)) and 2000 < v < 2100 for v in vals):
                year_row = r
                break
        if year_row is None:
            continue
        year_by_col = {
            c: int(ws.cell(row=year_row, column=c).value)
            for c in range(2, ws.max_column + 1)
            if isinstance(ws.cell(row=year_row, column=c).value, (int, float))
        }

        subtotal_row = None
        for r in range(year_row + 1, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value or "").strip() == "손익소계":
                subtotal_row = r
                break
        last_row = subtotal_row if subtotal_row else ws.max_row

        for r in range(year_row + 1, last_row + 1):
            label = ws.cell(row=r, column=1).value
            if label is None:
                continue
            label = str(label).strip()
            if not label or label in ("자본소계", "손익소계"):
                continue
            normalized = normalize_account_name(label)
            if normalized not in ALL_ACCOUNTS:
                continue
            for col, year in year_by_col.items():
                amount = ws.cell(row=r, column=col).value
                if amount in (None, ""):
                    continue
                try:
                    amount = float(amount)
                except (TypeError, ValueError):
                    continue
                if amount == 0:
                    continue
                rows.append({"사업장": site, "연도": year, "예산과목": normalized, "금액": amount * 1000})

    return pd.DataFrame(rows, columns=["사업장", "연도", "예산과목", "금액"])


# ---------------------------------------------------------------- 산출방식 맵
def load_method_map() -> pd.DataFrame:
    if os.path.exists(METHOD_MAP_PATH):
        return pd.read_csv(METHOD_MAP_PATH, dtype={"사업장": str})
    return pd.DataFrame(columns=["사업장", "예산과목", "방식"])


def save_method_map(df: pd.DataFrame):
    df.to_csv(METHOD_MAP_PATH, index=False)


def get_method_dict(method_map: pd.DataFrame) -> dict:
    """{(사업장, 예산과목): 방식} 딕셔너리로 변환. 없는 조합은 호출측에서 DEFAULT_METHOD를 쓰면 된다."""
    if method_map.empty:
        return {}
    return {(r["사업장"], r["예산과목"]): r["방식"] for _, r in method_map.iterrows()}


# ---------------------------------------------------------------- 수동 수정값
def load_overrides() -> pd.DataFrame:
    if os.path.exists(OVERRIDE_PATH):
        return pd.read_csv(OVERRIDE_PATH, dtype={"사업장": str})
    return pd.DataFrame(columns=["사업장", "예산과목", "등급", "표준금액"])


def save_overrides(df: pd.DataFrame):
    df.to_csv(OVERRIDE_PATH, index=False)


# ---------------------------------------------------------------- 계산
def _apply_method(yearly: pd.Series, grade: str, grade_by_year: dict, method: str) -> tuple:
    """yearly: 연도(index)->금액. 반환: (표준금액, 비고)"""
    if yearly.empty:
        return 0.0, "실적없음"

    if method == "최근실적":
        latest_year = yearly.index.max()
        return float(yearly.loc[latest_year]), "최근실적"

    if method == "3개년 평균":
        recent = yearly.sort_index().iloc[-3:]
        return float(recent.mean()), "3개년 평균"

    if method == "5개년 평균":
        recent = yearly.sort_index().iloc[-5:]
        return float(recent.mean()), "5개년 평균"

    # 기본: 등급별 평균 - 이 등급이 적용된 연도만 평균, 없으면 전체 평균으로 대체
    matching_years = [y for y, g in grade_by_year.items() if g == grade and y in yearly.index]
    if matching_years:
        return float(yearly.loc[matching_years].mean()), "등급별 평균"
    return float(yearly.mean()), "등급별 평균(해당등급 실적없음 - 전체평균 대체)"


ETA_SQUARED_THRESHOLD = 0.3  # 이 이상이면 '등급이 금액 차이를 설명한다'고 본다
CV_THRESHOLD = 0.15  # 변동계수(표준편차/평균)가 이 미만이면 '변동이 작다'고 본다


def recommend_method(yearly: pd.Series, grade_by_year: dict) -> tuple:
    """항목별 실적 패턴(연도 수, 변동성, 등급별 편차)을 통계적으로 분석해 산출방식을 추천한다.
    사람이 미리 방식을 정해둔 항목이 아닐 때 기본값으로 쓰인다.

    - 등급별로 나눴을 때 그룹 간 평균 차이가 전체 분산의 상당 부분을 설명하면(η² 기준)
      '등급별 평균'을 추천한다 - 실제로 등급이 금액을 가르는 경우다.
    - 등급 차이가 뚜렷하지 않아도 연도별 변동이 작으면 '최근실적'으로 충분하다고 본다.
    - 등급도 안 갈리고 변동도 있으면, 데이터가 충분할 때(5개년 이상) '5개년 평균'으로
      완만화하고, 데이터가 적으면 '3개년 평균'을 추천한다.

    반환: (방식, 추천사유)
    """
    n_years = len(yearly)
    if n_years == 0:
        return DEFAULT_METHOD, "실적 없음"
    if n_years == 1:
        return "최근실적", "실적이 1개년뿐이라 그대로 사용"

    mean = yearly.mean()
    std = yearly.std(ddof=0)
    cv = (std / mean) if mean else 0.0

    groups = {}
    for year, amount in yearly.items():
        grade = grade_by_year.get(year)
        if grade is not None:
            groups.setdefault(grade, []).append(amount)
    multi_grade_groups = {g: v for g, v in groups.items() if v}

    eta_sq = None
    if len(multi_grade_groups) >= 2:
        ss_total = ((yearly - mean) ** 2).sum()
        if ss_total > 0:
            ss_between = sum(len(v) * (sum(v) / len(v) - mean) ** 2 for v in multi_grade_groups.values())
            eta_sq = ss_between / ss_total

    if eta_sq is not None and eta_sq >= ETA_SQUARED_THRESHOLD:
        return "등급별 평균", f"등급별 그룹 간 편차가 뚜렷함(설명력 {eta_sq*100:.0f}%) - 등급별로 나눠서 계산"

    if cv < CV_THRESHOLD:
        return "최근실적", f"연도별 변동이 작음(변동계수 {cv*100:.0f}%) - 최근실적만으로 충분"

    if n_years >= 5:
        return "5개년 평균", f"등급 차이는 뚜렷하지 않고 변동성이 있음(변동계수 {cv*100:.0f}%) - 5개년 평균으로 완만화"

    return "3개년 평균", f"데이터가 {n_years}개년뿐이고 변동성이 있음(변동계수 {cv*100:.0f}%) - 3개년 평균"


def get_current_grade(site: str, grade_history_df: pd.DataFrame):
    """해당 지사의 정비등급 이력 중 가장 최근 연도의 등급(없으면 None).
    연도 정보가 없는 데이터(예: '투자유형 예측 테스트' 업로드분)에 임시 등급을 붙일 때 쓴다."""
    site_hist = grade_history_df[grade_history_df["사업장"] == site]
    if site_hist.empty:
        return None
    latest_year = site_hist["연도"].max()
    return site_hist.loc[site_hist["연도"] == latest_year, "등급"].iloc[0]


def has_ltsa_detail(actuals_df: pd.DataFrame) -> bool:
    """'기계장치'에서 LTSA/CRI를 실제로 분리할 수 있는지(투자유형세부_확정 컬럼 존재 여부)."""
    return "투자유형세부_확정" in actuals_df.columns


def compute_standard_amounts(actuals_df: pd.DataFrame, grade_history_df: pd.DataFrame,
                              method_map: pd.DataFrame = None, overrides: pd.DataFrame = None) -> pd.DataFrame:
    """
    actuals_df: 사업장, 연도, 예산과목, 금액 컬럼 필요. '투자유형세부_확정' 컬럼이 있으면
                기계장치 항목에서 LTSA/CRI 성격의 실적을 제외한 '기타기계장치'만 표준화 대상으로 삼는다.
    grade_history_df: 사업장, 연도, 등급 (load_grade_history()/import_grade_history_from_workbook() 결과)
    method_map에 사람이 지정한 방식이 없는 (사업장,예산과목)은 recommend_method()가 데이터 패턴을
    분석해 자동으로 방식을 고른다(방식출처="자동추천").
    반환: 사업장, 구분(손익/자본), 예산과목, 등급, 표준금액, 산출방식, 방식출처, 추천사유, 비고
    """
    result_cols = ["사업장", "구분", "예산과목", "등급", "표준금액", "산출방식", "방식출처", "추천사유", "비고"]
    if actuals_df.empty:
        return pd.DataFrame(columns=result_cols)

    df = actuals_df.copy()
    df["예산과목"] = df["예산과목"].apply(normalize_account_name)
    df = df[df["예산과목"].isin(ALL_ACCOUNTS)]

    if "투자유형세부_확정" in df.columns:
        machine_mask = df["예산과목"] == "기계장치"
        ltsa_mask = machine_mask & df["투자유형세부_확정"].isin(LTSA_LIKE_DETAILS)
        df = df[~ltsa_mask]
    # 컬럼이 없으면(정식 classified_*.csv 모드) '기계장치'를 LTSA/CRI 분리 없이 총액 그대로 쓴다.
    # 이 경우의 안내는 화면 쪽에서 한 번만 표시한다(has_ltsa_detail() 참고).

    method_dict = get_method_dict(method_map) if method_map is not None else {}
    override_dict = {}
    if overrides is not None and not overrides.empty:
        override_dict = {(r["사업장"], r["예산과목"], r["등급"]): r["표준금액"] for _, r in overrides.iterrows()}

    rows = []
    for (site, account), g in df.groupby(["사업장", "예산과목"]):
        category = "손익" if account in PROFIT_LOSS_ACCOUNTS else "자본"
        yearly = g.groupby("연도")["금액"].sum()

        site_grades = grade_history_df[grade_history_df["사업장"] == site]
        grade_by_year = dict(zip(site_grades["연도"], site_grades["등급"]))
        grades_present = sorted(set(grade_by_year.values())) or ["표준"]

        explicit_method = method_dict.get((site, account))
        if explicit_method:
            method, method_source, reco_reason = explicit_method, "사용자지정", ""
        else:
            method, reco_reason = recommend_method(yearly, grade_by_year)
            method_source = "자동추천"

        for grade in grades_present:
            amount, note = _apply_method(yearly, grade, grade_by_year, method)

            override_key = (site, account, grade)
            is_override = override_key in override_dict
            if is_override:
                amount = override_dict[override_key]
                note = "사용자가 직접 입력"

            rows.append({
                "사업장": site, "구분": category, "예산과목": account, "등급": grade,
                "표준금액": amount,
                "산출방식": "수동수정" if is_override else method,
                "방식출처": "수동수정" if is_override else method_source,
                "추천사유": reco_reason if (not is_override and method_source == "자동추천") else "",
                "비고": note,
            })

    return pd.DataFrame(rows, columns=result_cols) if rows else pd.DataFrame(columns=result_cols)
