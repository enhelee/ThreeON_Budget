"""
손익예산/자본예산 실적집계표 양식(내장)에 계획대비실적 결과를 채워 넣어 다운로드용 엑셀을 만든다.
양식 파일을 사용자가 매번 업로드할 필요 없이, 앱에 내장된 두 양식(손익/자본)을 바로 사용한다.
필요하면 추가 양식을 더 등록할 수도 있다.
"""
import os
import json
from copy import copy
import pandas as pd
from openpyxl import load_workbook
from plan_vs_actual import build_plan_vs_actual

MANIFEST_PATH = "templates_manifest.json"

# 앱에 내장된 기본 양식 (사용자가 업로드하지 않아도 바로 사용 가능)
BUILTIN_TEMPLATES = {
    "손익": "builtin_template_손익.xlsx",
    "자본": "builtin_template_자본.xlsx",
}

HEADER_ROW = 2
DATA_START_ROW = 4  # 3행은 '예시)' 행이라 보존하고 4행부터 채운다

BUDGET_TO_TEMPLATE_COLS = {
    "예산과목": "예산과목",
    "속성": "속성",
    "주관부서명": "주관부서명",
    "예산귀속\n부서명(처.지사)": "예산귀속\n부서명(처.지사)",
    "예산귀속 부서명(부)": "예산귀속\n부서명(팀)",
}
TEMPLATE_COL_ORDER = ["연번", "예산과목", "속성", "주관부서명", "예산귀속\n부서명(처.지사)",
                      "예산귀속\n부서명(팀)", "사업명", "연예산(A)", "최종 실적금액(B)"]


def _load_manifest() -> dict:
    if os.path.exists(MANIFEST_PATH):
        with open(MANIFEST_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {}


def _save_manifest(manifest: dict):
    with open(MANIFEST_PATH, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)


def list_templates() -> dict:
    """내장 양식(손익/자본) + 사용자가 추가 등록한 양식을 합쳐서 {표시이름: 파일경로} 반환"""
    templates = {}
    for name, path in BUILTIN_TEMPLATES.items():
        if os.path.exists(path):
            templates[name] = path
    templates.update(_load_manifest())
    return templates


def has_templates() -> bool:
    return len(list_templates()) > 0


def save_template(display_name: str, file_bytes) -> str:
    """추가 양식을 이름 붙여 저장(손익/자본 내장 양식과 별개로 더 등록하고 싶을 때 사용)"""
    safe = "".join(c for c in display_name if c.isalnum() or c in " _-()").strip() or "양식"
    path = f"template_{safe}.xlsx"
    with open(path, "wb") as f:
        f.write(file_bytes)
    manifest = _load_manifest()
    manifest[display_name] = path
    _save_manifest(manifest)
    return path


def delete_template(display_name: str):
    if display_name in BUILTIN_TEMPLATES:
        return  # 내장 양식은 삭제 불가
    manifest = _load_manifest()
    path = manifest.pop(display_name, None)
    if path and os.path.exists(path):
        os.remove(path)
    _save_manifest(manifest)


def _find_template_sheet(wb):
    """'(작성필요)25년 계획 대비 실적' 또는 '25년 계획 대비 실적' 등 이름 차이를 흡수해서 찾는다."""
    for name in wb.sheetnames:
        if "계획 대비 실적" in name:
            return name
    return None


def build_output_rows(year: int, category_filter: str = None) -> pd.DataFrame:
    """사업명 기준 메타정보 + 연예산(A) + 최종실적금액(B)을 합친 표 (손익/자본 필터 가능)"""
    budget_path = f"budget_{year}.csv"
    matched_path = f"matched_{year}.csv"

    if not os.path.exists(budget_path):
        return pd.DataFrame()

    pva = build_plan_vs_actual(year, category_filter=category_filter)
    if pva.empty:
        return pva

    budget = pd.read_csv(budget_path)
    meta_source_cols = [c for c in BUDGET_TO_TEMPLATE_COLS if c in budget.columns]
    grouped_meta = budget.groupby("사업명")[meta_source_cols].first().reset_index() if meta_source_cols else \
        budget[["사업명"]].drop_duplicates().reset_index(drop=True)

    merged = pva.merge(grouped_meta, on="사업명", how="left")

    # 예산계획에 없는 사업명(신규 등록·소액 자동확정·미분류묶음 등)은
    # '예산과목'과 '예산귀속 부서명(처.지사)'이 비어있다. 종합표(고정 컬럼 기준 집계)에서
    # 이 항목들이 누락되지 않도록, 실적데이터의 계정과목/사업장 정보를 대신 채워넣는다.
    dept_col = "예산귀속\n부서명(처.지사)"
    if os.path.exists(matched_path):
        matched_raw = pd.read_csv(matched_path)

        if "예산과목" in merged.columns and merged["예산과목"].isna().any() and "계정과목" in matched_raw.columns:
            acct_by_project = matched_raw.groupby("사업명")["계정과목"].first()
            missing_mask = merged["예산과목"].isna()
            merged.loc[missing_mask, "예산과목"] = merged.loc[missing_mask, "사업명"].map(acct_by_project)

        if dept_col in merged.columns and merged[dept_col].isna().any() and "사업장" in matched_raw.columns:
            from mapping_config import get_site_display_name
            site_by_project = matched_raw.groupby("사업명")["사업장"].first().apply(get_site_display_name)
            missing_mask = merged[dept_col].isna()
            merged.loc[missing_mask, dept_col] = merged.loc[missing_mask, "사업명"].map(site_by_project)

    # 예산계획에 부서'코드'가 그대로 남아있는 경우(예: '2023.0')도 실제 지사명으로 정규화한다.
    # 종합표는 24개 지사/부서 '이름'을 기준으로 집계하므로, 코드 그대로 남아있으면 누락된다.
    # 등록되지 않은 코드라도 앞 3자리가 같은 코드가 있으면 같은 지사로 묶인다.
    if dept_col in merged.columns:
        from mapping_config import get_site_display_name
        merged[dept_col] = merged[dept_col].apply(lambda v: get_site_display_name(v) if pd.notna(v) else v)

    merged = merged.rename(columns={
        "최종실적금액(B)": "최종 실적금액(B)",
        **{k: v for k, v in BUDGET_TO_TEMPLATE_COLS.items() if k in merged.columns and k != v},
    })
    return merged.sort_values("연예산(A)", ascending=False).reset_index(drop=True)


def _find_last_styled_row(ws, start_row: int = DATA_START_ROW, num_cols: int = 9) -> int:
    """양식에 원래 서식(배경색, 줄 높이 등)이 입혀진 마지막 행을 찾는다. 그 이후 행은 서식이 없다."""
    last = start_row - 1
    for r in range(start_row, ws.max_row + 1):
        styled = ws.row_dimensions[r].height is not None
        if not styled:
            for c in range(1, num_cols + 1):
                cell = ws.cell(row=r, column=c)
                if cell.fill and cell.fill.patternType is not None:
                    styled = True
                    break
        if styled:
            last = r
    return last


def _copy_row_style(ws, src_row: int, dst_row: int, num_cols: int = 9):
    """src_row의 서식(배경색·글꼴·테두리·정렬·줄높이)을 dst_row에 그대로 복사한다."""
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for c in range(1, num_cols + 1):
        src_cell = ws.cell(row=src_row, column=c)
        dst_cell = ws.cell(row=dst_row, column=c)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format


def get_known_locations(template_name: str, header_row_candidates: tuple = (3, 5)) -> set:
    """
    등록된 양식의 종합표에서 SUMIFS 기준으로 쓰이는 지사/부서명 목록을 그대로 읽어온다.
    이 목록에 없는 '예산귀속 부서명(처.지사)' 값을 가진 항목은 종합표에서 집계되지 않는다.
    """
    templates = list_templates()
    if template_name not in templates:
        return set()
    wb = load_workbook(templates[template_name])
    if "(작성불필요)종합표" not in wb.sheetnames:
        return set()
    ws = wb["(작성불필요)종합표"]

    locations = set()
    for header_row in header_row_candidates:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v and str(v).strip() not in ("소계", "구 분", ""):
                locations.add(str(v).strip())
    return locations


def diagnose_unmatched_locations(year: int, template_name: str, category_filter: str = None) -> dict:
    """
    build_output_rows 결과 중, 종합표가 인식하는 지사/부서명 목록에 없는 항목을 찾아
    금액 합계와 목록을 반환한다. (종합표에서 누락되는 항목 진단용)
    """
    rows_df = build_output_rows(year, category_filter=category_filter)
    if rows_df.empty:
        return {"unmatched_amount": 0, "unmatched_count": 0, "items": pd.DataFrame(), "known_locations": set()}

    known = get_known_locations(template_name)
    dept_col = "예산귀속\n부서명(처.지사)"
    if dept_col not in rows_df.columns or not known:
        return {"unmatched_amount": 0, "unmatched_count": 0, "items": pd.DataFrame(), "known_locations": known}

    unmatched_mask = ~rows_df[dept_col].astype(str).str.strip().isin(known)
    unmatched = rows_df[unmatched_mask]
    return {
        "unmatched_amount": unmatched["최종 실적금액(B)"].sum() if "최종 실적금액(B)" in unmatched.columns else 0,
        "unmatched_count": len(unmatched),
        "items": unmatched[["사업명", dept_col, "최종 실적금액(B)"]] if not unmatched.empty else pd.DataFrame(),
        "known_locations": known,
    }


def fill_template(year: int, template_name: str, output_path: str, category_filter: str = None) -> str:
    """등록된 양식 중 하나(template_name)를 열어 데이터 행을 채우고 새 파일로 저장한다."""
    templates = list_templates()
    if template_name not in templates:
        raise FileNotFoundError(f"'{template_name}' 양식을 찾을 수 없습니다.")
    template_path = templates[template_name]
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"'{template_name}' 양식 파일을 찾을 수 없습니다.")

    rows_df = build_output_rows(year, category_filter=category_filter)
    wb = load_workbook(template_path)
    sheet_name = _find_template_sheet(wb)
    if sheet_name is None:
        raise ValueError("양식에서 '계획 대비 실적' 시트를 찾을 수 없습니다.")
    ws = wb[sheet_name]

    # 시스템 내부는 원(원) 단위로 저장하지만, 이 양식은 '[단위: 천원]'을 쓰므로 채울 때 다시 나눠준다.
    AMOUNT_COLS_IN_TEMPLATE = ["연예산(A)", "최종 실적금액(B)"]

    last_styled_row = _find_last_styled_row(ws)

    row_idx = DATA_START_ROW
    for seq, (_, row) in enumerate(rows_df.iterrows(), start=1):
        if row_idx > last_styled_row:
            _copy_row_style(ws, last_styled_row, row_idx)

        ws.cell(row=row_idx, column=1, value=seq)
        for col_idx, col_name in enumerate(TEMPLATE_COL_ORDER[1:], start=2):
            value = row.get(col_name)
            if pd.isna(value):
                value = None
            elif col_name in AMOUNT_COLS_IN_TEMPLATE:
                value = value / 1000
            ws.cell(row=row_idx, column=col_idx, value=value)
        row_idx += 1

    wb.save(output_path)
    return output_path
